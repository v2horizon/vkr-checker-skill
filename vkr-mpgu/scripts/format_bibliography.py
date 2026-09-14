#!/usr/bin/env python3
"""Форматирование списка литературы по образцам методички МПГУ (vkr-mpgu 6.33).

Пунктуация записей следует примерам методички и близка к ГОСТ 7.1-2003;
полной автоматической реализацией ГОСТ Р 7.0.100-2018 скрипт не является.

Вход — реестр источников ``sources.json`` (раздел 3 SPEC, схема описана в
``references/gost-citations.md``): массив объектов или объект с ключом
``sources``. Выход — нумерованный список в алфавитном порядке по методичке:
по первому элементу записи (фамилия первого автора, иначе заглавие),
кириллица раньше латиницы, «ё» = «е», регистр не важен; при равенстве —
заглавие, затем год. Нормативные акты сортируются вместе со всеми; опция
``--group-normative-first`` ставит нормативные акты и стандарты в начало.

Использование:
    python format_bibliography.py sources.json
    python format_bibliography.py sources.json -o exports/bibliography.txt
    python format_bibliography.py sources.json -o exports/bibliography.docx
    python format_bibliography.py sources.json --json --mapping exports/map.json

Коды выхода: 0 — список сформирован (предупреждения допустимы);
1 — ошибки в записях реестра (запись нельзя оформить, повторяющиеся id);
2 — ошибка использования, чтения файла или отсутствует зависимость.

Номера ссылок ``[N]`` в тексте работы после сортировки меняются. Черновики
``drafts/*.md`` со ссылками ``[@id]`` пересчитывает ``build_vkr.py``
автоматически; ``--mapping`` нужен только для ручных сценариев.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

TOOL_VERSION = "6.33"


class BibliographyError(ValueError):
    """Ошибка во входных данных, которую нужно показать без трассировки."""


# ========== ТИПЫ И АЛИАСЫ ==========

CANONICAL_TYPES = (
    "book", "article", "collection_article", "conference", "newspaper",
    "electronic", "dissertation", "autoreferat", "normative", "standard",
    "software_doc", "repository", "preprint", "dataset", "patent", "video",
)

TYPE_ALIASES = {
    "книга": "book", "monograph": "book", "монография": "book", "учебник": "book",
    "статья": "article", "journal_article": "article",
    "collection": "collection_article", "сборник": "collection_article",
    "chapter": "collection_article", "book_chapter": "collection_article",
    "конференция": "conference", "proceedings": "conference",
    "газета": "newspaper",
    "web": "electronic", "online": "electronic", "internet": "electronic",
    "url": "electronic", "website": "electronic", "site": "electronic",
    "электронный": "electronic", "электронный_ресурс": "electronic",
    "thesis": "dissertation", "дисс": "dissertation", "диссертация": "dissertation",
    "autoref": "autoreferat", "автореферат": "autoreferat",
    "нормативный": "normative", "закон": "normative", "law": "normative",
    "gost": "standard", "гост": "standard", "стандарт": "standard",
    "documentation": "software_doc", "docs": "software_doc", "документация": "software_doc",
    "repo": "repository", "github": "repository", "репозиторий": "repository",
    "arxiv": "preprint", "препринт": "preprint",
    "датасет": "dataset", "набор_данных": "dataset",
    "патент": "patent",
    "видео": "video", "youtube": "video",
}

ELECTRONIC_TYPES = {"electronic", "software_doc", "repository", "preprint", "dataset", "video"}
NORMATIVE_GROUP = {"normative", "standard"}

# Уточнение после маркера ресурса, если в записи нет поля title_note.
DEFAULT_TITLE_NOTES = {"repository": "репозиторий", "preprint": "препринт", "dataset": "датасет"}

# Алиасы полей — распространённые синонимы из пользовательских JSON.
FIELD_ALIASES = {
    "city": "place", "город": "place", "место": "place",
    "issue": "number", "номер": "number", "выпуск": "number",
    "pages_range": "pages", "pages_total": "pages", "страницы": "pages", "стр": "pages",
    "name": "title", "название": "title", "заглавие": "title",
    "год": "year",
    "автор": "authors", "авторы": "authors", "author": "authors",
    "editor": "editors", "редактор": "editors", "редакторы": "editors",
    "journal_name": "journal", "журнал": "journal",
    "издательство": "publisher",
    "ссылка": "url", "link": "url",
    "accessed": "access_date", "accessed_date": "access_date", "дата_обращения": "access_date",
    "container": "journal", "контейнер": "journal", "платформа": "journal", "platform": "journal",
    "тип_диссертации": "thesis_type",
    "lang": "language", "язык": "language",
}

STATUS_VALUES = ("pending", "confirmed", "suspicious", "rejected")

ID_RE = re.compile(r"^[A-Za-z0-9_.:-]+$")


def canonical_type(src: Dict[str, Any]) -> str:
    raw = str(src.get("type") or "book").strip()
    key = raw.casefold()
    if key in CANONICAL_TYPES:
        return key
    return TYPE_ALIASES.get(key, key)


def normalize_status(value: Any) -> str:
    """Статус записи реестра; ``verified`` читается как ``confirmed``."""
    text = str(value or "pending").strip().casefold()
    if text == "verified":
        return "confirmed"
    return text


def is_structured(src: Dict[str, Any]) -> bool:
    """Запись можно оформить, только если есть заглавие (не только raw_text)."""
    return bool(str(src.get("title") or "").strip())


# ========== ЛЮДИ И ОРГАНИЗАЦИИ ==========

# Буквы — любые Unicode-буквы: латиница с диакритикой («Rößling G.», «Müller J.», «Łukasz K.») разбирается
# так же, как кириллица и базовая латиница. Инициал — заглавная буква с точкой (проверяется в _initials_ok).
_LETTER = r"[^\W\d_]"
_INITIAL = _LETTER + r"\.(?:\s?-?" + _LETTER + r"\.)?"
_SURNAME = _LETTER + r"(?:" + _LETTER + r"|['’\-])*"


def _clean_initials(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def _initials_ok(value: str) -> bool:
    return all(ch.isupper() for ch in value if ch.isalpha())


def _parse_person_string(text: str) -> Optional[Dict[str, Any]]:
    """Разбирает «Иванов А.Б.», «Иванов, А. Б.», «А.Б. Иванов», «Иванов Иван Иванович», «Rößling G.»."""
    s = re.sub(r"\s+", " ", str(text or "")).strip()
    if not s:
        return None
    m = re.match(r"^(" + _SURNAME + r"),?\s+(" + _INITIAL + r")$", s)
    if m and _initials_ok(m.group(2)):
        return {"last": m.group(1), "initials": _clean_initials(m.group(2))}
    m = re.match(r"^(" + _INITIAL + r")\s*(" + _SURNAME + r")$", s)
    if m and _initials_ok(m.group(1)):
        return {"last": m.group(2), "initials": _clean_initials(m.group(1))}
    m = re.match(r"^([А-ЯЁ][а-яё\-]+)\s+([А-ЯЁ])[а-яё]+\s+([А-ЯЁ])[а-яё]+$", s)
    if m and re.search(r"(вич|вна|чна|ична|оглы|кызы)$", s):
        return {"last": m.group(1), "initials": f"{m.group(2)}.{m.group(3)}."}
    if re.match(r"^\S+$", s) and (re.search(r"\d", s) or re.search(r"[a-zа-яё][A-ZА-ЯЁ]", s)
                                  or re.match(r"^[A-ZА-ЯЁ&]{2,}$", s)):
        return {"last": s, "initials": "", "is_organization": True}  # W3C, IEEE, JetBrains, 1С
    m = re.match(r"^(" + _SURNAME + r")$", s)
    if m:
        return {"last": m.group(1), "initials": ""}
    if re.search(r"[А-ЯЁа-яёA-Za-z]", s) and not re.search(r"(?<![А-ЯЁа-яёA-Za-z])[А-ЯЁA-Z]\.", s):
        return {"last": s, "initials": "", "is_organization": True}
    return None


def normalize_people(value: Any, idx: int, field: str = "authors") -> List[Dict[str, Any]]:
    """Список авторов/редакторов: строки, объекты или одна строка."""
    if value is None or value == "":
        return []
    if isinstance(value, (str, dict)):
        value = [value]
    if not isinstance(value, list):
        raise BibliographyError(
            f"Источник #{idx}: поле '{field}' должно быть списком, получено {type(value).__name__}"
        )
    result: List[Dict[str, Any]] = []
    for pos, person in enumerate(value, 1):
        if isinstance(person, dict):
            item = dict(person)
            if "name" in item and not item.get("last"):
                parsed = _parse_person_string(str(item.get("name")))
                if parsed:
                    item.update(parsed)
            item["last"] = str(item.get("last") or "").strip()
            item["initials"] = _clean_initials(str(item.get("initials") or ""))
            if not item["last"]:
                raise BibliographyError(
                    f"Источник #{idx}, {field} #{pos}: у объекта нет поля 'last' (фамилия или организация)"
                )
            result.append(item)
        elif isinstance(person, str):
            parsed = _parse_person_string(person)
            if parsed is None:
                raise BibliographyError(
                    f"Источник #{idx}, {field} #{pos}: не удалось разобрать «{person}». "
                    "Ожидается «Фамилия И.О.», «И.О. Фамилия» или объект "
                    "{\"last\": \"Фамилия\", \"initials\": \"И.О.\"}"
                )
            result.append(parsed)
        else:
            raise BibliographyError(
                f"Источник #{idx}, {field} #{pos}: ожидается строка или объект, "
                f"получено {type(person).__name__}"
            )
    return result


def normalize_authors(authors: Any, idx: int) -> List[Dict[str, Any]]:
    """Совместимость с 6.32: нормализация поля authors."""
    return normalize_people(authors, idx, "authors")


def _person_after_slash(person: Dict[str, Any]) -> str:
    if person.get("is_organization"):
        return person["last"]
    return f"{person.get('initials', '')} {person['last']}".strip()


def _genitive_surname(surname: str) -> Optional[str]:
    """Родительный падеж русской фамилии для «под ред.»; None — не уверены."""
    s = surname
    low = s.casefold()
    rules = (
        ("ский", "ского"), ("цкий", "цкого"), ("ская", "ской"), ("цкая", "цкой"),
        ("ова", "овой"), ("ева", "евой"), ("ёва", "ёвой"), ("ина", "иной"), ("ына", "ыной"),
        ("ов", "ова"), ("ев", "ева"), ("ёв", "ёва"), ("ин", "ина"), ("ын", "ына"),
    )
    for ending, replacement in rules:
        if low.endswith(ending) and len(s) > len(ending) + 1:
            return s[: len(s) - len(ending)] + replacement
    if re.search(r"(ко|их|ых|аго|яго|ово|ич|юк|ук|ак|ян|дзе|швили|сон|ман|ер)$", low):
        return s  # несклоняемые или склоняемые без изменения окончания в И.О.-форме не угадываем
    return None


def format_editors(src: Dict[str, Any], is_ru: bool, warnings: List[str]) -> str:
    """«под ред. В.Г. Петрова» / «ed. by J. Smith». Готовый текст — editors_text."""
    ready = str(src.get("editors_text") or "").strip()
    if ready:
        return ready
    editors = src.get("editors") or []
    if not editors:
        return ""
    if not is_ru:
        return "ed. by " + ", ".join(_person_after_slash(p) for p in editors)
    names = []
    for person in editors:
        if person.get("is_organization"):
            names.append(person["last"])
            continue
        genitive = _genitive_surname(person["last"])
        if genitive is None:
            warnings.append(
                f"«{src.get('title', '')}»: не удалось поставить фамилию редактора «{person['last']}» "
                "в родительный падеж — задайте поле editors_text, например «под ред. В.Г. Петрова»"
            )
            genitive = person["last"]
        names.append(f"{person.get('initials', '')} {genitive}".strip())
    return "под ред. " + ", ".join(names)


# ========== ЯЗЫК ЗАПИСИ ==========

def _letter_counts(text: str) -> Tuple[int, int]:
    low = text.casefold()
    cyr = sum(1 for ch in low if "а" <= ch <= "я" or ch == "ё")
    lat = sum(1 for ch in low if "a" <= ch <= "z")
    return cyr, lat


CYRILLIC_SHARE_THRESHOLD = 0.3


def _is_cyrillic_source(src: Dict[str, Any]) -> bool:
    """Русская запись? Явное поле language важнее; иначе доля кириллицы.

    Считается доля кириллических букв среди всех букв заглавия, подзаголовка,
    фамилий авторов и редакторов. Запись русская, если доля не меньше 0,3:
    «ASP.NET Core MVC: руководство» и «Docker и Kubernetes: CI/CD для
    микросервисов» — русские. Если в этих полях нет букв, учитываются
    журнал/сборник, издательство и место издания.
    """
    language = str(src.get("language") or "").strip().casefold()
    if language:
        return language in {"ru", "rus", "russian", "русский", "ru-ru"}
    chunks = [str(src.get("title") or ""), str(src.get("subtitle") or "")]
    for field in ("authors", "editors"):
        for person in src.get(field) or []:
            if isinstance(person, dict):
                chunks.append(str(person.get("last") or ""))
            else:
                chunks.append(str(person))
    cyr, lat = _letter_counts(" ".join(chunks))
    if cyr + lat == 0:
        extra = " ".join(str(src.get(k) or "") for k in (
            "journal", "container_title", "book_title", "proceedings", "newspaper", "publisher", "place"))
        cyr, lat = _letter_counts(extra)
        if cyr + lat == 0:
            return True
    return cyr / float(cyr + lat) >= CYRILLIC_SHARE_THRESHOLD


def _labels(is_ru: bool) -> Dict[str, str]:
    if is_ru:
        return {"text": " [Текст]", "pages_total": "с.", "pages": "С.", "volume": "Т.",
                "number": "№", "no_place": "[Б. м.]", "no_publisher": "[б. и.]",
                "no_year": "[б. г.]", "et_al": " [и др.]"}
    return {"text": "", "pages_total": "p.", "pages": "P.", "volume": "Vol.", "number": "No.",
            "no_place": "[S. l.]", "no_publisher": "[s. n.]", "no_year": "[s. a.]", "et_al": " [et al.]"}


E_MARKER = "[Электронный ресурс]"


# ========== СБОРКА ЗАПИСИ ==========

def _join_parts(parts: List[str]) -> str:
    """Соединяет области записи через «. – », не удваивая точки."""
    cleaned = [p.strip() for p in parts if p and p.strip()]
    out = ""
    for part in cleaned:
        part = part.rstrip(".").rstrip()
        if not out:
            out = part
        elif out.endswith(("?", "!")):
            out += " – " + part
        else:
            out += ". – " + part
    if out and not out.endswith(("?", "!", ".")):
        out += "."
    return out


def _pages_range(value: Any) -> str:
    return re.sub(r"\s*[-—]\s*", "–", str(value).strip())


def _str(src: Dict[str, Any], key: str) -> str:
    value = src.get(key)
    return "" if value is None else str(value).strip()


def _people_count_for_heading(src: Dict[str, Any]) -> int:
    authors = src.get("authors") or []
    return 99 if src.get("et_al") else len(authors)


def format_authors_for_title(authors: list) -> str:
    """«Фамилия, И.О.» для заголовка записи (1–3 автора-человека), иначе ''."""
    if not authors or len(authors) >= 4:
        return ""
    first = authors[0]
    if not isinstance(first, dict):
        parsed = _parse_person_string(str(first))
        if not parsed:
            return ""
        first = parsed
    last = str(first.get("last") or "").strip()
    initials = _clean_initials(str(first.get("initials") or ""))
    if first.get("is_organization") or not last:
        return ""
    if not initials:
        return f"{last}."
    return f"{last}, {initials}"


def format_authors_after_slash(authors: list, is_ru: bool = True, et_al: bool = False) -> str:
    """«И.О. Фамилия, И.О. Фамилия»; пять и более — первые три и [и др.]."""
    people = []
    for person in authors or []:
        if isinstance(person, dict):
            people.append(person)
        else:
            parsed = _parse_person_string(str(person))
            if parsed:
                people.append(parsed)
    if not people:
        return ""
    labels = _labels(is_ru)
    if len(people) >= 5:
        return ", ".join(_person_after_slash(p) for p in people[:3]) + labels["et_al"]
    text = ", ".join(_person_after_slash(p) for p in people)
    if et_al:
        text += labels["et_al"]
    return text


def _heading(src: Dict[str, Any], warnings: List[str]) -> str:
    authors = src.get("authors") or []
    if _people_count_for_heading(src) >= 4 or not authors:
        return ""
    first = authors[0]
    if not isinstance(first, dict):
        first = _parse_person_string(str(first)) or {"last": str(first), "is_organization": True}
        authors = [first] + list(authors[1:])
    if first.get("is_organization"):
        return ""
    heading = format_authors_for_title(authors)
    if heading.endswith(".") and "," not in heading:
        warnings.append(
            f"«{src.get('title', '')}»: у автора «{first['last']}» нет инициалов — заголовок записи неполный; "
            f"если это организация, задайте {{\"last\": \"{first['last']}\", \"is_organization\": true}}"
        )
    return heading


def _full_title(src: Dict[str, Any]) -> str:
    title = _str(src, "title")
    subtitle = _str(src, "subtitle")
    return f"{title}: {subtitle}" if subtitle else title


def _responsibility(src: Dict[str, Any], is_ru: bool) -> str:
    ready = _str(src, "responsibility")
    if ready:
        return ready
    return format_authors_after_slash(src.get("authors") or [], is_ru, bool(src.get("et_al")))


def _first_area(src: Dict[str, Any], marker: str, warnings: List[str], title_note: str = "",
                with_editors: bool = True) -> str:
    """Заголовок, заглавие, маркер, сведения об ответственности."""
    is_ru = _is_cyrillic_source(src)
    heading = _heading(src, warnings)
    text = f"{heading} {_full_title(src)}".strip() if heading else _full_title(src)
    if marker:
        text += f" {marker.strip()}"
    if title_note:
        text += f": {title_note}"
    resp_parts = []
    resp = _responsibility(src, is_ru)
    if resp:
        resp_parts.append(resp)
    if with_editors:
        editors = format_editors(src, is_ru, warnings)
        if editors:
            resp_parts.append(editors)
    translation = _str(src, "translation")
    if translation:
        resp_parts.append(translation)
    if resp_parts:
        text += " / " + "; ".join(resp_parts)
    return text


def _imprint(src: Dict[str, Any], is_ru: bool, warnings: List[str], required: bool) -> str:
    labels = _labels(is_ru)
    place = _str(src, "place")
    publisher = _str(src, "publisher")
    year = _str(src, "year")
    if not (place or publisher or year):
        if not required:
            return ""
        warnings.append(f"«{src.get('title', '')}»: нет места, издательства и года издания")
        return f"{labels['no_place']}: {labels['no_publisher']}, {labels['no_year']}"
    if required and not (place and publisher and year):
        missing = [name for name, value in (("место", place), ("издательство", publisher), ("год", year)) if not value]
        warnings.append(f"«{src.get('title', '')}»: в выходных данных нет: {', '.join(missing)}")
    head = ""
    if publisher:
        head = f"{place or labels['no_place']}: {publisher}"
    elif place and not required:
        head = place  # «Уфа, 2003», «М., 1975» — издательство не указано
    elif place:
        head = f"{place}: {labels['no_publisher']}"
    if year:
        head = f"{head}, {year}" if head else year
    elif required:
        head = f"{head}, {labels['no_year']}"
    return head


def _volume_number(src: Dict[str, Any], is_ru: bool) -> str:
    labels = _labels(is_ru)
    items = []
    if _str(src, "volume"):
        items.append(f"{labels['volume']} {_str(src, 'volume')}")
    if _str(src, "number"):
        items.append(f"{labels['number']} {_str(src, 'number')}")
    return ", ".join(items)


def _url_part(src: Dict[str, Any], warnings: List[str]) -> str:
    url = _str(src, "url")
    if not url:
        return ""
    access = _str(src, "access_date")
    if not access:
        warnings.append(f"«{src.get('title', '')}»: у электронного ресурса нет даты обращения (access_date)")
        return f"URL: {url}"
    if not re.match(r"^\d{2}\.\d{2}\.\d{4}$", access):
        warnings.append(f"«{src.get('title', '')}»: дата обращения «{access}» не в формате ДД.ММ.ГГГГ")
    return f"URL: {url} (дата обращения: {access})"


def _require(src: Dict[str, Any], *fields: str) -> None:
    for field in fields:
        if not _str(src, field):
            raise BibliographyError(
                f"Источник «{src.get('title') or src.get('id') or '?'}» (type={canonical_type(src)}): "
                f"не заполнено обязательное поле '{field}'"
            )


def format_book(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Книга: Фамилия, И.О. Заглавие [Текст] / И.О. Фамилия. – Место: Изд-во, Год. – N с."""
    warnings = [] if warnings is None else warnings
    _require(src, "title")
    is_ru = _is_cyrillic_source(src)
    labels = _labels(is_ru)
    marker = E_MARKER if _str(src, "url") else labels["text"]
    parts = [_first_area(src, marker, warnings, _str(src, "title_note"))]
    if _str(src, "edition"):
        parts.append(_str(src, "edition"))
    parts.append(_imprint(src, is_ru, warnings, required=True))
    if _str(src, "pages"):
        parts.append(f"{_str(src, 'pages')} {labels['pages_total']}")
    if _str(src, "doi"):
        parts.append(f"DOI: {_str(src, 'doi')}")
    parts.append(_url_part(src, warnings))
    return _join_parts(parts)


def format_article(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Статья в журнале: … [Текст] / … // Журнал. – Год. – Т. N, № N. – С. X–Y."""
    warnings = [] if warnings is None else warnings
    _require(src, "title", "journal")
    is_ru = _is_cyrillic_source(src)
    labels = _labels(is_ru)
    marker = E_MARKER if _str(src, "url") else labels["text"]
    first = _first_area(src, marker, warnings, _str(src, "title_note")) + f" // {_str(src, 'journal')}"
    parts = [first]
    if _str(src, "place") or _str(src, "publisher"):
        parts.append(_imprint(src, is_ru, warnings, required=False))
    elif _str(src, "year"):
        parts.append(_str(src, "year"))
    vol = _volume_number(src, is_ru)
    if vol:
        parts.append(vol)
    if _str(src, "pages"):
        parts.append(f"{labels['pages']} {_pages_range(src['pages'])}")
    if _str(src, "doi"):
        parts.append(f"DOI: {_str(src, 'doi')}")
    parts.append(_url_part(src, warnings))
    return _join_parts(parts)


def format_collection_article(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Статья в сборнике или материалах конференции."""
    warnings = [] if warnings is None else warnings
    _require(src, "title")
    container = (_str(src, "container_title") or _str(src, "book_title")
                 or _str(src, "proceedings") or _str(src, "journal"))
    if not container:
        raise BibliographyError(
            f"Источник «{src.get('title')}» (type={canonical_type(src)}): нужно поле container_title "
            "(название сборника или материалов конференции)"
        )
    is_ru = _is_cyrillic_source(src)
    labels = _labels(is_ru)
    marker = E_MARKER if _str(src, "url") else labels["text"]
    first = _first_area(src, marker, warnings, _str(src, "title_note"), with_editors=False)
    first += f" // {container}"
    editors = format_editors(src, is_ru, warnings)
    if editors:
        first += f" / {editors}"
    parts = [first, _imprint(src, is_ru, warnings, required=False)]
    if _str(src, "volume"):
        parts.append(f"{labels['volume']} {_str(src, 'volume')}")
    if _str(src, "number"):
        parts.append(f"{labels['number']} {_str(src, 'number')}")
    if _str(src, "pages"):
        parts.append(f"{labels['pages']} {_pages_range(src['pages'])}")
    if _str(src, "doi"):
        parts.append(f"DOI: {_str(src, 'doi')}")
    parts.append(_url_part(src, warnings))
    return _join_parts(parts)


def format_newspaper(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Газета: … // Газета. – Год. – 15 марта (№ N). – С. 3."""
    warnings = [] if warnings is None else warnings
    _require(src, "title")
    newspaper = _str(src, "newspaper") or _str(src, "journal") or _str(src, "container_title")
    if not newspaper:
        raise BibliographyError(f"Источник «{src.get('title')}» (type=newspaper): нужно поле newspaper")
    is_ru = _is_cyrillic_source(src)
    labels = _labels(is_ru)
    marker = E_MARKER if _str(src, "url") else labels["text"]
    parts = [_first_area(src, marker, warnings) + f" // {newspaper}"]
    if _str(src, "year"):
        parts.append(_str(src, "year"))
    date = _str(src, "date")
    number = _str(src, "number")
    if date and number:
        parts.append(f"{date} ({labels['number']} {number})")
    elif date:
        parts.append(date)
    elif number:
        parts.append(f"{labels['number']} {number}")
    if _str(src, "pages"):
        parts.append(f"{labels['pages']} {_pages_range(src['pages'])}")
    parts.append(_url_part(src, warnings))
    return _join_parts(parts)


def format_electronic(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Электронный ресурс (сайт, документация, репозиторий, препринт, датасет, видео)."""
    warnings = [] if warnings is None else warnings
    _require(src, "title")
    kind = canonical_type(src)
    if not (_str(src, "url") or _str(src, "doi")):
        raise BibliographyError(
            f"Источник «{src.get('title')}» (type={kind}): электронному ресурсу нужно поле url или doi"
        )
    is_ru = _is_cyrillic_source(src)
    labels = _labels(is_ru)
    note = _str(src, "title_note") or DEFAULT_TITLE_NOTES.get(kind, "")
    first = _first_area(src, E_MARKER, warnings, note)
    container = _str(src, "journal") or _str(src, "container_title")
    if container:
        first += f" // {container}"
    parts = [first]
    if _str(src, "place") or _str(src, "publisher"):
        parts.append(_imprint(src, is_ru, warnings, required=False))
    elif _str(src, "year"):
        parts.append(_str(src, "year"))
    if _str(src, "published"):
        parts.append(f"Дата публикации: {_str(src, 'published')}")
    vol = _volume_number(src, is_ru)
    if vol:
        parts.append(vol)
    if _str(src, "pages"):
        parts.append(f"{labels['pages']} {_pages_range(src['pages'])}")
    if _str(src, "doi"):
        parts.append(f"DOI: {_str(src, 'doi')}")
    parts.append(_url_part(src, warnings))
    return _join_parts(parts)


def format_dissertation(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Диссертация / автореферат: …: Дис. … канд. пед. наук [Текст] / … – Место: Орг., Год. – N с."""
    warnings = [] if warnings is None else warnings
    _require(src, "title")
    if not src.get("authors"):
        raise BibliographyError(f"Источник «{src.get('title')}» (type={canonical_type(src)}): нужен автор (authors)")
    is_autoref = canonical_type(src) == "autoreferat" or bool(src.get("is_autoref"))
    thesis_type = _str(src, "thesis_type")
    if not thesis_type:
        degree = _str(src, "degree") or "канд."
        field = _str(src, "field") or "пед."
        base = degree if "наук" in degree.casefold() else f"{degree} {field} наук"
        thesis_type = f"автореф. дис. … {base}" if is_autoref else f"Дис. … {base}"
    heading = _heading(src, warnings)
    first = f"{heading} {_full_title(src)}".strip() + f": {thesis_type} [Текст]"
    resp = _responsibility(src, True)
    if resp:
        first += f" / {resp}"
    place = _str(src, "place")
    organization = _str(src, "organization") or _str(src, "publisher")
    year = _str(src, "year")
    if not place:
        warnings.append(f"«{src.get('title')}»: не указано место защиты/издания")
        place = "[Б. м.]"
    imprint = f"{place}: {organization}" if organization else place
    if year:
        imprint += f", {year}"
    parts = [first, imprint]
    if _str(src, "pages"):
        parts.append(f"{_str(src, 'pages')} с.")
    return _join_parts(parts)


def format_normative(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Нормативный акт: Заглавие [Текст]: вид акта от ДД.ММ.ГГГГ № N. – Источник."""
    warnings = [] if warnings is None else warnings
    _require(src, "title")
    marker = E_MARKER if _str(src, "url") else "[Текст]"
    first = f"{_full_title(src)} {marker}"
    details = []
    if _str(src, "doc_type"):
        details.append(_str(src, "doc_type"))
    if _str(src, "date"):
        details.append(f"от {_str(src, 'date')}")
    if _str(src, "number"):
        details.append(f"№ {_str(src, 'number')}")
    if details:
        first += ": " + " ".join(details)
    if _str(src, "adoption_note"):
        first += f": {_str(src, 'adoption_note')}"
    parts = [first]
    if _str(src, "source"):
        parts.append(_str(src, "source"))
    parts.append(_url_part(src, warnings))
    return _join_parts(parts)


def format_standard(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Стандарт: ГОСТ Р 7.0.5-2008. Заглавие [Текст]. – Введ. 2009-01-01. – М.: Изд-во, 2008. – 19 с."""
    warnings = [] if warnings is None else warnings
    _require(src, "title")
    designation = _str(src, "designation")
    if not designation:
        warnings.append(f"«{src.get('title')}»: у стандарта нет обозначения (designation), например «ГОСТ Р 7.0.5-2008»")
    marker = E_MARKER if _str(src, "url") else "[Текст]"
    first = f"{designation}. {_full_title(src)}" if designation else _full_title(src)
    first = f"{first} {marker}"
    parts = [first]
    if _str(src, "introduced"):
        parts.append(f"Введ. {_str(src, 'introduced')}")
    imprint = _imprint(src, True, warnings, required=False)
    if imprint:
        parts.append(imprint)
    if _str(src, "pages"):
        parts.append(f"{_str(src, 'pages')} с.")
    parts.append(_url_part(src, warnings))
    return _join_parts(parts)


def format_patent(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Патент: Пат. N Страна, МПК …. Заглавие [Текст] / Авторы; заявитель. – № заявки; заявл. …; опубл. …."""
    warnings = [] if warnings is None else warnings
    _require(src, "title", "number")
    head = f"Пат. {_str(src, 'number')}"
    if _str(src, "country"):
        head += f" {_str(src, 'country')}"
    if _str(src, "ipc"):
        head += f", МПК {_str(src, 'ipc')}"
    first = f"{head}. {_full_title(src)} [Текст]"
    resp = []
    authors = format_authors_after_slash(src.get("authors") or [], True)
    if authors:
        resp.append(authors)
    if _str(src, "holder"):
        resp.append(_str(src, "holder"))
    if resp:
        first += " / " + "; ".join(resp)
    tail = []
    if _str(src, "application_number"):
        tail.append(f"№ {_str(src, 'application_number')}")
    if _str(src, "filed"):
        tail.append(f"заявл. {_str(src, 'filed')}")
    if _str(src, "published"):
        tail.append(f"опубл. {_str(src, 'published')}")
    parts = [first]
    if tail:
        parts.append("; ".join(tail))
    return _join_parts(parts)


FORMATTERS: Dict[str, Callable[..., str]] = {
    "book": format_book,
    "article": format_article,
    "collection_article": format_collection_article,
    "conference": format_collection_article,
    "newspaper": format_newspaper,
    "electronic": format_electronic,
    "software_doc": format_electronic,
    "repository": format_electronic,
    "preprint": format_electronic,
    "dataset": format_electronic,
    "video": format_electronic,
    "dissertation": format_dissertation,
    "autoreferat": format_dissertation,
    "normative": format_normative,
    "standard": format_standard,
    "patent": format_patent,
}


def format_source(src: dict, warnings: Optional[List[str]] = None) -> str:
    """Оформляет одну запись; неизвестный тип — понятная ошибка, а не молчаливый book."""
    if not isinstance(src, dict):
        raise BibliographyError(f"Источник должен быть объектом, получено {type(src).__name__}")
    kind = canonical_type(src)
    formatter = FORMATTERS.get(kind)
    if formatter is None:
        raise BibliographyError(
            f"Источник «{src.get('title') or src.get('id') or '?'}»: неизвестный тип «{src.get('type')}». "
            f"Допустимые типы: {', '.join(CANONICAL_TYPES)}"
        )
    if not is_structured(src):
        raise BibliographyError(
            f"Источник id={src.get('id', '?')}: запись не структурирована (нет title, есть только raw_text) — "
            "разберите её на поля по references/gost-citations.md"
        )
    for field in ("authors", "editors"):
        people = src.get(field)
        if people and (not isinstance(people, list) or not all(isinstance(p, dict) for p in people)):
            src = normalize_source(src, 0)
            break
    return formatter(src, warnings)


# ========== СОРТИРОВКА ==========

def _char_weight(ch: str) -> int:
    if ch.isspace():
        return 0
    if "0" <= ch <= "9":
        return 1 + ord(ch) - ord("0")
    if ch == "ё":
        ch = "е"
    if "а" <= ch <= "я":
        return 20 + ord(ch) - ord("а")
    if "a" <= ch <= "z":
        return 60 + ord(ch) - ord("a")
    if ch.isalpha():
        import unicodedata

        base = unicodedata.normalize("NFKD", ch)[0]
        if "a" <= base <= "z":
            return 60 + ord(base) - ord("a")
        return 1000 + ord(ch)
    return -1  # пунктуация — разделитель


def sort_text_key(text: str) -> Tuple[int, ...]:
    """Ключ алфавитной сортировки: пробел < цифры < кириллица (ё = е) < латиница."""
    weights: List[int] = []
    for ch in str(text or "").casefold():
        weight = _char_weight(ch)
        if weight <= 0:
            if weights and weights[-1] != 0:
                weights.append(0)
            continue
        weights.append(weight)
    while weights and weights[-1] == 0:
        weights.pop()
    return tuple(weights)


def sort_head(src: Dict[str, Any]) -> str:
    """Первый элемент записи: заголовок по автору или заглавие."""
    kind = canonical_type(src)
    if kind == "standard":
        return f"{_str(src, 'designation')} {_full_title(src)}".strip()
    if kind == "patent":
        return f"Пат. {_str(src, 'number')}"
    if kind not in NORMATIVE_GROUP:
        heading = _heading(src, [])
        if heading:
            return heading
    return _full_title(src)


def get_sort_key(src: dict, group_normative_first: bool = False) -> tuple:
    """Ключ сортировки по разделу 3 SPEC (совместимость: один аргумент)."""
    group = 0 if (group_normative_first and canonical_type(src) in NORMATIVE_GROUP) else 1
    year_text = _str(src, "year")
    year = int(year_text) if year_text.isdigit() else 0
    return (group, sort_text_key(sort_head(src)), sort_text_key(_full_title(src)), year)


# ========== НОРМАЛИЗАЦИЯ ВХОДНЫХ ДАННЫХ ==========

def normalize_source(src: Any, idx: int) -> Dict[str, Any]:
    """Нормализует одну запись: алиасы полей, authors/editors, тип."""
    if not isinstance(src, dict):
        raise BibliographyError(
            f"Источник #{idx}: ожидается объект (dict), получено {type(src).__name__}: {src!r}"
        )
    normalized: Dict[str, Any] = {}
    for key, value in src.items():
        canonical = FIELD_ALIASES.get(key, key)
        if canonical not in normalized or key == canonical:
            normalized[canonical] = value
    normalized["authors"] = normalize_people(normalized.get("authors"), idx, "authors")
    if "editors" in normalized:
        normalized["editors"] = normalize_people(normalized.get("editors"), idx, "editors")
    if "type" in normalized:
        normalized["type"] = canonical_type(normalized)
    if "id" in normalized:
        value = normalized["id"]
        if isinstance(value, bool) or not isinstance(value, (int, str)):
            raise BibliographyError(f"Источник #{idx}: id должен быть строкой или целым числом")
        if isinstance(value, str) and not ID_RE.match(value):
            raise BibliographyError(
                f"Источник #{idx}: id «{value}» допускает только латинские буквы, цифры и символы _ . : -"
            )
    return normalized


def normalize_input_data(data: Any) -> List[Dict[str, Any]]:
    """Массив источников или объект с ключом sources/items/bibliography/refs/references."""
    if isinstance(data, list):
        sources = data
    elif isinstance(data, dict):
        for key in ("sources", "items", "bibliography", "refs", "references"):
            if key in data and isinstance(data[key], list):
                sources = data[key]
                break
        else:
            raise BibliographyError(
                "Неподдерживаемый формат JSON: ожидается массив источников или объект с ключом "
                f"'sources'. Получен объект с ключами {sorted(data.keys())}"
            )
    else:
        raise BibliographyError(
            f"Неподдерживаемый тип входных данных: {type(data).__name__}; ожидается массив источников"
        )
    return [normalize_source(src, i) for i, src in enumerate(sources, 1)]


def source_ids(sources: List[Dict[str, Any]]) -> List[Any]:
    """id записей; у записи без id неявный id — её номер в реестре. Дубли — ошибка."""
    ids: List[Any] = []
    seen: Dict[str, int] = {}
    for pos, src in enumerate(sources, 1):
        value = src.get("id", pos)
        key = str(value)
        if key in seen:
            raise BibliographyError(
                f"Повторяющийся id «{value}» у записей #{seen[key]} и #{pos} реестра источников"
            )
        seen[key] = pos
        ids.append(value)
    return ids


def number_sources(
    sources: List[Dict[str, Any]],
    *,
    keep_order: bool = False,
    group_normative_first: bool = False,
) -> Dict[str, Any]:
    """Оформляет и нумерует список.

    Возвращает {"entries": [{"number", "id", "text", "source"}], "mapping": {str(id): number},
    "warnings": [...]}; ошибки записей собираются и выбрасываются одним BibliographyError.
    """
    ids = source_ids(sources)
    warnings: List[str] = []
    errors: List[str] = []
    formatted: List[Tuple[Any, Dict[str, Any], str]] = []
    for value, src in zip(ids, sources):
        try:
            formatted.append((value, src, format_source(src, warnings)))
        except BibliographyError as error:
            errors.append(str(error))
    if errors:
        raise BibliographyError("\n".join(errors))
    if keep_order:
        ordered = formatted
    else:
        ordered = [item for _, item in sorted(
            enumerate(formatted), key=lambda pair: (get_sort_key(pair[1][1], group_normative_first), pair[0])
        )]
    entries = []
    mapping: Dict[str, int] = {}
    for number, (value, src, text) in enumerate(ordered, 1):
        entries.append({"number": number, "id": value, "text": text, "source": src})
        mapping[str(value)] = number
    return {"entries": entries, "mapping": mapping, "warnings": warnings}


def alphabetical_problems(entries: List[Dict[str, Any]], group_normative_first: bool = False) -> List[str]:
    """Пары соседних записей, нарушающих алфавитный порядок."""
    problems = []
    for prev, cur in zip(entries, entries[1:]):
        if get_sort_key(prev["source"], group_normative_first) > get_sort_key(cur["source"], group_normative_first):
            problems.append(f"«{prev['text'][:40]}…» стоит перед «{cur['text'][:40]}…»")
    return problems


def check_preserved_ids(sources: List[Dict[str, Any]]) -> List[int]:
    """Для режима без сортировки: у всех записей целые положительные уникальные id."""
    ids = source_ids(sources)
    bad = [value for value in ids
           if isinstance(value, bool) or not str(value).isdigit() or int(str(value)) < 1]
    if bad:
        raise BibliographyError(
            "без сортировки номер записи = её id, поэтому нужны целые положительные id; неподходящие: "
            + ", ".join(str(v) for v in bad)
        )
    return [int(str(value)) for value in ids]


def format_bibliography(sources: list, preserve_ids: bool = False, group_normative_first: bool = False) -> tuple:
    """Совместимый интерфейс: (текст списка, {старый id: новый номер}).

    preserve_ids=False — алфавитная сортировка и сквозная нумерация; словарь
    содержит записи, чей номер отличается от прежнего id.
    preserve_ids=True — порядок входа, номер = id (целые уникальные id обязательны).
    """
    normalized = [normalize_source(src, i) for i, src in enumerate(sources, 1)]
    if preserve_ids:
        check_preserved_ids(normalized)
        result = number_sources(normalized, keep_order=True)
        lines = [f"{int(entry['id'])}. {entry['text']}" for entry in result["entries"]]
        return "\n\n".join(lines), {}
    result = number_sources(normalized, group_normative_first=group_normative_first)
    mapping = {}
    for entry in result["entries"]:
        if str(entry["id"]) != str(entry["number"]):
            mapping[entry["id"]] = entry["number"]
    lines = [f"{entry['number']}. {entry['text']}" for entry in result["entries"]]
    return "\n\n".join(lines), mapping


# ========== ССЫЛКИ В ТЕКСТЕ ==========

_BRACKET_RE = re.compile(r"\[([^\[\]\n]{1,400})\]")
_INLINE_CODE_RE = re.compile(r"`[^`\n]+`")
_PREFIX_RE = r"(?P<prefix>(?:Цит\.\s*по(?:\s*(?:кн|ст)\.)?|[СсCc]м\.|[Сс]р\.)\s*:?\s*)?"
_ID_TOKEN = r"@(?P<id>[A-Za-z0-9_](?:[A-Za-z0-9_.:\-]*[A-Za-z0-9_])?)"
_NUM_TOKEN = r"(?P<num>[1-9]\d{0,2})(?:\s*[–—-]\s*(?P<num2>[1-9]\d{0,2}))?"
_PART_RE = re.compile(r"^\s*" + _PREFIX_RE + r"(?:" + _ID_TOKEN + r"|" + _NUM_TOKEN + r")(?P<loc>\s*,\s*[^;]*?)?\s*$")
_LOCATOR_RE = re.compile(
    r"^\s*,\s*(?:[сСcCpP]{1,2}\.|стр\.|[тТ]\.|[гГ]л\.|разд\.|[рР]ис\.|[тТ]абл\.|[чЧ]\.|[сС]т\.|§|[vV]ol\.|[nN]o\.|pp?\.|[сС]\s*\d)"
)


def _mask_inline_code(text: str) -> str:
    return _INLINE_CODE_RE.sub(lambda m: " " * len(m.group(0)), text)


def parse_citation_group(inner: str) -> Optional[List[Dict[str, Any]]]:
    """Разбирает содержимое скобок; None — это не ссылка на источник."""
    parts: List[Dict[str, Any]] = []
    for raw_part in inner.split(";"):
        match = _PART_RE.match(raw_part)
        if not match:
            return None
        loc = match.group("loc") or ""
        if match.group("id"):
            parts.append({"prefix": match.group("prefix") or "", "kind": "id", "ref": match.group("id"), "loc": loc})
            continue
        nums = [int(match.group("num"))]
        if match.group("num2"):
            first, last = nums[0], int(match.group("num2"))
            if last < first or last - first > 50:
                return None
            parts.append({"prefix": match.group("prefix") or "", "kind": "range",
                          "ref": list(range(first, last + 1)), "loc": loc})
            continue
        if loc and re.match(r"^\s*,\s*\d[\d\s,]*$", loc):
            extra = [int(x) for x in re.findall(r"\d+", loc)]
            parts.append({"prefix": match.group("prefix") or "", "kind": "num", "ref": nums[0], "loc": ""})
            for value in extra:
                parts.append({"prefix": "", "kind": "num", "ref": value, "loc": "", "comma": True})
            continue
        if loc and not _LOCATOR_RE.match(loc):
            return None
        parts.append({"prefix": match.group("prefix") or "", "kind": "num", "ref": nums[0], "loc": loc})
    return parts or None


def find_citations(text: str) -> List[Dict[str, Any]]:
    """Все ссылки вида [@id…], [N…], [Цит. по: …] вне inline-кода."""
    masked = _mask_inline_code(text or "")
    found = []
    for match in _BRACKET_RE.finditer(masked):
        parts = parse_citation_group(match.group(1))
        if parts is None:
            continue
        found.append({"start": match.start(), "end": match.end(), "text": text[match.start():match.end()],
                      "parts": parts})
    return found


def citation_refs(text: str) -> List[Tuple[str, Any]]:
    """Список (kind, ref) всех ссылок: kind = 'id' (строка) или 'num' (целое)."""
    refs: List[Tuple[str, Any]] = []
    for citation in find_citations(text):
        for part in citation["parts"]:
            if part["kind"] == "range":
                refs.extend(("num", value) for value in part["ref"])
            else:
                refs.append((part["kind"], part["ref"]))
    return refs


def _render_numbers(numbers: List[int]) -> str:
    numbers = sorted(set(numbers))
    if len(numbers) > 2 and numbers == list(range(numbers[0], numbers[-1] + 1)):
        return f"{numbers[0]}–{numbers[-1]}"
    return "; ".join(str(n) for n in numbers)


def rewrite_citations(text: str, resolve: Callable[[str, Any], Optional[int]],
                      problems: Optional[List[str]] = None) -> str:
    """Переписывает ссылки в числовой вид [N, с. …] по функции resolve(kind, ref).

    resolve возвращает номер или None (ссылка не найдена — остаётся как есть,
    сообщение добавляется в problems).
    """
    if not text:
        return text
    problems = [] if problems is None else problems
    citations = find_citations(text)
    if not citations:
        return text
    out = []
    pos = 0
    for citation in citations:
        rendered_parts = []
        ok = True
        for part in citation["parts"]:
            if part["kind"] == "range":
                numbers = []
                for value in part["ref"]:
                    number = resolve("num", value)
                    if number is None:
                        ok = False
                        problems.append(f"ссылка {citation['text']}: источник {value} не найден")
                        break
                    numbers.append(number)
                if not ok:
                    break
                rendered_parts.append(f"{part['prefix']}{_render_numbers(numbers)}{part['loc']}")
                continue
            number = resolve(part["kind"], part["ref"])
            if number is None:
                ok = False
                shown = f"@{part['ref']}" if part["kind"] == "id" else str(part["ref"])
                problems.append(f"ссылка {citation['text']}: источник {shown} не найден")
                break
            rendered_parts.append(f"{part['prefix']}{number}{part['loc']}")
        out.append(text[pos:citation["start"]])
        if ok:
            joined = ""
            for part, rendered in zip(citation["parts"], rendered_parts):
                if not joined:
                    joined = rendered
                elif part.get("comma"):
                    joined += ", " + rendered
                else:
                    joined += "; " + rendered
            out.append(f"[{joined}]")
        else:
            out.append(citation["text"])
        pos = citation["end"]
    out.append(text[pos:])
    return "".join(out)


def make_resolver(mapping: Dict[str, int], ids: List[Any]) -> Callable[[str, Any], Optional[int]]:
    """resolve(kind, ref): @id → номер; legacy N → номер записи с целым id N."""
    by_int: Dict[int, int] = {}
    for value in ids:
        text = str(value)
        if text.isdigit():
            by_int[int(text)] = mapping[text]

    def resolve(kind: str, ref: Any) -> Optional[int]:
        if kind == "id":
            return mapping.get(str(ref))
        return by_int.get(int(ref))

    return resolve


# ========== I/O ==========

def load_sources(path: Path) -> Any:
    """Загрузить JSON (utf-8-sig) или YAML."""
    text = path.read_text(encoding="utf-8-sig")
    if path.suffix.lower() in (".yaml", ".yml"):
        try:
            import yaml  # type: ignore
        except ImportError as error:
            raise ImportError("для YAML нужен пакет pyyaml (pip install pyyaml)") from error
        return yaml.safe_load(text)
    return json.loads(text)


def save_as_docx(text: str, output_path: Path) -> None:
    """Сохранить список в .docx: A4, поля МПГУ, TNR 14, интервал 1,5, отступ 1,25 см."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Cm, Mm, Pt, RGBColor

    doc = Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(35)
        section.right_margin = Mm(10)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("Список использованной литературы")
    run.bold = True
    for block in text.split("\n\n"):
        if not block.strip():
            continue
        paragraph = doc.add_paragraph(block)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def _configure_streams() -> None:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8", errors="replace")
            except (ValueError, OSError):
                pass


def main(argv: Optional[List[str]] = None) -> int:
    _configure_streams()
    parser = argparse.ArgumentParser(
        description="Список литературы по образцам методички МПГУ: алфавитный порядок, сквозная нумерация"
    )
    parser.add_argument("sources_file", help="sources.json (или YAML) — реестр источников")
    parser.add_argument("-o", "--output", help="Файл результата: .txt или .docx")
    parser.add_argument("--json", action="store_true", help="Вывести результат в stdout как JSON")
    parser.add_argument("--preserve-ids", action="store_true",
                        help="Не сортировать: порядок реестра, номер = целый id записи (предупреждает о нарушении алфавита)")
    parser.add_argument("--group-normative-first", action="store_true",
                        help="Нормативные акты и стандарты — в начало списка (только если так требует программа)")
    parser.add_argument("--mapping", help="JSON-файл с соответствием {id: номер в списке}")
    args = parser.parse_args(argv)

    def fail(code: int, message: str) -> int:
        if args.json:
            print(json.dumps({"status": "error", "exit_code": code, "errors": message.split("\n")},
                             ensure_ascii=False, indent=2))
        else:
            print(f"ОШИБКА: {message}", file=sys.stderr)
        return code

    src_path = Path(args.sources_file)
    if not src_path.is_file():
        return fail(2, f"файл не найден: {src_path}")
    try:
        raw = load_sources(src_path)
    except ImportError as error:
        return fail(2, str(error))
    except (OSError, UnicodeDecodeError) as error:
        return fail(2, f"не удалось прочитать {src_path}: {error}")
    except json.JSONDecodeError as error:
        return fail(2, f"{src_path} — некорректный JSON: {error}")
    try:
        sources = normalize_input_data(raw)
        if not sources:
            raise BibliographyError("список источников пуст")
        skipped = [src for src in sources if normalize_status(src.get("status")) == "rejected"]
        sources = [src for src in sources if normalize_status(src.get("status")) != "rejected"]
        if args.preserve_ids:
            check_preserved_ids(sources)
            result = number_sources(sources, keep_order=True)
            text = "\n\n".join(f"{int(str(e['id']))}. {e['text']}" for e in result["entries"])
            numbers = {str(entry["id"]): int(str(entry["id"])) for entry in result["entries"]}
        else:
            result = number_sources(sources, group_normative_first=args.group_normative_first)
            text = "\n\n".join(f"{e['number']}. {e['text']}" for e in result["entries"])
            numbers = result["mapping"]
    except BibliographyError as error:
        return fail(1, str(error))
    except (KeyError, TypeError, AttributeError) as error:
        return fail(1, f"некорректная запись реестра: {error!r}")

    warnings = list(result["warnings"])
    for src in skipped:
        warnings.append(f"источник id={src.get('id', '?')} со статусом rejected в список не включён")
    if args.preserve_ids:
        ids = [int(str(entry["id"])) for entry in result["entries"]]
        if ids != list(range(1, len(ids) + 1)):
            warnings.append("--preserve-ids: номера не идут подряд с 1 — нумерация списка не сквозная")
        problems = alphabetical_problems(result["entries"], args.group_normative_first)
        if problems:
            warnings.append(
                "--preserve-ids: порядок не алфавитный (методичка: «Список составляется в алфавитном порядке»): "
                + problems[0]
            )
    else:
        changed = {k: v for k, v in numbers.items() if k != str(v)}
        if changed:
            warnings.append(
                "номера источников после сортировки отличаются от id: "
                + ", ".join(f"[{k}] → [{v}]" for k, v in list(changed.items())[:20])
                + ". Ссылки [@id] в drafts/*.md пересчитывает build_vkr.py; числовые ссылки в готовом "
                  "тексте нужно исправить по --mapping"
            )

    try:
        if args.mapping:
            mapping_path = Path(args.mapping)
            mapping_path.parent.mkdir(parents=True, exist_ok=True)
            mapping_path.write_text(json.dumps(numbers, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        if args.output:
            out_path = Path(args.output)
            if out_path.suffix.lower() == ".docx":
                try:
                    save_as_docx(text, out_path)
                except ImportError:
                    return fail(2, "для вывода .docx нужен python-docx (pip install python-docx)")
            else:
                out_path.parent.mkdir(parents=True, exist_ok=True)
                out_path.write_text(text + "\n", encoding="utf-8")
    except OSError as error:
        return fail(2, f"не удалось записать результат: {error}")

    if args.json:
        payload = {
            "status": "ok",
            "tool_version": TOOL_VERSION,
            "entries": [{"number": (int(e["id"]) if args.preserve_ids else e["number"]), "id": e["id"],
                         "text": e["text"]} for e in result["entries"]],
            "mapping": numbers,
            "warnings": warnings,
            "errors": [],
        }
        print(json.dumps(payload, ensure_ascii=False, indent=2))
    else:
        if not args.output:
            print(text)
        else:
            print(f"OK: сохранено: {args.output}")
        for warning in warnings:
            print(f"ПРЕДУПРЕЖДЕНИЕ: {warning}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
