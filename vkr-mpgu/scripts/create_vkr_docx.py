#!/usr/bin/env python3
"""Генератор DOCX ВКР по методичке МПГУ (vkr-mpgu 6.33) — рендерер JSON.

Обычно его вызывает ``build_vkr.py``: черновики ``drafts/*.md`` + реестр
``sources.json`` → ``exports/vkr-content.json`` → ``final/vkr.docx``.
Прямой вызов:

    python create_vkr_docx.py vkr-content.json -o vkr.docx [--base-dir PROJECT]
           [--keep-source-order] [--group-normative-first] [--json]

Полная схема входного JSON — ``references/docx-input-schema.md``; допустимые
ключи перечислены в константах ``*_KEYS`` ниже, неизвестный ключ — ошибка.

Оформление: A4, поля 20/20/35/10 мм, Times New Roman 14, чёрный, интервал 1,5,
абзацный отступ 1,25 см, выравнивание по ширине; номер страницы внизу по
центру, на титуле не виден. Стили: Heading 1/2 (TNR, чёрные, без тематических
шрифтов), VKR Title, VKR Listing, VKR Caption, VKR Bibliography,
VKR Appendix Label, VKR Appendix Title, toc 1–3, List Paragraph + numPr.

Структурированные ``sources`` сортируются по разделу 3 SPEC (алфавит), ссылки
``[@id]`` и числовые ``[N]`` во всех текстовых блоках, кроме кода листингов,
переписываются в новые номера. ``keep_source_order`` отключает сортировку:
номер записи = её целый id.

Коды выхода: 0 — создан; 1 — содержание нельзя оформить (ссылка на
отсутствующий источник, запись реестра без обязательных полей, нет файла
рисунка); 2 — ошибка использования, схемы JSON, ввода-вывода или нет python-docx.
"""

from __future__ import annotations

import argparse
import copy
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Tuple

GENERATOR_VERSION = "6.33"

# ---------------------------------------------------------------------------
# Схема входа (docx-input-schema.md перечисляет те же ключи; тест сверяет)
# ---------------------------------------------------------------------------

TOP_LEVEL_KEYS = (
    "title", "title_page", "annotation", "introduction", "chapters", "conclusion",
    "sources", "bibliography", "appendices", "skip_last_page", "keep_source_order",
    "group_normative_first",
)
TITLE_PAGE_KEYS = (
    "university", "institute", "department", "program_code", "program_name",
    "program_profile", "work_type", "author", "authors", "group", "supervisor",
    "supervisor_title", "head_of_department", "head_title", "city", "year",
)
# Ключи 6.32, которые читаются с предупреждением и переводятся в ключи SPEC 2.
LEGACY_TITLE_PAGE_KEYS = {
    "direction_code": "program_code",
    "direction_name": "program_name",
    "profile": "program_profile",
    "program": "program_code + program_name",
    "supervisor_name": "supervisor",
    "supervisor_position": "supervisor_title",
    "supervisor_degree": "supervisor_title",
    "head_name": "head_of_department",
    "head_degree": "head_title",
}
CHAPTER_KEYS = ("title", "blocks", "paragraphs", "conclusion")
PARAGRAPH_KEYS = ("title", "text", "blocks")
APPENDIX_KEYS = ("title", "label", "content", "blocks")
BLOCK_KEYS = {
    "text": ("type", "content"),
    "list": ("type", "items", "ordered"),
    "table": ("type", "number", "title", "headers", "rows", "source"),
    "listing": ("type", "number", "caption", "code", "language"),
    "figure": ("type", "number", "caption", "path", "placeholder", "width_cm"),
}
SECTION_KEYS = ("annotation", "introduction", "conclusion")

STYLE_TITLE = "VKR Title"
STYLE_LISTING = "VKR Listing"
STYLE_CAPTION = "VKR Caption"
STYLE_BIBLIOGRAPHY = "VKR Bibliography"
STYLE_APPENDIX_LABEL = "VKR Appendix Label"
STYLE_APPENDIX_TITLE = "VKR Appendix Title"
STYLE_LIST = "List Paragraph"

TOC_PLACEHOLDER = "[Для обновления содержания: щёлкните правой кнопкой → Обновить поле]"
FIGURE_PLACEHOLDER = "[ Здесь вставить рисунок вручную в Word ]"
BIBLIOGRAPHY_HEADING = "Список использованной литературы"
DEFAULT_UNIVERSITY = "Московский педагогический государственный университет"
TEXT_WIDTH_MM = 165.0  # 210 − 35 − 10
MAX_FIGURE_HEIGHT_MM = 200.0

CHAPTER_PREFIX_RE = re.compile(
    r"^\s*глава\s+(?:[IVXLCDM]+|\d+)\s*(?:[.:)]|[-–—])?\s*", re.IGNORECASE
)
PARAGRAPH_NUMBER_RE = re.compile(r"^\s*\d+(?:\.\d+)+\.?\s+|^\s*\d+\.\s+")
APPENDIX_PREFIX_RE = re.compile(
    r"^\s*приложение\s+(?:[А-ЯЁA-Z]|\d+)\s*(?:[.:)]|[-–—])?\s*", re.IGNORECASE
)

ROMAN = {1: "I", 2: "II", 3: "III", 4: "IV", 5: "V", 6: "VI", 7: "VII", 8: "VIII", 9: "IX", 10: "X"}


class InputError(ValueError):
    """Ошибка схемы входного JSON (код 2)."""


class ContentError(ValueError):
    """Содержание нельзя оформить корректно (код 1)."""


def to_roman(n: int) -> str:
    return ROMAN.get(n, str(n))


def strip_chapter_prefix(title: str) -> str:
    """«Глава I. Теоретические основы» → «Теоретические основы» (R2-16)."""
    return CHAPTER_PREFIX_RE.sub("", str(title or ""), count=1).strip().rstrip(".").strip()


def strip_paragraph_number(title: str) -> str:
    """«1.2. Название» → «Название»; номер ставит генератор."""
    return PARAGRAPH_NUMBER_RE.sub("", str(title or ""), count=1).strip().rstrip(".").strip()


def strip_appendix_prefix(title: str) -> str:
    return APPENDIX_PREFIX_RE.sub("", str(title or ""), count=1).strip().rstrip(".").strip()


def _scripts_dir() -> Path:
    return Path(__file__).resolve().parent


def _import_bibliography():
    sys.dont_write_bytecode = True
    scripts = str(_scripts_dir())
    if scripts not in sys.path:
        sys.path.insert(0, scripts)
    import format_bibliography  # noqa: WPS433

    return format_bibliography


# ---------------------------------------------------------------------------
# Нормализация и проверка входа
# ---------------------------------------------------------------------------

def _check_keys(obj: Dict[str, Any], allowed: Tuple[str, ...], where: str, errors: List[str]) -> None:
    unknown = sorted(k for k in obj if not str(k).startswith("_") and k not in allowed)
    if unknown:
        errors.append(
            f"{where}: неизвестные ключи {', '.join(unknown)}; допустимые: {', '.join(allowed)}"
        )


def _normalize_title_page(raw: Any, errors: List[str], warnings: List[str]) -> Dict[str, Any]:
    if raw is None:
        raw = {}
    if not isinstance(raw, dict):
        errors.append("title_page: ожидается объект")
        return {}
    tp: Dict[str, Any] = {}
    legacy_used = []
    for key, value in raw.items():
        if str(key).startswith("_"):
            continue
        if key in TITLE_PAGE_KEYS:
            tp[key] = value
        elif key in LEGACY_TITLE_PAGE_KEYS:
            legacy_used.append(key)
        else:
            errors.append(
                f"title_page: неизвестный ключ «{key}»; допустимые: {', '.join(TITLE_PAGE_KEYS)}"
            )
    if legacy_used:
        warnings.append(
            "title_page: устаревшие ключи 6.32 ("
            + ", ".join(f"{k} → {LEGACY_TITLE_PAGE_KEYS[k]}" for k in legacy_used)
            + ") переведены в ключи 6.33; обновите JSON"
        )

        def take(key: str) -> str:
            return str(raw.get(key) or "").strip()

        if take("direction_code") and not tp.get("program_code"):
            tp["program_code"] = take("direction_code")
        if take("direction_name") and not tp.get("program_name"):
            tp["program_name"] = take("direction_name")
        if take("program") and not (tp.get("program_code") or tp.get("program_name")):
            match = re.match(r"^(\d{2}\.\d{2}\.\d{2})\.?\s+(.+)$", take("program"))
            if match:
                tp["program_code"], tp["program_name"] = match.group(1), match.group(2)
            else:
                tp["program_name"] = take("program")
        if take("profile") and not tp.get("program_profile"):
            tp["program_profile"] = take("profile")
        if take("supervisor_name") and not tp.get("supervisor"):
            tp["supervisor"] = take("supervisor_name")
        position = ", ".join(x for x in (take("supervisor_position"), take("supervisor_degree")) if x)
        if position and not tp.get("supervisor_title"):
            tp["supervisor_title"] = position
        if take("head_name") and not tp.get("head_of_department"):
            tp["head_of_department"] = take("head_name")
        if take("head_degree") and not tp.get("head_title"):
            tp["head_title"] = take("head_degree")
    authors = tp.get("authors")
    if authors is not None:
        if not isinstance(authors, list) or not all(isinstance(a, str) and a.strip() for a in authors):
            errors.append("title_page.authors: ожидается массив непустых строк (ФИО участников)")
    author = tp.get("author")
    if author is not None and not isinstance(author, str):
        errors.append("title_page.author: ожидается строка (ФИО полностью)")
    if not (str(author or "").strip() or (isinstance(authors, list) and authors)):
        errors.append("title_page: обязательно поле author (ФИО полностью) или authors (массив ФИО)")
    for key in TITLE_PAGE_KEYS:
        if key in ("authors",):
            continue
        value = tp.get(key)
        if value is not None and not isinstance(value, (str, int)):
            errors.append(f"title_page.{key}: ожидается строка")
    return tp


def _normalize_blocks(value: Any, where: str, errors: List[str]) -> List[Dict[str, Any]]:
    """Строка → один text-блок; массив → проверенные блоки."""
    if value is None or value == "":
        return []
    if isinstance(value, str):
        return [{"type": "text", "content": value}]
    if not isinstance(value, list):
        errors.append(f"{where}: ожидается строка или массив блоков")
        return []
    blocks = []
    for i, block in enumerate(value, 1):
        label = f"{where}[{i}]"
        if not isinstance(block, dict):
            errors.append(f"{label}: блок должен быть объектом")
            continue
        btype = block.get("type", "text")
        if btype not in BLOCK_KEYS:
            errors.append(f"{label}: неизвестный тип блока «{btype}»; допустимые: {', '.join(BLOCK_KEYS)}")
            continue
        _check_keys(block, BLOCK_KEYS[btype], label, errors)
        item = {k: v for k, v in block.items() if not str(k).startswith("_")}
        item["type"] = btype
        if btype == "text" and not isinstance(item.get("content", ""), str):
            errors.append(f"{label}: content должен быть строкой")
        if btype == "list":
            items = item.get("items")
            if not isinstance(items, list) or not items or not all(isinstance(x, (str, int, float)) for x in items):
                errors.append(f"{label}: items должен быть непустым массивом строк")
        if btype == "table":
            headers = item.get("headers", [])
            rows = item.get("rows", [])
            if not isinstance(headers, list) or not isinstance(rows, list) or not all(isinstance(r, list) for r in rows):
                errors.append(f"{label}: headers — массив, rows — массив массивов")
            elif not headers and not rows:
                errors.append(f"{label}: таблица без заголовков и строк")
        if btype == "listing" and not isinstance(item.get("code", ""), str):
            errors.append(f"{label}: code должен быть строкой")
        if btype == "figure":
            path = item.get("path")
            if path is not None and not isinstance(path, str):
                errors.append(f"{label}: path должен быть строкой")
        blocks.append(item)
    return blocks


def normalize_content(data: Any) -> Tuple[Dict[str, Any], List[str]]:
    """Проверяет схему и приводит вход к единому виду. InputError — список всех нарушений."""
    errors: List[str] = []
    warnings: List[str] = []
    if not isinstance(data, dict):
        raise InputError("входной JSON должен быть объектом")
    _check_keys(data, TOP_LEVEL_KEYS, "корень JSON", errors)
    out: Dict[str, Any] = {}
    title = data.get("title")
    if not isinstance(title, str) or not title.strip():
        errors.append("title: обязательна непустая строка — тема работы")
    out["title"] = str(title or "").strip()
    out["title_page"] = _normalize_title_page(data.get("title_page"), errors, warnings)
    for key in SECTION_KEYS:
        out[key] = _normalize_blocks(data.get(key), key, errors)
    chapters = data.get("chapters")
    out["chapters"] = []
    if not isinstance(chapters, list) or not chapters:
        errors.append("chapters: обязателен непустой массив глав")
        chapters = []
    for ci, chapter in enumerate(chapters, 1):
        where = f"chapters[{ci}]"
        if not isinstance(chapter, dict):
            errors.append(f"{where}: глава должна быть объектом")
            continue
        _check_keys(chapter, CHAPTER_KEYS, where, errors)
        ctitle = strip_chapter_prefix(chapter.get("title", ""))
        if not ctitle:
            errors.append(f"{where}.title: обязательно название главы")
        paragraphs = chapter.get("paragraphs", [])
        norm_paragraphs = []
        if not isinstance(paragraphs, list):
            errors.append(f"{where}.paragraphs: ожидается массив")
            paragraphs = []
        for pi, para in enumerate(paragraphs, 1):
            pwhere = f"{where}.paragraphs[{pi}]"
            if not isinstance(para, dict):
                errors.append(f"{pwhere}: параграф должен быть объектом")
                continue
            _check_keys(para, PARAGRAPH_KEYS, pwhere, errors)
            ptitle = strip_paragraph_number(para.get("title", ""))
            if not ptitle:
                errors.append(f"{pwhere}.title: обязательно название параграфа")
            if "text" in para and "blocks" in para:
                errors.append(f"{pwhere}: используйте либо text, либо blocks")
            blocks = _normalize_blocks(para.get("blocks", para.get("text")), pwhere, errors)
            norm_paragraphs.append({"title": ptitle, "blocks": blocks})
        out["chapters"].append({
            "title": ctitle,
            "blocks": _normalize_blocks(chapter.get("blocks"), f"{where}.blocks", errors),
            "paragraphs": norm_paragraphs,
            "conclusion": _normalize_blocks(chapter.get("conclusion"), f"{where}.conclusion", errors),
        })
    if data.get("sources") and data.get("bibliography"):
        errors.append("sources и bibliography вместе недопустимы: оставьте один список литературы")
    sources = data.get("sources")
    if sources is not None and not (isinstance(sources, list) and all(isinstance(s, dict) for s in sources)):
        errors.append("sources: ожидается массив объектов реестра (схема sources.json)")
    bibliography = data.get("bibliography")
    if bibliography is not None and not (isinstance(bibliography, list) and all(isinstance(s, str) for s in bibliography)):
        errors.append("bibliography: ожидается массив готовых строк")
    out["sources"] = list(sources or [])
    out["bibliography"] = list(bibliography or [])
    appendices = data.get("appendices", [])
    out["appendices"] = []
    if not isinstance(appendices, list):
        errors.append("appendices: ожидается массив")
        appendices = []
    for ai, app in enumerate(appendices, 1):
        where = f"appendices[{ai}]"
        if not isinstance(app, dict):
            errors.append(f"{where}: приложение должно быть объектом")
            continue
        _check_keys(app, APPENDIX_KEYS, where, errors)
        atitle = strip_appendix_prefix(app.get("title", ""))
        if not atitle:
            errors.append(f"{where}.title: приложение должно иметь заголовок (методичка, раздел «Оформление приложений»)")
        if "content" in app and "blocks" in app:
            errors.append(f"{where}: используйте либо content, либо blocks")
        out["appendices"].append({
            "title": atitle,
            "label": str(app.get("label") or f"Приложение {ai}").strip(),
            "blocks": _normalize_blocks(app.get("blocks", app.get("content")), where, errors),
        })
    for flag in ("skip_last_page", "keep_source_order", "group_normative_first"):
        value = data.get(flag, False)
        if not isinstance(value, bool):
            errors.append(f"{flag}: ожидается true или false")
        out[flag] = bool(value) if isinstance(value, bool) else False
    if errors:
        raise InputError("\n".join(errors))
    return out, warnings


# ---------------------------------------------------------------------------
# Ссылки и список литературы
# ---------------------------------------------------------------------------

def map_text_fields(content: Dict[str, Any], fn: Callable[[str], str]) -> None:
    """Применяет fn ко всем текстовым полям, кроме кода листингов и заголовков."""

    def blocks(items: List[Dict[str, Any]]) -> None:
        for block in items:
            btype = block["type"]
            if btype == "text":
                block["content"] = fn(str(block.get("content") or ""))
            elif btype == "list":
                block["items"] = [fn(str(x)) for x in block.get("items", [])]
            elif btype == "table":
                block["title"] = fn(str(block.get("title") or ""))
                block["headers"] = [fn(str(x)) for x in block.get("headers", [])]
                block["rows"] = [[fn(str(x)) for x in row] for row in block.get("rows", [])]
                if block.get("source"):
                    block["source"] = fn(str(block["source"]))
            elif btype == "listing":
                block["caption"] = fn(str(block.get("caption") or ""))
            elif btype == "figure":
                block["caption"] = fn(str(block.get("caption") or ""))

    for key in SECTION_KEYS:
        blocks(content[key])
    for chapter in content["chapters"]:
        blocks(chapter["blocks"])
        for para in chapter["paragraphs"]:
            blocks(para["blocks"])
        blocks(chapter["conclusion"])
    for app in content["appendices"]:
        blocks(app["blocks"])


def prepare_bibliography(content: Dict[str, Any], warnings: List[str]) -> List[str]:
    """Возвращает строки списка литературы и переписывает ссылки в тексте."""
    fb = _import_bibliography()
    if content["sources"]:
        try:
            sources = fb.normalize_input_data(content["sources"])
            ids = fb.source_ids(sources)
            if content["keep_source_order"]:
                numbers = fb.check_preserved_ids(sources)
                result = fb.number_sources(sources, keep_order=True)
                mapping = {str(v): n for v, n in zip(ids, numbers)}
                labels = numbers
                if numbers != list(range(1, len(numbers) + 1)):
                    warnings.append("keep_source_order: номера списка не идут подряд с 1")
                problems = fb.alphabetical_problems(result["entries"])
                if problems:
                    warnings.append("keep_source_order: список не в алфавитном порядке: " + problems[0])
            else:
                result = fb.number_sources(sources, group_normative_first=content["group_normative_first"])
                mapping = result["mapping"]
                labels = [e["number"] for e in result["entries"]]
        except fb.BibliographyError as error:
            raise ContentError(f"список литературы: {error}") from error
        warnings.extend(result["warnings"])
        resolver = fb.make_resolver(mapping, ids)
        problems: List[str] = []
        map_text_fields(content, lambda text: fb.rewrite_citations(text, resolver, problems))
        if problems:
            raise ContentError("ссылки на отсутствующие источники: " + "; ".join(sorted(set(problems))))
        return [f"{label}. {entry['text']}" for label, entry in zip(labels, result["entries"])]
    unresolved: List[str] = []

    def check(text: str) -> str:
        for kind, ref in fb.citation_refs(text):
            if kind == "id":
                unresolved.append(f"[@{ref}]")
        return text

    map_text_fields(content, check)
    if unresolved:
        raise ContentError(
            "ссылки " + ", ".join(sorted(set(unresolved))) + " нельзя разрешить без структурированного sources"
        )
    lines = [line.strip() for line in content["bibliography"] if line.strip()]
    if not lines:
        warnings.append("список литературы пуст: нет ни sources, ни bibliography")
        return []
    if not any(re.match(r"^\d+\.\s", line) for line in lines):
        lines = [f"{i}. {line}" for i, line in enumerate(lines, 1)]
    if not content["group_normative_first"]:
        # Сравниваются заголовок и заглавие — часть записи до маркера или сведений об ответственности.
        heads = [re.split(r" \[| / | // ", re.sub(r"^\d+\.\s*", "", line), maxsplit=1)[0] for line in lines]
        keys = [fb.sort_text_key(x) for x in heads]
        if keys != sorted(keys):
            warnings.append("bibliography: готовые строки списка не в алфавитном порядке — генератор их не сортирует")
    return lines


# ---------------------------------------------------------------------------
# python-docx (ленивый импорт)
# ---------------------------------------------------------------------------

Document = Pt = Mm = Cm = Emu = RGBColor = None  # type: ignore
WD_ALIGN_PARAGRAPH = WD_LINE_SPACING = WD_STYLE_TYPE = None  # type: ignore
qn = OxmlElement = parse_xml = nsdecls = None  # type: ignore


def _load_docx() -> None:
    global Document, Pt, Mm, Cm, Emu, RGBColor, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_STYLE_TYPE
    global qn, OxmlElement, parse_xml, nsdecls
    if Document is not None:
        return
    from docx import Document as _Document
    from docx.enum.style import WD_STYLE_TYPE as _WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH as _WD_ALIGN, WD_LINE_SPACING as _WD_LS
    from docx.oxml import OxmlElement as _OxmlElement, parse_xml as _parse_xml
    from docx.oxml.ns import nsdecls as _nsdecls, qn as _qn
    from docx.shared import Cm as _Cm, Emu as _Emu, Mm as _Mm, Pt as _Pt, RGBColor as _RGB

    Document, Pt, Mm, Cm, Emu, RGBColor = _Document, _Pt, _Mm, _Cm, _Emu, _RGB
    WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_STYLE_TYPE = _WD_ALIGN, _WD_LS, _WD_STYLE_TYPE
    qn, OxmlElement, parse_xml, nsdecls = _qn, _OxmlElement, _parse_xml, _nsdecls


# ---------------------------------------------------------------------------
# Страница и стили
# ---------------------------------------------------------------------------

def setup_page_margins(doc) -> None:
    """A4 и поля МПГУ на всех секциях."""
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(35)
        section.right_margin = Mm(10)


def _set_fonts(rpr_parent, name: str) -> None:
    rpr = rpr_parent.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _set_lang(rpr_parent, value: str) -> None:
    rpr = rpr_parent.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), value)
    lang.set(qn("w:eastAsia"), value)


def _get_or_add_style(doc, name: str, based_on: str = "Normal", builtin: bool = False, style_id: str = ""):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=builtin)
        if style_id:
            style.element.styleId = style_id
    if based_on:
        style.base_style = doc.styles[based_on]
    style.quick_style = True
    return style


def _para_format(style, *, align=None, first=None, left=None, before=None, after=None,
                 spacing=None, keep_next=None) -> None:
    pf = style.paragraph_format
    if align is not None:
        pf.alignment = align
    if first is not None:
        pf.first_line_indent = first
    if left is not None:
        pf.left_indent = left
    if before is not None:
        pf.space_before = before
    if after is not None:
        pf.space_after = after
    if spacing == "single":
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    elif spacing == "1.5":
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if keep_next is not None:
        pf.keep_with_next = keep_next


def setup_styles(doc) -> None:
    """Normal, заголовки, оглавление и стили VKR * (SPEC 8)."""
    styles_el = doc.styles.element
    defaults = styles_el.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None and rpr_default.find(qn("w:rPr")) is not None:
            rfonts = rpr_default.find(qn("w:rPr")).find(qn("w:rFonts"))
            if rfonts is not None:
                for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    rfonts.attrib.pop(qn(attr), None)
                for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                    rfonts.set(qn(attr), "Times New Roman")

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    _set_fonts(normal.element, "Times New Roman")
    _set_lang(normal.element, "ru-RU")
    normal.font.size = Pt(14)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    _para_format(normal, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=Cm(1.25), before=Pt(0), after=Pt(0), spacing="1.5")
    normal.paragraph_format.widow_control = True

    for name, align, first, before in (
        ("Heading 1", WD_ALIGN_PARAGRAPH.CENTER, Cm(0), Pt(0)),
        ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT, Cm(1.25), Pt(18)),
        ("Heading 3", WD_ALIGN_PARAGRAPH.LEFT, Cm(1.25), Pt(12)),
    ):
        style = doc.styles[name]
        style.base_style = normal
        style.font.name = "Times New Roman"
        _set_fonts(style.element, "Times New Roman")
        style.font.size = Pt(14)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        _para_format(style, align=align, first=first, before=before, after=Pt(42), spacing="1.5", keep_next=True)
    for name in ("Heading 1 Char", "Heading 2 Char", "Heading 3 Char"):
        try:
            char_style = doc.styles[name]
        except KeyError:
            continue
        char_style.font.name = "Times New Roman"
        _set_fonts(char_style.element, "Times New Roman")
        char_style.font.color.rgb = RGBColor(0, 0, 0)
        char_style.font.size = Pt(14)

    for level, left in ((1, 0.0), (2, 0.5), (3, 1.0)):
        style = _get_or_add_style(doc, f"toc {level}", builtin=True, style_id=f"TOC{level}")
        _para_format(style, align=WD_ALIGN_PARAGRAPH.LEFT, first=Cm(0), left=Cm(left), before=Pt(0), after=Pt(0))
        style.element.attrib.pop(qn("w:customStyle"), None)

    title = _get_or_add_style(doc, STYLE_TITLE)
    _para_format(title, align=WD_ALIGN_PARAGRAPH.CENTER, first=Cm(0), left=Cm(0), before=Pt(0), after=Pt(0), spacing="single")

    listing = _get_or_add_style(doc, STYLE_LISTING)
    listing.font.name = "Consolas"
    _set_fonts(listing.element, "Consolas")
    listing.font.size = Pt(11)
    rpr = listing.element.get_or_add_rPr()
    if rpr.find(qn("w:noProof")) is None:
        rpr.append(OxmlElement("w:noProof"))
    _para_format(listing, align=WD_ALIGN_PARAGRAPH.LEFT, first=Cm(0), left=Cm(1.0), before=Pt(0), after=Pt(0), spacing="single")
    listing.paragraph_format.widow_control = False

    caption = _get_or_add_style(doc, STYLE_CAPTION)
    _para_format(caption, align=WD_ALIGN_PARAGRAPH.CENTER, first=Cm(0), left=Cm(0), before=Pt(6), after=Pt(12), spacing="single")

    bibliography = _get_or_add_style(doc, STYLE_BIBLIOGRAPHY)
    _para_format(bibliography, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=Cm(1.25), left=Cm(0), before=Pt(0), after=Pt(0), spacing="1.5")

    label = _get_or_add_style(doc, STYLE_APPENDIX_LABEL)
    label.font.bold = True
    _para_format(label, align=WD_ALIGN_PARAGRAPH.RIGHT, first=Cm(0), left=Cm(0), before=Pt(0), after=Pt(0), keep_next=True)

    app_title = _get_or_add_style(doc, STYLE_APPENDIX_TITLE)
    app_title.font.bold = True
    _para_format(app_title, align=WD_ALIGN_PARAGRAPH.CENTER, first=Cm(0), left=Cm(0), before=Pt(0), after=Pt(42), keep_next=True)

    list_style = doc.styles[STYLE_LIST]
    pPr = list_style.element.get_or_add_pPr()
    for child in list(pPr):
        if child.tag in (qn("w:contextualSpacing"), qn("w:ind")):
            pPr.remove(child)
    _para_format(list_style, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=Pt(0), after=Pt(0), spacing="1.5")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "1069")
    ind.set(qn("w:hanging"), "360")
    list_style.element.get_or_add_pPr().append(ind)


def add_page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._element.append(element)


def setup_page_numbers(doc) -> None:
    """Номер внизу по центру; на титуле (первая страница, titlePg) не виден.

    Первая секция и так считает страницы с 1, поэтому ``w:start`` не пишется:
    Word копирует его в каждую новую секцию, и нумерация после разрыва раздела,
    вставленного студентом, начиналась бы заново.
    """
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    pg_num = OxmlElement("w:pgNumType")
    sect_pr.insert_element_before(
        pg_num, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange",
    )
    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.paragraph_format.first_line_indent = Cm(0)
    add_page_number_field(footer_para)
    section.first_page_footer.paragraphs[0].text = ""


# ---------------------------------------------------------------------------
# Рендерер
# ---------------------------------------------------------------------------

_INLINE_CODE = re.compile(r"`([^`\n]+)`")


def _join_lines(text: str) -> List[str]:
    """Абзацы по пустой строке; одиночный перевод строки — пробел (D16)."""
    paragraphs = []
    for chunk in re.split(r"\n\s*\n", str(text or "").replace("\r\n", "\n")):
        joined = " ".join(line.strip() for line in chunk.split("\n") if line.strip())
        if joined:
            paragraphs.append(joined)
    return paragraphs


class Renderer:
    def __init__(self, doc, base_dir: Optional[Path], warnings: List[str]):
        self.doc = doc
        self.base_dir = base_dir
        self.warnings = warnings
        self.table_no = 0
        self.figure_no = 0
        self.listing_no = 0
        self.pending_break = False
        self._abstract_ids: Dict[str, int] = {}

    # --- базовые абзацы -----------------------------------------------------
    def paragraph(self, text: str = "", style: Optional[str] = None, *, align=None, bold: bool = False,
                  italic: bool = False, size: Optional[float] = None, page_break: bool = False,
                  inline_code: bool = True):
        para = self.doc.add_paragraph(style=style) if style else self.doc.add_paragraph()
        if align is not None:
            para.alignment = align
        if page_break or self.pending_break:
            para.paragraph_format.page_break_before = True
            self.pending_break = False
        if text:
            self.add_runs(para, text, bold=bold, italic=italic, size=size, inline_code=inline_code)
        return para

    def add_runs(self, para, text: str, *, bold: bool = False, italic: bool = False,
                 size: Optional[float] = None, inline_code: bool = True) -> None:
        pieces: List[Tuple[str, bool]] = []
        if inline_code:
            pos = 0
            for match in _INLINE_CODE.finditer(text):
                if match.start() > pos:
                    pieces.append((text[pos:match.start()], False))
                pieces.append((match.group(1), True))
                pos = match.end()
            pieces.append((text[pos:], False))
        else:
            pieces.append((text, False))
        for piece, is_code in pieces:
            if not piece:
                continue
            run = para.add_run(piece)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if size:
                run.font.size = Pt(size)
            if is_code:
                run.font.name = "Consolas"
                _set_fonts(run._element, "Consolas")

    def section_heading(self, text: str, *, in_toc: bool) -> None:
        if in_toc:
            para = self.doc.add_paragraph(text, style="Heading 1")
            para.paragraph_format.page_break_before = True
        else:
            para = self.paragraph(text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, page_break=True, inline_code=False)
            pf = para.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.space_after = Pt(42)
            pf.keep_with_next = True
        self.pending_break = False

    # --- блоки --------------------------------------------------------------
    def blocks(self, blocks: List[Dict[str, Any]]) -> None:
        for block in blocks:
            btype = block["type"]
            if btype == "text":
                for text in _join_lines(block.get("content", "")):
                    self.paragraph(text)
            elif btype == "list":
                self.list_block(block)
            elif btype == "table":
                self.table_block(block)
            elif btype == "listing":
                self.listing_block(block)
            elif btype == "figure":
                self.figure_block(block)

    def _abstract_num(self, ordered: bool) -> int:
        key = "decimal" if ordered else "bullet"
        if key in self._abstract_ids:
            return self._abstract_ids[key]
        numbering = self.doc.part.numbering_part.element
        existing = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        abstract_id = max(existing + [99]) + 1
        fmt, text = ("decimal", "%1.") if ordered else ("bullet", "–")
        xml = (
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
            '<w:multiLevelType w:val="singleLevel"/>'
            '<w:lvl w:ilvl="0"><w:start w:val="1"/>'
            f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
            '<w:pPr><w:tabs><w:tab w:val="num" w:pos="1069"/></w:tabs><w:ind w:left="1069" w:hanging="360"/></w:pPr>'
            '<w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman" w:hint="default"/></w:rPr>'
            "</w:lvl></w:abstractNum>"
        )
        abstract = parse_xml(xml)
        nums = numbering.findall(qn("w:num"))
        if nums:
            nums[0].addprevious(abstract)
        else:
            numbering.append(abstract)
        self._abstract_ids[key] = abstract_id
        return abstract_id

    def list_block(self, block: Dict[str, Any]) -> None:
        """Каждый список — отдельный w:num с startOverride=1 (REG-D-05)."""
        ordered = bool(block.get("ordered", False))
        numbering = self.doc.part.numbering_part.element
        num = numbering.add_num(self._abstract_num(ordered))
        override = num.add_lvlOverride(ilvl=0)
        override.add_startOverride(1)
        for item in block.get("items", []):
            para = self.paragraph(str(item).strip(), STYLE_LIST)
            num_pr = para._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = num.numId

    def table_block(self, block: Dict[str, Any]) -> None:
        self.table_no += 1
        number = block.get("number") or self.table_no
        title = str(block.get("title") or "").strip().rstrip(".")
        label = self.paragraph(f"Таблица {number}", STYLE_CAPTION, align=WD_ALIGN_PARAGRAPH.RIGHT, inline_code=False)
        label.paragraph_format.space_before = Pt(12)
        label.paragraph_format.space_after = Pt(0)
        label.paragraph_format.keep_with_next = True
        if title:
            head = self.paragraph(title, STYLE_CAPTION, bold=True)
            head.paragraph_format.space_before = Pt(0)
            head.paragraph_format.space_after = Pt(6)
            head.paragraph_format.keep_with_next = True
        else:
            self.warnings.append(f"таблица {number}: нет заголовка (методичка требует заголовок над таблицей)")
        headers = [str(h) for h in block.get("headers", [])]
        rows = [[str(c) for c in row] for row in block.get("rows", [])]
        ncols = max([1, len(headers)] + [len(r) for r in rows])
        table = self.doc.add_table(rows=(1 if headers else 0) + len(rows), cols=ncols)
        table.style = "Table Grid"
        all_rows = ([headers] if headers else []) + rows
        for r_index, row_values in enumerate(all_rows):
            is_header = bool(headers) and r_index == 0
            row = table.rows[r_index]
            if is_header:
                tr_pr = row._tr.get_or_add_trPr()
                tr_pr.append(OxmlElement("w:tblHeader"))
            for c_index in range(ncols):
                value = row_values[c_index] if c_index < len(row_values) else ""
                cell = row.cells[c_index]
                lines = [x.strip() for x in value.split("\n")] or [""]
                cell.text = ""
                for l_index, line in enumerate(lines):
                    para = cell.paragraphs[0] if l_index == 0 else cell.add_paragraph()
                    pf = para.paragraph_format
                    pf.first_line_indent = Cm(0)
                    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
                    if line:
                        self.add_runs(para, line, bold=is_header, size=12)
        if block.get("source"):
            source = str(block["source"]).strip()
            if not source.lower().startswith("источник"):
                source = f"Источник: {source}"
            para = self.paragraph(source, STYLE_CAPTION, align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
            para.paragraph_format.space_before = Pt(6)

    def listing_block(self, block: Dict[str, Any]) -> None:
        code = str(block.get("code") or "").replace("\r\n", "\n").rstrip("\n")
        if not code.strip():
            self.warnings.append("пустой листинг пропущен")
            return
        self.listing_no += 1
        number = block.get("number") or self.listing_no
        lines = code.split("\n")
        for index, line in enumerate(lines):
            para = self.doc.add_paragraph(style=STYLE_LISTING)
            if self.pending_break:
                para.paragraph_format.page_break_before = True
                self.pending_break = False
            run = para.add_run(line if line else " ")
            run.font.name = "Consolas"
            if index == len(lines) - 1:
                para.paragraph_format.keep_with_next = True
        caption = str(block.get("caption") or "").strip().rstrip(".")
        text = f"Листинг {number}" + (f" — {caption}" if caption else "")
        if not caption:
            self.warnings.append(f"листинг {number}: нет подписи")
        self.paragraph(text, STYLE_CAPTION, italic=True, size=12)

    def figure_block(self, block: Dict[str, Any]) -> None:
        self.figure_no += 1
        number = block.get("number") or self.figure_no
        caption = str(block.get("caption") or "").strip().rstrip(".")
        path_text = str(block.get("path") or "").strip()
        holder = self.paragraph(style=STYLE_CAPTION, align=WD_ALIGN_PARAGRAPH.CENTER)
        holder.paragraph_format.space_before = Pt(12)
        holder.paragraph_format.space_after = Pt(0)
        holder.paragraph_format.keep_with_next = True
        if path_text and block.get("placeholder") is not True:
            path = Path(path_text)
            if not path.is_absolute():
                path = (self.base_dir or Path.cwd()) / path
            if not path.is_file():
                raise ContentError(f"рисунок {number}: файл не найден: {path_text}")
            width, height = self._figure_size(path, block)
            try:
                holder.add_run().add_picture(str(path), width=width, height=height)
            except Exception as error:  # UnrecognizedImageError и ошибки чтения
                raise ContentError(
                    f"рисунок {number}: не удалось вставить {path_text} ({type(error).__name__}); "
                    "используйте PNG или JPEG"
                ) from error
        else:
            holder.add_run(FIGURE_PLACEHOLDER).italic = True
            self.warnings.append(f"рисунок {number}: вместо изображения заглушка")
        text = f"Рисунок {number}" + (f" — {caption}" if caption else "")
        if not caption:
            self.warnings.append(f"рисунок {number}: нет подписи")
        self.paragraph(text, STYLE_CAPTION)

    def _figure_size(self, path: Path, block: Dict[str, Any]):
        from docx.image.image import Image

        try:
            image = Image.from_file(str(path))
        except Exception as error:
            raise ContentError(f"не удалось прочитать изображение {path.name}: {type(error).__name__}; используйте PNG или JPEG") from error
        native_w, native_h = int(image.width), int(image.height)
        max_w = int(Mm(TEXT_WIDTH_MM))
        if block.get("width_cm"):
            try:
                max_w = min(max_w, int(Cm(float(block["width_cm"]))))
            except (TypeError, ValueError):
                self.warnings.append(f"рисунок {path.name}: width_cm должен быть числом")
        max_h = int(Mm(MAX_FIGURE_HEIGHT_MM))
        scale = min(1.0, max_w / float(native_w or 1), max_h / float(native_h or 1))
        if block.get("width_cm"):
            scale = min(max_w / float(native_w or 1), max_h / float(native_h or 1))
        return Emu(int(native_w * scale)), Emu(int(native_h * scale))

    # --- разделы ------------------------------------------------------------
    def title_page(self, tp: Dict[str, Any], topic: str) -> List[str]:
        """Титул по образцу ИМО 09.03.02; пустое обязательное поле — маркер [ЗАПОЛНИТЬ: …]."""
        missing: List[str] = []

        def value(key: str, label: str, default: str = "") -> str:
            text = str(tp.get(key) or "").strip()
            if text:
                return text
            if default:
                return default
            missing.append(key)
            return f"[ЗАПОЛНИТЬ: {label}]"

        def line(text: str, align, *, bold: bool = False, spacing: Optional[float] = None):
            para = self.doc.add_paragraph(style=STYLE_TITLE)
            para.alignment = align
            if spacing:
                para.paragraph_format.line_spacing = spacing
            if text:
                run = para.add_run(text)
                if bold:
                    run.bold = True
            return para

        center, right, left = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT
        university = str(tp.get("university") or DEFAULT_UNIVERSITY).strip().strip("«»\"")
        for text in ("Министерство просвещения Российской Федерации", "федеральное государственное бюджетное",
                     "образовательное учреждение высшего образования", f"«{university}»"):
            line(text, center, spacing=1.15)
        line("", center)
        authors = tp.get("authors") if isinstance(tp.get("authors"), list) and tp.get("authors") else [tp.get("author")]
        for name in authors:
            line(str(name or "").strip() or "[ЗАПОЛНИТЬ: ФИО автора]", right, bold=True)
        group = str(tp.get("group") or "").strip()
        if group:
            group = re.sub(r"^\s*группа\s*:?\s*", "", group, flags=re.IGNORECASE)
            line(f"Группа {group}", right)
        line("", center)
        line(value("institute", "институт"), center)
        line("", center)
        line(topic, center, bold=True)
        line("", center)
        code = value("program_code", "код направления")
        name = value("program_name", "направление подготовки")
        line(f"Код и направление подготовки: {code.rstrip('.')}. {name}", center)
        line("Направленность (профиль) образовательной программы:", center)
        profile = value("program_profile", "профиль")
        line(profile if profile.endswith((".", "]")) else f"{profile}.", center)
        line("", center)
        work_type = str(tp.get("work_type") or "").strip()
        base = "Выпускная квалификационная работа"
        rest = work_type
        if work_type.casefold().startswith(base.casefold()):
            rest = work_type[len(base):].strip()
        rest = rest.strip().strip("()").strip() or "бакалаврская работа"
        line(base, center)
        line(f"({rest})", center)
        line("", center)
        line("Научный руководитель –", left)
        line(value("supervisor_title", "должность, учёная степень научного руководителя"), left)
        line(value("supervisor", "И.О. Фамилия научного руководителя"), left)
        line("", center)
        line("«Допустить к защите»", left)
        line("Заведующий кафедрой", left)
        dept = value("department", "кафедра")
        dept = re.sub(r"^\s*кафедр(?:а|ой|ы)\s+", "", dept, flags=re.IGNORECASE)
        line(dept, left)
        line("______________________________", left)
        line(value("head_title", "учёная степень, звание заведующего кафедрой"), left)
        line(value("head_of_department", "И.О. Фамилия заведующего кафедрой"), left)
        line("", center)
        line("Проверка на объем заимствований:", left)
        line("_______% авторского текста", left)
        line("", center)
        year = str(tp.get("year") or dt.date.today().year).strip()
        city = str(tp.get("city") or "Москва").strip()
        line(f"{city} – {year} год", center)
        return missing

    def contents(self) -> None:
        self.section_heading("Содержание", in_toc=False)
        para = self.doc.add_paragraph()
        para.paragraph_format.first_line_indent = Cm(0)
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = para.add_run()
        for kind, text in (("begin", None), (None, r' TOC \o "1-3" \h \z \u \f '), ("separate", None)):
            if kind:
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), kind)
            else:
                element = OxmlElement("w:instrText")
                element.set(qn("xml:space"), "preserve")
                element.text = text
            run._element.append(element)
        placeholder = para.add_run(TOC_PLACEHOLDER)
        placeholder.font.italic = True
        placeholder.font.size = Pt(12)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        placeholder._element.append(end)

    def chapter(self, chapter: Dict[str, Any], index: int) -> None:
        roman = to_roman(index)
        heading = self.doc.add_paragraph(f"Глава {roman}. {chapter['title']}", style="Heading 1")
        heading.paragraph_format.page_break_before = True
        self.blocks(chapter["blocks"])
        for p_index, para in enumerate(chapter["paragraphs"], 1):
            self.doc.add_paragraph(f"{index}.{p_index}. {para['title']}", style="Heading 2")
            self.blocks(para["blocks"])
        if chapter["conclusion"]:
            head = self.paragraph(f"Выводы по главе {roman}", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, inline_code=False)
            pf = head.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.space_before = Pt(18)
            pf.space_after = Pt(12)
            pf.keep_with_next = True
            self.blocks(chapter["conclusion"])

    def bibliography(self, lines: List[str]) -> None:
        heading = self.doc.add_paragraph(BIBLIOGRAPHY_HEADING, style="Heading 1")
        heading.paragraph_format.page_break_before = True
        for entry in lines:
            self.doc.add_paragraph(entry, style=STYLE_BIBLIOGRAPHY)

    def appendix(self, app: Dict[str, Any]) -> None:
        label = self.doc.add_paragraph(app["label"], style=STYLE_APPENDIX_LABEL)
        label.paragraph_format.page_break_before = True
        title = self.doc.add_paragraph(app["title"], style=STYLE_APPENDIX_TITLE)
        toc_text = f"{app['label']}. {app['title']}".replace('"', "»").replace("\\", "/")
        # Поле TC без скрытого форматирования: Word (проверено через COM) включает его в
        # оглавление с ключом \f и не печатает; скрытое поле TC в оглавление не попадает.
        run = title.add_run()
        for kind, text in (("begin", None), (None, f' TC "{toc_text}" \\l 1 '), ("end", None)):
            if kind:
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), kind)
            else:
                element = OxmlElement("w:instrText")
                element.set(qn("xml:space"), "preserve")
                element.text = text
            run._element.append(element)
        self.blocks(app["blocks"])

    def last_page(self, authors: List[str]) -> None:
        """«Последний лист ВКР» — клятва о самостоятельном выполнении (с новой страницы)."""
        head = self.doc.add_paragraph("Выпускная квалификационная работа", style=STYLE_TITLE)
        head.runs[0].bold = True
        head.paragraph_format.page_break_before = True
        head.paragraph_format.space_after = Pt(18)
        authors = authors or [""]
        oath = (
            "Выпускная квалификационная работа выполнена мной совершенно самостоятельно. "
            "Все использованные в работе материалы и концепции из опубликованной научной "
            "литературы и других источников имеют ссылки на них."
        )
        table = self.doc.add_table(rows=len(authors), cols=1)
        table.style = "Table Grid"
        for index in range(len(authors)):
            cell_para = table.rows[index].cells[0].paragraphs[0]
            cell_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            cell_para.paragraph_format.first_line_indent = Cm(1.25)
            cell_para.add_run(oath)
        self.doc.add_paragraph(style=STYLE_TITLE)
        signatures = self.doc.add_table(rows=len(authors), cols=3)
        signatures.style = "Table Grid"
        for index, name in enumerate(authors):
            for c_index, text in enumerate((name or "___________________", "___________________", "___________________")):
                cell_para = signatures.rows[index].cells[c_index].paragraphs[0]
                cell_para.paragraph_format.first_line_indent = Cm(0)
                cell_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                cell_para.add_run(text).font.size = Pt(12)
        note = self.doc.add_paragraph("(Ф.И.О.)        (подпись)        (дата)", style=STYLE_TITLE)
        note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        note.runs[0].italic = True
        note.runs[0].font.size = Pt(10)


def build_vkr_docx(data: dict, output_path, *, base_dir=None, keep_source_order: Optional[bool] = None,
                   group_normative_first: Optional[bool] = None, warnings: Optional[List[str]] = None):
    """Построить DOCX. InputError — схема входа, ContentError — содержание."""
    warnings = [] if warnings is None else warnings
    content, schema_warnings = normalize_content(copy.deepcopy(data))
    warnings.extend(schema_warnings)
    if keep_source_order is not None:
        content["keep_source_order"] = bool(keep_source_order)
    if group_normative_first is not None:
        content["group_normative_first"] = bool(group_normative_first)
    lines = prepare_bibliography(content, warnings)

    _load_docx()
    output_path = Path(output_path)
    doc = Document()
    setup_page_margins(doc)
    setup_styles(doc)
    setup_page_numbers(doc)
    tp = content["title_page"]
    now = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None, microsecond=0)
    authors = tp.get("authors") if isinstance(tp.get("authors"), list) and tp.get("authors") else [str(tp.get("author") or "")]
    doc.core_properties.author = ", ".join(a for a in authors if a)
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = content["title"]
    doc.core_properties.subject = "Выпускная квалификационная работа"
    doc.core_properties.comments = ""
    doc.core_properties.created = now
    doc.core_properties.modified = now

    renderer = Renderer(doc, Path(base_dir) if base_dir else None, warnings)
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            body.remove(child)
    missing = renderer.title_page(tp, content["title"])
    if missing:
        warnings.append("титульный лист: не заполнены поля " + ", ".join(missing) + " — на титуле стоят маркеры [ЗАПОЛНИТЬ: …]")
    if content["annotation"]:
        renderer.section_heading("Аннотация", in_toc=False)
        renderer.blocks(content["annotation"])
    else:
        warnings.append("нет аннотации")
    renderer.contents()
    if content["introduction"]:
        renderer.section_heading("Введение", in_toc=True)
        renderer.blocks(content["introduction"])
    else:
        warnings.append("нет введения")
    for index, chapter in enumerate(content["chapters"], 1):
        renderer.chapter(chapter, index)
    if content["conclusion"]:
        renderer.section_heading("Заключение", in_toc=True)
        renderer.blocks(content["conclusion"])
    else:
        warnings.append("нет заключения")
    if lines:
        renderer.bibliography(lines)
    for app in content["appendices"]:
        renderer.appendix(app)
    if not content["skip_last_page"]:
        renderer.last_page([a for a in authors if a])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return doc


def _configure_streams() -> None:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8", errors="replace")
            except (ValueError, OSError):
                pass


def default_base_dir(input_path: Path) -> Path:
    parent = input_path.resolve().parent
    return parent.parent if parent.name == "exports" else parent


def main(argv: Optional[List[str]] = None) -> int:
    _configure_streams()
    parser = argparse.ArgumentParser(description="Генератор DOCX ВКР по методичке МПГУ (рендерер JSON)")
    parser.add_argument("input_json", help="JSON с содержимым ВКР (references/docx-input-schema.md)")
    parser.add_argument("-o", "--output", default="vkr.docx", help="Выходной DOCX (по умолчанию vkr.docx)")
    parser.add_argument("--base-dir", help="Каталог, от которого считаются пути рисунков "
                                           "(по умолчанию каталог JSON; для exports/ — корень проекта)")
    parser.add_argument("--keep-source-order", action="store_true",
                        help="Не сортировать sources: номер записи = её целый id")
    parser.add_argument("--group-normative-first", action="store_true",
                        help="Нормативные акты и стандарты — в начало списка литературы")
    parser.add_argument("--json", action="store_true", help="Результат в stdout как JSON")
    args = parser.parse_args(argv)

    def finish(code: int, errors: List[str], warnings: List[str]) -> int:
        if args.json:
            print(json.dumps({"status": "ok" if code == 0 else "error", "exit_code": code,
                              "output": str(args.output) if code == 0 else None,
                              "errors": errors, "warnings": warnings}, ensure_ascii=False, indent=2))
        else:
            for message in errors:
                print(f"ОШИБКА: {message}", file=sys.stderr)
            for message in warnings:
                print(f"ПРЕДУПРЕЖДЕНИЕ: {message}", file=sys.stderr)
            if code == 0:
                print(f"OK: создан {args.output}")
                print("Дальше: откройте файл в Word и обновите поля (Ctrl+A, F9) — оглавление появится; "
                      "затем validate_vkr.py. Подпись и дату на последнем листе ставят от руки.")
        return code

    in_path = Path(args.input_json)
    if not in_path.is_file():
        return finish(2, [f"файл не найден: {in_path}"], [])
    try:
        data = json.loads(in_path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as error:
        return finish(2, [f"не удалось прочитать JSON {in_path}: {error}"], [])
    base_dir = Path(args.base_dir) if args.base_dir else default_base_dir(in_path)
    warnings: List[str] = []
    try:
        build_vkr_docx(
            data, Path(args.output), base_dir=base_dir,
            keep_source_order=True if args.keep_source_order else None,
            group_normative_first=True if args.group_normative_first else None,
            warnings=warnings,
        )
    except InputError as error:
        return finish(2, str(error).split("\n"), warnings)
    except ContentError as error:
        return finish(1, str(error).split("\n"), warnings)
    except ImportError:
        return finish(2, ["не установлен python-docx: pip install python-docx"], warnings)
    except OSError as error:
        return finish(2, [f"не удалось записать {args.output}: {error}"], warnings)
    return finish(0, [], warnings)


if __name__ == "__main__":
    raise SystemExit(main())
