"""Тесты update_docx_fields.py: коды выхода, ручной путь, открытый в Word документ, процесс Word
этого запуска, exports/docx-fields.json и (по желанию, VKR_TEST_WORD=1) реальный Word."""

from __future__ import annotations

import hashlib
import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import types
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "vkr-mpgu" / "scripts"
SCRIPT = SCRIPTS / "update_docx_fields.py"
BUILD = SCRIPTS / "build_vkr.py"
FIXTURE_PROJECT = ROOT / "tests" / "fixtures" / "build" / "project"
ENV = dict(os.environ, PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8")

sys.dont_write_bytecode = True
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))


def load_module():
    spec = importlib.util.spec_from_file_location("update_docx_fields_under_test", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def make_minimal_docx(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w:document xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'/>")


def make_project_docx(project: Path) -> Path:
    """Проект с final/vkr.docx (нужен python-docx)."""
    from docx import Document

    (project / "final").mkdir(parents=True)
    (project / "vkr-project.json").write_text("{}", encoding="utf-8")
    doc = Document()
    doc.add_paragraph("Введение")
    doc.add_paragraph("Текст работы.")
    target = project / "final" / "vkr.docx"
    doc.save(str(target))
    return target


def sha(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class word_like_lock:
    """Открывает файл так, как Word держит открытый документ (общий доступ только на чтение)."""

    def __init__(self, path: Path) -> None:
        self.path = Path(path)

    def __enter__(self):
        import ctypes
        from ctypes import wintypes

        self.kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        create = self.kernel32.CreateFileW
        create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                           wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        create.restype = wintypes.HANDLE
        self.kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
        self.handle = create(str(self.path), 0x80000000 | 0x40000000, 0x1, None, 3, 0x80, None)
        if self.handle is None or self.handle == ctypes.c_void_p(-1).value:
            raise OSError(ctypes.get_last_error(), f"не удалось открыть {self.path}")
        return self

    def __exit__(self, *exc):
        self.kernel32.CloseHandle(self.handle)
        return False


class FakeWord:
    """Подмена subprocess внутри модуля: «PowerShell с Word» сохраняет копию документа."""

    def __init__(self, pages: int = 12, pid_line: str = "window:4242", timeout: bool = False,
                 open_error: str = "") -> None:
        self.pages = pages
        self.pid_line = pid_line
        self.timeout = timeout
        self.open_error = open_error
        self.calls: list = []

    def run(self, command, **kwargs):
        command = [str(part) for part in command]
        self.calls.append(command)
        if "-File" in command:
            source = Path(command[command.index("-Source") + 1])
            saved = Path(command[command.index("-Saved") + 1])
            Path(command[command.index("-PidFile") + 1]).write_text(self.pid_line, encoding="ascii")
            if self.timeout:
                raise subprocess.TimeoutExpired(command, kwargs.get("timeout"))
            if self.open_error:
                # Word запустился, но Documents.Open дал COM-исключение (как на повреждённом DOCX).
                return subprocess.CompletedProcess(command, 3, stdout=f"open_error={self.open_error}\n".encode(), stderr=b"")
            shutil.copyfile(str(source), str(saved))
            return subprocess.CompletedProcess(command, 0, stdout=f"opened=1\nword_pid=4242\npages={self.pages}\n".encode(), stderr=b"")
        if command[0] == "tasklist":
            pid = command[2].split()[-1]
            return subprocess.CompletedProcess(command, 0, stdout=f'"WINWORD.EXE","{pid}","Console","1","90 K"\r\n'.encode(), stderr=b"")
        return subprocess.CompletedProcess(command, 0, stdout=b"", stderr=b"")

    def install(self, module) -> None:
        module.subprocess = types.SimpleNamespace(run=self.run, PIPE=subprocess.PIPE, DEVNULL=subprocess.DEVNULL,
                                                  TimeoutExpired=subprocess.TimeoutExpired,
                                                  CompletedProcess=subprocess.CompletedProcess)
        module.find_powershell = lambda: "powershell.exe"

    def word_started(self) -> bool:
        return any("-File" in call for call in self.calls)


def run_main(module, *args):
    buffer = io.StringIO()
    with redirect_stdout(buffer):
        code = module.main([str(arg) for arg in args])
    return code, buffer.getvalue()


class UpdateFieldsCliTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="vkr-fields-test-"))

    def tearDown(self) -> None:
        shutil.rmtree(self.tmp, ignore_errors=True)

    def run_cli(self, *args: str) -> subprocess.CompletedProcess:
        return subprocess.run([sys.executable, str(SCRIPT), *args], capture_output=True, env=ENV, check=False)

    def test_missing_file_is_usage_error(self) -> None:
        result = self.run_cli(str(self.tmp / "absent.docx"), "--json")
        self.assertEqual(2, result.returncode, result.stdout.decode("utf-8", "replace"))
        self.assertIn('"status": "error"', result.stdout.decode("utf-8"))

    def test_not_a_docx_is_usage_error(self) -> None:
        bogus = self.tmp / "text.docx"
        bogus.write_text("not a zip", encoding="utf-8")
        self.assertEqual(2, self.run_cli(str(bogus)).returncode)

    def test_manual_path_when_word_unavailable(self) -> None:
        module = load_module()
        docx = self.tmp / "vkr.docx"
        make_minimal_docx(docx)
        before = docx.read_bytes()
        module.find_powershell = lambda: None
        buffer = io.StringIO()
        with redirect_stdout(buffer):
            code = module.main([str(docx), "--json"])
        self.assertEqual(1, code)
        self.assertIn("manual_required", buffer.getvalue())
        self.assertIn("F9", buffer.getvalue())
        self.assertEqual(before, docx.read_bytes(), "без Word файл не должен меняться")

    @unittest.skipUnless(os.name == "nt", "блокировка файла, как у открытого в Word документа, — только Windows")
    def test_docx_open_in_word_is_refused_before_word_starts(self) -> None:
        # A4: код 2 «закрой файл в Word», Word не запускается, в final/ ничего лишнего.
        module = load_module()
        project = self.tmp / "project"
        docx = make_project_docx(project)
        fake = FakeWord()
        fake.install(module)
        before = docx.read_bytes()
        with word_like_lock(docx):
            code, output = run_main(module, docx, "--json")
        self.assertEqual(2, code, output)
        self.assertIn("закрой файл в Word", json.loads(output)["error"])
        self.assertFalse(fake.word_started(), "Word не должен запускаться для заблокированного файла")
        self.assertEqual(before, docx.read_bytes())
        self.assertEqual(["vkr.docx"], sorted(p.name for p in (project / "final").iterdir()))
        self.assertFalse((project / "backups").exists())
        self.assertFalse((project / "exports").exists())

    def test_failed_replace_removes_intermediate_file_and_backup(self) -> None:
        # A4: промежуточный .vkr.fields-* удаляется и при сбое замены (прежде оставался в final/).
        module = load_module()
        project = self.tmp / "project"
        docx = make_project_docx(project)
        FakeWord().install(module)
        before = docx.read_bytes()

        def refuse(source, target):
            raise PermissionError(13, "Access is denied", str(target))

        with mock.patch.object(module.os, "replace", refuse):
            code, output = run_main(module, docx, "--json")
        self.assertEqual(2, code, output)
        self.assertIn("закрой файл в Word", json.loads(output)["error"])
        self.assertEqual(before, docx.read_bytes())
        self.assertEqual(["vkr.docx"], sorted(p.name for p in (project / "final").iterdir()))
        checkpoints = project / "backups" / "checkpoints"
        self.assertEqual([], sorted(checkpoints.iterdir()) if checkpoints.is_dir() else [])

    def test_fields_record_written_after_update(self) -> None:
        # A8: exports/docx-fields.json {docx_sha256, pages, updated_at} после успешного обновления.
        import vkr_common

        module = load_module()
        project = self.tmp / "project"
        docx = make_project_docx(project)
        FakeWord(pages=37).install(module)
        code, output = run_main(module, docx, "--json")
        self.assertEqual(0, code, output)
        payload = json.loads(output)
        self.assertEqual("exports/docx-fields.json", payload["fields_record"])
        self.assertEqual({"source": "window", "pids": [4242]}, payload["word_pid"])
        record = json.loads((project / "exports" / "docx-fields.json").read_text(encoding="utf-8"))
        self.assertEqual({"schema", "docx", "docx_sha256", "pages", "word_opened", "updated_at",
                          "fingerprint_schema", "text_fingerprint"}, set(record))
        self.assertIs(True, record["word_opened"])  # T2: число страниц — после успешного открытия в Word
        self.assertEqual(("final/vkr.docx", 37), (record["docx"], record["pages"]))
        self.assertEqual(sha(docx), record["docx_sha256"])
        self.assertIsNotNone(vkr_common.parse_iso8601(record["updated_at"]))
        self.assertEqual(vkr_common.docx_text_fingerprint(docx)[0], record["text_fingerprint"])
        self.assertEqual(["docx-fields.json"], sorted(p.name for p in (project / "exports").iterdir()))
        self.assertEqual(["vkr.docx"], sorted(p.name for p in (project / "final").iterdir()))
        outside = self.tmp / "loose.docx"  # вне проекта запись не делается
        shutil.copyfile(str(docx), str(outside))
        code, output = run_main(module, outside, "--json")
        self.assertEqual(0, code, output)
        self.assertNotIn("fields_record", json.loads(output))

    def test_word_refusing_to_open_is_not_the_same_as_no_word(self) -> None:
        # T2: «Word не открыл файл» — код 2, status open_failed и подсказка про docx_integrity.py;
        # «нет Word» остаётся кодом 1 с ручной инструкцией. Прежде оба случая давали 1/manual_required.
        module = load_module()
        project = self.tmp / "project"
        docx = make_project_docx(project)
        fake = FakeWord(open_error="Ошибка при попытке открытия файла.")
        fake.install(module)
        before = docx.read_bytes()
        code, output = run_main(module, docx, "--json")
        self.assertEqual(2, code, output)
        payload = json.loads(output)
        self.assertEqual("open_failed", payload["status"])
        self.assertIs(False, payload["word_opened"])
        self.assertIn("docx_integrity.py", payload["reason"])
        self.assertIn("Ошибка при попытке открытия файла", payload["word_error"])
        self.assertEqual(before, docx.read_bytes(), "повреждённый файл не должен меняться")
        self.assertFalse((project / "exports").exists(), "без открытия в Word число страниц не записывается")
        self.assertEqual([["taskkill", "/PID", "4242", "/F"]], [call for call in fake.calls if call[0] == "taskkill"])
        # «Нет Word» — по-прежнему другой исход.
        module = load_module()
        module.find_powershell = lambda: None
        code, output = run_main(module, docx, "--json")
        self.assertEqual(1, code, output)
        self.assertEqual("manual_required", json.loads(output)["status"])
        self.assertIn("$doc.Fields.Update()", module.POWERSHELL_SCRIPT)
        self.assertIn("open_error=", module.POWERSHELL_SCRIPT)  # исход открытия различается в самом скрипте

    @unittest.skipUnless(os.name == "nt" and os.environ.get("VKR_TEST_WORD") == "1",
                         "реальный Word проверяется только при VKR_TEST_WORD=1 на Windows")
    def test_real_word_refuses_broken_docx(self) -> None:
        # T2: настоящий Word на DOCX с текстом вне run — open_failed и код 2 (а не «нет Word»).
        broken = self.tmp / "broken.docx"
        source = ROOT / "tests" / "fixtures" / "build" / "vkr-word-saved.docx"
        with zipfile.ZipFile(source) as archive:
            parts = [(item.filename, archive.read(item.filename)) for item in archive.infolist()]
        with zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as out:
            for name, data in parts:
                if name == "word/document.xml":
                    data = data.decode("utf-8").replace("<w:body>", "<w:body><w:p><w:t>висячий</w:t></w:p>", 1).encode("utf-8")
                out.writestr(name, data)
        result = self.run_cli(str(broken), "--json")
        payload = json.loads(result.stdout.decode("utf-8"))
        self.assertEqual(2, result.returncode, result.stdout.decode("utf-8", "replace"))
        self.assertEqual("open_failed", payload["status"], payload)
        self.assertIn("docx_integrity.py", payload["reason"])
        # Тот же файл ловит структурная проверка без Word.
        import docx_integrity  # noqa: WPS433

        self.assertTrue([p for p in docx_integrity.docx_package_problems(broken) if "WordprocessingML" in p])

    def test_only_word_of_this_run_is_stopped(self) -> None:
        # A5: PID окна документа; «shared» (Word пользователя) и не-WINWORD не останавливаются.
        module = load_module()
        calls: list = []
        alive = {111, 333, 444}

        def fake_run(command, **kwargs):
            command = [str(part) for part in command]
            calls.append(command)
            if command[0] == "tasklist":
                pid = int(command[2].split()[-1])
                out = f'"WINWORD.EXE","{pid}","Console","1","90 K"\r\n' if pid in alive else "INFO: No tasks are running.\r\n"
                return subprocess.CompletedProcess(command, 0, stdout=out.encode("ascii"), stderr=b"")
            return subprocess.CompletedProcess(command, 0, stdout=b"", stderr=b"")

        module.subprocess = types.SimpleNamespace(run=fake_run, PIPE=subprocess.PIPE, DEVNULL=subprocess.DEVNULL)
        pid_file = self.tmp / "word-pid.txt"
        cases = {"window:111": [111], "diff:333\nwindow:111": [111], "shared:111": [], "diff:333,444": [333, 444],
                 "window:555": [], "": [], "111": []}
        for content, expected in cases.items():
            with self.subTest(content=content):
                pid_file.write_text(content, encoding="ascii")
                calls.clear()
                self.assertEqual(expected, module.kill_recorded_word(pid_file))
                self.assertEqual(expected, [int(call[2]) for call in calls if call[0] == "taskkill"])
        self.assertIn("GetWindowThreadProcessId", module.POWERSHELL_SCRIPT)
        self.assertIn("$doc.ActiveWindow.Hwnd", module.POWERSHELL_SCRIPT)

    def test_timeout_reports_and_stops_window_process(self) -> None:
        module = load_module()
        project = self.tmp / "project"
        docx = make_project_docx(project)
        fake = FakeWord(pid_line="diff:777,778\nwindow:4242", timeout=True)
        fake.install(module)
        before = docx.read_bytes()
        code, output = run_main(module, docx, "--json")
        self.assertEqual(1, code, output)
        payload = json.loads(output)
        self.assertEqual({"source": "window", "pids": [4242], "stopped": [4242]}, payload["word_pid"])
        self.assertEqual([["taskkill", "/PID", "4242", "/F"]], [call for call in fake.calls if call[0] == "taskkill"])
        self.assertEqual(before, docx.read_bytes())

    @unittest.skipUnless(os.name == "nt" and os.environ.get("VKR_TEST_WORD") == "1",
                         "реальный Word проверяется только при VKR_TEST_WORD=1 на Windows")
    def test_real_word_updates_toc(self) -> None:
        project = self.tmp / "project"
        shutil.copytree(FIXTURE_PROJECT, project)
        build = subprocess.run([sys.executable, str(BUILD), str(project)], capture_output=True, env=ENV, check=False)
        self.assertEqual(0, build.returncode, build.stdout.decode("utf-8", "replace"))
        docx = project / "final" / "vkr.docx"
        result = self.run_cli(str(docx), "--json")
        self.assertEqual(0, result.returncode, result.stdout.decode("utf-8", "replace"))
        payload = json.loads(result.stdout.decode("utf-8"))
        with zipfile.ZipFile(docx) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        self.assertNotIn("Для обновления содержания", xml)
        self.assertTrue(list((project / "backups" / "checkpoints").glob("*-fields/vkr.docx")))
        self.assertEqual("window", payload["word_pid"]["source"], payload)
        module = load_module()
        time.sleep(2)
        self.assertFalse(any(module.is_word_process(pid) for pid in payload["word_pid"]["pids"]), "Word этого запуска закрыт")
        record = json.loads((project / "exports" / "docx-fields.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["pages"], record["pages"])
        self.assertGreater(record["pages"], 5)
        self.assertEqual(sha(docx), record["docx_sha256"])
        self.assertIs(True, record["word_opened"])
        # T1: то же число Word записал в сам DOCX (docProps/app.xml) — doctor сверяет запись с файлом
        import vkr_common  # noqa: WPS433

        self.assertEqual((record["pages"], None), vkr_common.docx_app_pages(docx))
        # и очистка метаданных его сохраняет
        clean = subprocess.run([sys.executable, str(SCRIPTS / "clean_docx_metadata.py"), str(docx), "--in-place", "--json"],
                               capture_output=True, env=ENV, check=False)
        self.assertEqual(0, clean.returncode, clean.stdout.decode("utf-8", "replace"))
        self.assertEqual((record["pages"], None), vkr_common.docx_app_pages(docx))

    @unittest.skipUnless(os.name == "nt" and os.environ.get("VKR_TEST_WORD") == "1",
                         "реальный Word проверяется только при VKR_TEST_WORD=1 на Windows")
    def test_real_word_timeout_stops_only_own_instance(self) -> None:
        # A5: другой Word автоматизации запускается между снимком процессов и New-Object (гонка, которую
        # разница списков принимала за «свой» Word). При тайм-ауте останавливается только процесс окна документа.
        module = load_module()
        project = self.tmp / "project"
        docx = make_project_docx(project)
        bystander_pid_file = self.tmp / "bystander-pid.txt"
        diff_file = self.tmp / "diff.txt"
        bystander = self.tmp / "bystander.ps1"
        bystander.write_text(
            "param([string]$PidFile)\n$ErrorActionPreference = 'Stop'\n"
            "Add-Type -TypeDefinition 'using System; using System.Runtime.InteropServices; public static class ByUser32 "
            "{ [DllImport(\"user32.dll\")] public static extern uint GetWindowThreadProcessId(IntPtr hWnd, out uint processId); }'\n"
            "$w = New-Object -ComObject Word.Application\n$w.Visible = $false\n$d = $w.Documents.Add()\n"
            "[uint32]$found = 0\n[void][ByUser32]::GetWindowThreadProcessId([IntPtr][int64]$d.ActiveWindow.Hwnd, [ref]$found)\n"
            "Set-Content -LiteralPath $PidFile -Value $found -Encoding ASCII\nStart-Sleep -Seconds 300\n"
            "$d.Close(0)\n$w.Quit()\n", encoding="utf-8-sig")
        start_bystander = (
            "Start-Process -FilePath powershell.exe -WindowStyle Hidden -ArgumentList @('-NoProfile','-NonInteractive',"
            f"'-ExecutionPolicy','Bypass','-File','{bystander}','-PidFile','{bystander_pid_file}')\n"
            f"$waited = 0; while (-not (Test-Path -LiteralPath '{bystander_pid_file}') -and $waited -lt 240) "
            "{ Start-Sleep -Milliseconds 500; $waited++ }\nStart-Sleep -Seconds 1\n")
        script = module.POWERSHELL_SCRIPT.replace(
            "$word = New-Object -ComObject Word.Application\n", start_bystander + "$word = New-Object -ComObject Word.Application\n", 1)
        script = script.replace(
            'Set-Content -LiteralPath $PidFile -Value ("diff:" + ($fresh -join ",")) -Encoding ASCII\n',
            'Set-Content -LiteralPath $PidFile -Value ("diff:" + ($fresh -join ",")) -Encoding ASCII\n'
            f"Set-Content -LiteralPath '{diff_file}' -Value ($fresh -join ',') -Encoding ASCII\n", 1)
        script = script.replace('Write-Output "word_pid=$windowPid"\n', 'Write-Output "word_pid=$windowPid"\nStart-Sleep -Seconds 600\n', 1)
        self.assertEqual(3, sum(marker in script for marker in (str(bystander), str(diff_file), "Start-Sleep -Seconds 600")))
        module.POWERSHELL_SCRIPT = script
        bystander_pid = None
        try:
            result = module.update_fields(docx, 150, module.find_powershell())
            bystander_pid = int(bystander_pid_file.read_text(encoding="ascii").strip())
            diff = [int(x) for x in diff_file.read_text(encoding="ascii").strip().split(",") if x.strip()]
            self.assertEqual("manual_required", result["status"], result)
            self.assertEqual("window", result["word_pid"]["source"], result)
            own = result["word_pid"]["pids"]
            self.assertEqual(own, result["word_pid"]["stopped"])
            self.assertIn(bystander_pid, diff, "гонка воспроизведена: разница списков захватила чужой Word")
            self.assertNotIn(bystander_pid, own)
            time.sleep(2)
            self.assertFalse(any(module.is_word_process(pid) for pid in own), "Word этого запуска остановлен")
            self.assertTrue(module.is_word_process(bystander_pid), "чужой Word не тронут")
        finally:
            if bystander_pid is None and bystander_pid_file.is_file():
                bystander_pid = int(bystander_pid_file.read_text(encoding="ascii").strip() or 0)
            if bystander_pid and module.is_word_process(bystander_pid):
                subprocess.run(["taskkill", "/PID", str(bystander_pid), "/F"], stdout=subprocess.DEVNULL,
                               stderr=subprocess.DEVNULL, check=False)


if __name__ == "__main__":
    unittest.main()
