"""WP5 6.33: verify_sources (--extract/--report/--queries), ai_detection_heuristic и
content_ownership_check — границы разделов DOCX, INSUFFICIENT_DATA и коды выхода,
cliche_allowlist, выборка вопросов.

Запуск: python -m unittest discover -s tests -t .   или   python tests/test_sources_analyzers_633.py
DOCX-фикстуры создаются python-docx во временных каталогах; Word-фикстура —
tests/fixtures/build/vkr-word-saved.docx (оглавление обновлено в Word).
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("PYTHONDONTWRITEBYTECODE", "1")
sys.dont_write_bytecode = True

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "vkr-mpgu"
SCRIPTS = SKILL / "scripts"
WORD_FIXTURE = ROOT / "tests" / "fixtures" / "build" / "vkr-word-saved.docx"

from docx import Document  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


def load_module(name: str):
    spec = importlib.util.spec_from_file_location(f"{name}_under_test_wp5", SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


ai = load_module("ai_detection_heuristic")
own = load_module("content_ownership_check")
vs = load_module("verify_sources")


def run_script(name: str, *args, env_extra=None):
    env = dict(os.environ)
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    env.update(env_extra or {})
    return subprocess.run(
        [sys.executable, str(SCRIPTS / name), *map(str, args)],
        capture_output=True, text=True, encoding="utf-8", env=env,
    )


def sha(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


# ---------------------------------------------------------------------------
# Построение DOCX
# ---------------------------------------------------------------------------

SENTENCES = [
    "Пространственное мышление школьников формируется при решении задач на построение сечений.",
    "Учитель отмечает, что наглядная модель сокращает время на объяснение новой темы.",
    "В пилотной группе двадцать учащихся выполнили по шесть заданий разной сложности.",
    "Тренажёр фиксирует ошибки и показывает подсказку после второй неверной попытки.",
    "Сравнение платформ проводилось по цене гарнитуры, доступности сборок и поддержке контроллеров.",
    "Анкета включала вопросы об удобстве интерфейса и понятности формулировок заданий.",
    "Результаты обработаны в электронной таблице и сведены в итоговую диаграмму.",
    "Методические рекомендации опираются на программу основной школы по геометрии.",
    "Для каждой модели многогранника подготовлено описание вершин, рёбер и граней.",
    "Время загрузки сцены на автономном шлеме не превышает четырёх секунд.",
]

ENTRIES = [
    "Абрамова, Т.Н. Цифровые тренажёры в школе [Текст] / Т.Н. Абрамова // Информатизация образования. – 2021. – № 3. – С. 15–22.",
    "Рапай, К. Культурный код: Как мы живем, что покупаем и почему [Текст] / К. Рапай. – М.: Альпина Бизнес Бук, 2008. – 167 с.",
    "Smith, J. Immersive learning environments / J. Smith // Computers & Education. – 2019. – Vol. 140. – P. 1–12.",
]
OATH = ("Выпускная квалификационная работа выполнена мной совершенно самостоятельно. Все использованные "
        "в работе материалы и концепции из опубликованной научной литературы и других источников имеют ссылки на них.")


def prose(tag: str, sentences: int = 4, offset: int = 0) -> str:
    return f"{tag}. " + " ".join(SENTENCES[(offset + i) % len(SENTENCES)] for i in range(sentences))


def ensure_style(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        return style


def _field_char(paragraph, kind: str) -> None:
    run = paragraph.add_run()
    element = OxmlElement("w:fldChar")
    element.set(qn("w:fldCharType"), kind)
    run._r.append(element)


def _instr(paragraph, text: str) -> None:
    run = paragraph.add_run()
    element = OxmlElement("w:instrText")
    element.set(qn("xml:space"), "preserve")
    element.text = text
    run._r.append(element)


def set_numbering(paragraph, num_id: int = 7) -> None:
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def add_toc_field(doc, lines, style_name=None) -> None:
    """Оглавление как после F9 в Word: поле TOC открыто в первой строке, закрыто отдельным абзацем."""
    style = ensure_style(doc, style_name) if style_name else None
    first = doc.add_paragraph(style=style)
    _field_char(first, "begin")
    _instr(first, ' TOC \\o "1-3" \\h \\z \\u ')
    _field_char(first, "separate")
    first.add_run(lines[0])
    for line in lines[1:]:
        doc.add_paragraph(line, style=style)
    last = doc.add_paragraph(style=style)
    _field_char(last, "end")


def add_toc_sdt(doc, lines) -> None:
    """sdt-оглавление (docPartGallery «Table of Contents») без табуляций и номеров страниц."""
    sdt = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    part = OxmlElement("w:docPartObj")
    gallery = OxmlElement("w:docPartGallery")
    gallery.set(qn("w:val"), "Table of Contents")
    part.append(gallery)
    properties.append(part)
    sdt.append(properties)
    content = OxmlElement("w:sdtContent")
    for line in lines:
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.set(qn("xml:space"), "preserve")
        text.text = line
        run.append(text)
        paragraph.append(run)
        content.append(paragraph)
    sdt.append(content)
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(sdt)
    else:
        body.append(sdt)


def add_last_sheet(doc, styled: bool = True) -> None:
    style = ensure_style(doc, "VKR Title") if styled else None
    doc.add_paragraph("Выпускная квалификационная работа", style=style)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = OATH
    doc.add_paragraph("")
    row = doc.add_table(rows=1, cols=3).rows[0]
    row.cells[0].text, row.cells[1].text, row.cells[2].text = "Иванов Иван Иванович", "_" * 19, "_" * 19
    doc.add_paragraph("(Ф.И.О.)        (подпись)        (дата)", style=style)


def build_docx(path: Path, items) -> Path:
    doc = Document()
    for item in items:
        kind = item[0]
        if kind == "p":
            doc.add_paragraph(item[1])
        elif kind == "h1":
            doc.add_paragraph(item[1], style="Heading 1")
        elif kind == "h2":
            doc.add_paragraph(item[1], style="Heading 2")
        elif kind == "title":
            doc.add_paragraph(item[1], style=ensure_style(doc, "VKR Title"))
        elif kind == "toc_field":
            add_toc_field(doc, item[1], item[2] if len(item) > 2 else None)
        elif kind == "toc_sdt":
            add_toc_sdt(doc, item[1])
        elif kind == "bib_text":
            for number, entry in enumerate(item[1], start=1):
                doc.add_paragraph(f"{number}. {entry}")
        elif kind == "bib_paren":
            for number, entry in enumerate(item[1], start=1):
                doc.add_paragraph(f"{number}) {entry}")
        elif kind == "bib_numpr":
            for entry in item[1]:
                set_numbering(doc.add_paragraph(entry, style="List Paragraph"), 7)
        elif kind == "bib_style":
            for number, entry in enumerate(item[1], start=1):
                doc.add_paragraph(f"{number}. {entry}", style=ensure_style(doc, "VKR Bibliography"))
        elif kind == "list":
            for entry in item[1]:
                set_numbering(doc.add_paragraph(entry, style="List Paragraph"), 8)
        elif kind == "code":
            for line in item[1]:
                doc.add_paragraph(line, style=ensure_style(doc, "VKR Listing"))
        elif kind == "caption":
            doc.add_paragraph(item[1], style=ensure_style(doc, "VKR Caption"))
        elif kind == "last_sheet":
            add_last_sheet(doc, styled=item[1] if len(item) > 1 else True)
        else:  # pragma: no cover - ошибка в тесте
            raise ValueError(kind)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def title_page():
    return [
        ("title", "Министерство просвещения Российской Федерации"),
        ("title", "Выпускная квалификационная работа"),
        ("title", "Москва – 2026 год"),
        ("p", "Аннотация"),
        ("p", prose("Аннотация работы", 3, 5)),
    ]


TOC_LINES = [
    "Введение\t3",
    "Глава I. Теоретические основы\t5",
    "Заключение\t30",
    "Список использованной литературы\t34",
    "Приложение 1. Анкета\t36",
]


def body(chapter_extra=None, conclusion_paragraphs=3):
    items = [
        ("h1", "Введение"),
        ("p", prose("Абзац введения один", 4, 0)),
        ("p", prose("Абзац введения два", 4, 3)),
        ("p", prose("Абзац введения три", 4, 6)),
        ("h1", "Глава I. Теоретические основы"),
        ("h2", "1.1. Пространственное мышление"),
        ("p", prose("Первая глава абзац один", 4, 1)),
        ("p", prose("Первая глава абзац два", 4, 4)),
    ]
    items.extend(chapter_extra or [])
    items.extend([
        ("p", prose("Первая глава абзац последний", 4, 7)),
        ("h1", "Глава II. Анализ платформ"),
        ("p", prose("Вторая глава абзац один", 4, 2)),
        ("p", prose("Вторая глава абзац два", 4, 5)),
        ("p", prose("Вторая глава абзац три", 4, 8)),
        ("h1", "Заключение"),
    ])
    for number in range(conclusion_paragraphs):
        items.append(("p", prose(f"Абзац заключения {number + 1}", 4, number)))
    return items


def full_document(bib_heading="Список использованной литературы", bib_kind="bib_text", after=None, toc=True):
    items = title_page()
    if toc:
        items.append(("p", "Содержание"))
        items.append(("toc_field", TOC_LINES, "toc 1"))
    items.extend(body())
    items.append(("h1", bib_heading))
    items.append((bib_kind, ENTRIES))
    items.extend(after if after is not None else [("last_sheet",)])
    return items


# ---------------------------------------------------------------------------
# verify_sources --extract
# ---------------------------------------------------------------------------


class ExtractBibliographyTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def extract(self, items, name="doc.docx"):
        return vs.extract_bibliography(build_docx(self.dir / name, items))

    def assert_three_entries(self, result):
        self.assertEqual([1, 2, 3], [record["id"] for record in result["records"]], result)
        self.assertTrue(result["records"][0]["raw_text"].startswith("Абрамова, Т.Н."))
        self.assertTrue(result["records"][2]["raw_text"].startswith("Smith, J."))
        for record in result["records"]:
            self.assertEqual("pending", record["status"])
            self.assertIn("hints", record)
            self.assertNotIn("raw", record["hints"])

    def test_text_numbers_updated_toc_and_last_sheet(self):
        result = self.extract(full_document())
        self.assertTrue(result["found"])
        self.assertEqual("Список использованной литературы", result["heading"])
        self.assert_three_entries(result)
        joined = json.dumps(result["records"], ensure_ascii=False)
        self.assertNotIn("Выпускная квалификационная работа", joined)
        self.assertNotIn("Ф.И.О.", joined)

    def test_plain_last_sheet_without_styles_ends_section(self):
        result = self.extract(full_document(after=[("last_sheet", False)]))
        self.assert_three_entries(result)
        self.assertEqual([], [w for w in result["warnings"] if w["code"] == "UNNUMBERED_PARAGRAPH"])

    def test_parenthesis_numbers(self):
        self.assert_three_entries(self.extract(full_document(bib_kind="bib_paren")))

    def test_word_autonumbering_numpr(self):
        self.assert_three_entries(self.extract(full_document(bib_kind="bib_numpr")))

    def test_list_number_style_numbering(self):
        # Нумерация приходит из стиля List Number (numPr в стиле), номеров в тексте нет.
        doc_path = self.dir / "list_number.docx"
        doc = Document()
        doc.add_paragraph("Библиографический список", style="Heading 1")
        for entry in ENTRIES:
            doc.add_paragraph(entry, style="List Number")
        doc.add_paragraph("Приложение 1", style="Heading 1")
        doc.add_paragraph("1. Нравится ли вам заниматься с тренажёром на уроке геометрии?")
        doc.save(str(doc_path))
        self.assert_three_entries(vs.extract_bibliography(doc_path))

    def test_vkr_bibliography_style(self):
        self.assert_three_entries(self.extract(full_document(bib_kind="bib_style")))

    def test_heading_variants(self):
        variants = [
            "Список использованной литературы", "Список использованных источников", "Список литературы",
            "Библиографический список", "Литература", "Список использованной литературы и источников",
            "СПИСОК ЛИТЕРАТУРЫ:", "литература:",
        ]
        for index, heading in enumerate(variants):
            with self.subTest(heading=heading):
                result = self.extract([("p", prose("Текст", 2)), ("h1", heading), ("bib_text", ENTRIES), ("last_sheet",)],
                                      name=f"variant{index}.docx")
                self.assert_three_entries(result)

    def test_plain_heading_without_style(self):
        result = self.extract(body() + [("p", "Литература"), ("bib_text", ENTRIES), ("last_sheet", False)])
        self.assert_three_entries(result)

    def test_appendices_plural_after_bibliography(self):
        after = [
            ("p", "ПРИЛОЖЕНИЯ"),
            ("p", "Приложение 1. Анкета"),
            ("bib_text", ["Нравится ли вам заниматься с тренажёром?", "Понятны ли формулировки заданий?"]),
            ("last_sheet",),
        ]
        result = self.extract(full_document(after=after))
        self.assert_three_entries(result)
        self.assertNotIn("Нравится", json.dumps(result, ensure_ascii=False))

    def test_toc_lines_are_not_section_heading(self):
        # Заголовок списка есть только в оглавлении: нумерованные задачи после него не источники.
        items = title_page() + [
            ("p", "Содержание"),
            ("toc_field", TOC_LINES, "toc 1"),
            ("h1", "Введение"),
            ("p", "Задачи исследования:"),
            ("bib_text", ["Проанализировать литературу по теме.", "Разработать прототип тренажёра."]),
        ]
        result = self.extract(items)
        self.assertFalse(result["found"])
        self.assertEqual([], result["records"])

    def test_toc_without_styles_field_only_and_sdt(self):
        items = [("p", "Содержание"), ("toc_field", TOC_LINES), ("toc_sdt", ["Список использованной литературы", "Литература"]),
                 ("h1", "Введение"), ("bib_text", ["Первая задача исследования.", "Вторая задача исследования."])]
        result = self.extract(items)
        self.assertFalse(result["found"], result)
        result = self.extract(items + [("h1", "Список литературы"), ("bib_text", ENTRIES)], name="with_bib.docx")
        self.assert_three_entries(result)

    def test_wrapped_entry_is_joined_and_stray_paragraph_is_warning(self):
        entries = [
            "1. " + ENTRIES[0],
            "2. Якобсон, Р. Лингвистика и поэтика [Электронный ресурс] / Р. Якобсон // Структурализм: за и против. – М., 1975.",
            "– URL: http://www.philology.ru/linguistics1/jakobson-75.htm (дата обращения: 16.02.2024).",
            "3. " + ENTRIES[2],
            "Примечание: список сокращён для примера.",
        ]
        items = [("h1", "Список использованных источников")] + [("p", text) for text in entries] + [("last_sheet",)]
        result = self.extract(items)
        ids = [record["id"] for record in result["records"]]
        self.assertEqual([1, 2, 3], ids)
        self.assertIn("jakobson-75.htm", result["records"][1]["raw_text"])
        self.assertEqual("http://www.philology.ru/linguistics1/jakobson-75.htm", result["records"][1]["hints"]["url"])
        codes = {warning["code"]: warning["text"] for warning in result["warnings"]}
        self.assertIn("CONTINUATION_JOINED", codes)
        self.assertIn("Примечание", codes.get("UNNUMBERED_PARAGRAPH", ""))

    def test_duplicate_numbers_get_unique_ids(self):
        result = self.extract([("h1", "Список литературы"), ("p", "1. " + ENTRIES[0]), ("p", "1. " + ENTRIES[1])])
        self.assertEqual([1, 2], [record["id"] for record in result["records"]])
        self.assertIn("NUMBER_DUPLICATE", [warning["code"] for warning in result["warnings"]])

    def test_group_headings_without_style_are_skipped(self):
        # Полужирный или прописной короткий абзац без года — подзаголовок группы, а не запись без номера.
        path = self.dir / "groups.docx"
        doc = Document()
        doc.add_paragraph("Список литературы", style="Heading 1")
        doc.add_paragraph().add_run("Научная литература").bold = True
        doc.add_paragraph("1. " + ENTRIES[0])
        doc.add_paragraph("ЭЛЕКТРОННЫЕ РЕСУРСЫ")
        doc.add_paragraph("2. " + ENTRIES[1])
        doc.add_paragraph().add_run("Иванов, И.И. Полужирная запись без номера. – М., 2020.").bold = True
        doc.save(str(path))
        result = vs.extract_bibliography(path)
        self.assertEqual([1, 2], [record["id"] for record in result["records"]])
        codes = [(warning["code"], warning["text"]) for warning in result["warnings"]]
        self.assertIn(("SUBHEADING_SKIPPED", "Научная литература"), codes)
        self.assertIn(("SUBHEADING_SKIPPED", "ЭЛЕКТРОННЫЕ РЕСУРСЫ"), codes)
        self.assertIn("UNNUMBERED_PARAGRAPH", [code for code, _ in codes])  # с годом — это запись, не подзаголовок

    def test_parse_bibliography_blocks_is_the_extract_parser(self):
        blocks = [{"type": "p", "text": "1. " + ENTRIES[0]}, {"type": "p", "text": "2. Якобсон, Р. Лингвистика и поэтика /"},
                  {"type": "p", "text": "Р. Якобсон. – М., 1975."}, {"type": "p", "text": "3. " + ENTRIES[2]}]
        parsed = vs.parse_bibliography_blocks(blocks)
        self.assertEqual([1, 2, 3], [record["id"] for record in parsed["records"]])
        self.assertEqual("Якобсон, Р. Лингвистика и поэтика / Р. Якобсон. – М., 1975.", parsed["records"][1]["raw_text"])
        self.assertEqual(["CONTINUATION_JOINED"], [warning["code"] for warning in parsed["warnings"]])
        self.assertEqual([1, 2, 3], [entry["number"] for entry in parsed["entries"]])

    def test_word_saved_fixture(self):
        result = vs.extract_bibliography(WORD_FIXTURE)
        self.assertEqual(list(range(1, 10)), [record["id"] for record in result["records"]])
        self.assertTrue(result["records"][0]["raw_text"].startswith("Абрамова, Т.Н."))
        joined = json.dumps(result["records"], ensure_ascii=False)
        for foreign in ("Выпускная квалификационная работа", "Исходный код", "Приложение", "Ф.И.О."):
            self.assertNotIn(foreign, joined)
        # Совместимость API 6.32
        self.assertEqual(9, len(vs.extract_sources_from_docx(WORD_FIXTURE)))


class SourceHintsTest(unittest.TestCase):
    def test_authors_with_hyphen_and_apostrophe(self):
        self.assertEqual("Римский-Корсаков", vs.parse_source_hints("Римский-Корсаков, Н.А. Мемуары [Текст]. – СПб., 1909.")["author_last"])
        hints = vs.parse_source_hints("O'Brien, J. Learning in virtual reality / J. O'Brien. – New York: ACM, 2019. – 200 p.")
        self.assertEqual(("O'Brien", "J."), (hints["author_last"], hints["author_initials"]))
        self.assertEqual("Learning in virtual reality", hints["title"])
        self.assertEqual("en", hints["language"])

    def test_title_is_clean_for_exact_phrase(self):
        cases = {
            "Якобсон, Р. Лингвистика и поэтика [Электронный ресурс] / Р. Якобсон // Структурализм. – М., 1975.": "Лингвистика и поэтика",
            "Сидоров, С.С. Методика обучения стереометрии: дис. … канд. пед. наук: 13.00.02 / С.С. Сидоров. – М., 2015. – 200 с.": "Методика обучения стереометрии",
            "КиберЛенинка: научная электронная библиотека [Электронный ресурс]. – URL: https://cyberleninka.ru (дата обращения: 16.02.2024).": "КиберЛенинка",
            "Smith, J. Title of the book: subtitle / J. Smith. – London: Routledge, 2020. – 250 p.": "Title of the book: subtitle",
            "Иванов, И.И. Цифровая школа : электронный учебник [Электронный ресурс] / И.И. Иванов. – URL: https://example.org": "Цифровая школа",
        }
        for raw, title in cases.items():
            with self.subTest(raw=raw[:30]):
                self.assertEqual(title, vs.parse_source_hints(raw)["title"])
        self.assertEqual("dissertation", vs.parse_source_hints(list(cases)[1])["type"])

    def test_normative_and_queries_are_separate(self):
        law = vs.parse_source_hints("Об образовании в Российской Федерации [Текст]: федер. закон от 29.12.2012 № 273-ФЗ. – Доступ из справ.-правовой системы.")
        self.assertEqual(("Об образовании в Российской Федерации", "normative"), (law["title"], law["type"]))
        hints = vs.source_hints({
            "type": "article", "title": "Immersive learning [Electronic resource]", "year": 2019,
            "authors": [{"last": "Smith", "initials": "J."}], "doi": "10.1016/j.compedu.2019.103603", "isbn": "978-5-4461-1234-5",
        })
        queries = [item["query"] for item in vs.build_search_queries(hints)]
        self.assertIn('"Immersive learning"', queries)
        self.assertIn('Smith J. "Immersive learning"', queries)
        self.assertIn('"Immersive learning" 2019', queries)
        self.assertIn("https://doi.org/10.1016/j.compedu.2019.103603", queries)
        self.assertIn("ISBN 978-5-4461-1234-5", queries)
        self.assertFalse(any("cyberleninka" in query or "elibrary" in query for query in queries), queries)
        self.assertFalse(any("[" in query for query in queries), queries)

    def test_queries_use_subtitle_and_author_title(self):
        # A10: общее заглавие «Информатика» без подзаголовка не находит издание.
        hints = vs.source_hints({"type": "book", "title": "Информатика", "subtitle": "учебник для 7 класса",
                                 "authors": ["Босова Л.Л.", "Босова А.Ю."], "year": 2017})
        self.assertEqual("учебник для 7 класса", hints["subtitle"])
        queries = [item["query"] for item in vs.build_search_queries(hints)]
        self.assertIn('"Информатика: учебник для 7 класса"', queries)
        self.assertIn('Босова Л.Л. "Информатика: учебник для 7 класса"', queries)
        self.assertIn("Босова Информатика учебник для 7 класса", queries)
        self.assertIn('"Информатика: учебник для 7 класса" 2017', queries)
        self.assertNotIn('"Информатика" 2017', queries)
        self.assertNotIn('"Информатика"', queries)
        plain = [item["query"] for item in vs.build_search_queries(vs.source_hints(
            {"type": "book", "title": "Мышление и речь", "authors": ["Выготский Л.С."], "year": 1934}))]
        self.assertEqual(['"Мышление и речь"', 'Выготский Л.С. "Мышление и речь"', "Выготский Мышление и речь",
                          '"Мышление и речь" 1934', '"Мышление и речь" site:elibrary.ru'], plain)


class VerifySourcesCliTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def write_json(self, name, value, bom=False):
        path = self.dir / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(("\ufeff" if bom else "") + json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def test_extract_force_and_registry_protection(self):
        docx_path = build_docx(self.dir / "vkr.docx", full_document())
        out = self.dir / "audit" / "extracted.json"
        proc = run_script("verify_sources.py", "--extract", docx_path, "-o", out)
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        self.assertEqual([1, 2, 3], [item["id"] for item in json.loads(out.read_text(encoding="utf-8"))])
        before = sha(out)
        proc = run_script("verify_sources.py", "--extract", docx_path, "-o", out)
        self.assertEqual(2, proc.returncode)
        self.assertIn("--force", proc.stderr)
        self.assertEqual(before, sha(out))
        out.write_text("[]", encoding="utf-8")
        proc = run_script("verify_sources.py", "--extract", docx_path, "-o", out, "--force")
        self.assertEqual(0, proc.returncode, proc.stderr)
        self.assertEqual(3, len(json.loads(out.read_text(encoding="utf-8"))))

        registry = self.write_json("project/sources.json", [
            {"id": "abramova2021", "type": "article", "title": "Цифровые тренажёры в школе", "journal": "Информатизация образования",
             "status": "confirmed", "verification": {"method": "catalog", "checked_at": "2026-09-01T10:00:00+03:00", "note": "eLibrary"}},
        ])
        before = sha(registry)
        for extra in ([], ["--force"]):
            proc = run_script("verify_sources.py", "--extract", docx_path, "-o", registry, *extra, "--json")
            self.assertEqual(2, proc.returncode, proc.stdout)
            payload = json.loads(proc.stdout)
            self.assertEqual("REGISTRY_PROTECTED", payload["error"]["code"])
            self.assertIn("sources-extracted.json", payload["error"]["message"])
            self.assertEqual(before, sha(registry))

        raw_only = self.write_json("raw/sources.json", [{"id": 1, "raw_text": "Старая запись", "status": "pending"}])
        proc = run_script("verify_sources.py", "--extract", docx_path, "-o", raw_only, "--force")
        self.assertEqual(0, proc.returncode, proc.stderr)

    def test_extract_default_output_is_project_audit(self):
        project = self.dir / "project"
        (project / "final").mkdir(parents=True)
        (project / "vkr-project.json").write_text("{}", encoding="utf-8")
        docx_path = build_docx(project / "final" / "vkr.docx", full_document())
        proc = run_script("verify_sources.py", "--extract", docx_path, "--json")
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        payload = json.loads(proc.stdout)
        self.assertEqual(3, payload["records"])
        self.assertTrue((project / "audit" / "sources-extracted.json").is_file())

    def test_extract_zero_records_exit_1(self):
        docx_path = build_docx(self.dir / "nobib.docx", title_page() + body())
        out = self.dir / "out.json"
        proc = run_script("verify_sources.py", "--extract", docx_path, "-o", out)
        self.assertEqual(1, proc.returncode, proc.stdout + proc.stderr)
        self.assertIn("0 записей", proc.stderr)
        self.assertFalse(out.exists())
        unnumbered = build_docx(self.dir / "unnumbered.docx", [("h1", "Литература")] + [("p", entry) for entry in ENTRIES])
        proc = run_script("verify_sources.py", "--extract", unnumbered, "-o", out, "--json")
        self.assertEqual(1, proc.returncode)
        payload = json.loads(proc.stdout)
        self.assertEqual("empty", payload["status"])
        self.assertIn("без номера", payload["message"])
        self.assertFalse(out.exists())

    def test_report_lists_problem_sources_with_notes(self):
        registry = self.write_json("sources.json", [
            {"id": 1, "raw_text": "Сидоров, С. Несуществующая книга. – М., 2020.", "status": "suspicious",
             "verification": {"method": "web_search", "checked_at": "2026-09-10", "note": "год не совпал"}},
            {"id": 2, "title": "Выдумка", "status": "rejected", "verification": {"note": "не найдено ни в одном каталоге"}},
            {"id": 3, "raw_text": "Петров, П. Ждёт проверки. – М., 2019.", "status": "pending"},
            {"id": 4, "title": "Мышление и речь", "authors": [{"last": "Выготский", "initials": "Л.С."}], "status": "Confirmed"},
            {"id": 5, "title": "Педагогика", "status": "verified",
             "verification": {"method": "catalog", "checked_at": "2026-09-01T10:00:00.1234567+0300", "url": "https://search.rsl.ru/"}},
            {"id": 6, "raw_text": "Фейк, Ф. Старый статус.", "status": "unconfirmed", "notes": "не найдено"},
        ], bom=True)
        proc = run_script("verify_sources.py", "--report", registry)
        self.assertEqual(1, proc.returncode, proc.stderr)
        for fragment in ("[1] Сидоров", "год не совпал", "[2] Выдумка", "не найдено ни в одном каталоге",
                         "[3] Петров", "[4] Выготский", "Нет объекта verification", "[6] Фейк", "unconfirmed"):
            self.assertIn(fragment, proc.stdout)
        report_path = self.dir / "reports" / "report.json"
        proc = run_script("verify_sources.py", "--report", registry, "--json", "-o", report_path)
        payload = json.loads(proc.stdout)
        self.assertEqual(json.loads(report_path.read_text(encoding="utf-8")), payload)
        self.assertEqual(
            {"confirmed": 2, "confirmed_without_verification": 1, "confirmed_stale": 0, "suspicious": 1, "pending": 1,
             "rejected": 1, "unknown_status": 1},
            payload["counts"],
        )
        self.assertEqual(["год не совпал"], [item["note"] for item in payload["suspicious"]])
        self.assertEqual([4], [item["id"] for item in payload["confirmed_without_verification"]])

        good = self.write_json("good.json", [{"id": "a", "title": "Педагогика", "status": "verified",
                                              "verification": {"method": "doi", "checked_at": "2026-09-01T10:00:00Z", "note": "doi.org"}}])
        self.assertEqual(0, run_script("verify_sources.py", "--report", good).returncode)
        empty = self.write_json("empty.json", [])
        self.assertEqual(1, run_script("verify_sources.py", "--report", empty).returncode)
        broken = self.dir / "broken.json"
        broken.write_text("[{", encoding="utf-8")
        self.assertEqual(2, run_script("verify_sources.py", "--report", broken).returncode)
        self.assertEqual(2, run_script("verify_sources.py", "--report", self.dir / "missing.json").returncode)

    def test_queries_on_registry_without_id(self):
        registry = self.write_json("noid.json", [
            {"type": "electronic", "authors": ["Якобсон Р."], "title": "Лингвистика и поэтика [Электронный ресурс]",
             "year": 1975, "url": "http://www.philology.ru/linguistics1/jakobson-75.htm"},
            {"raw_text": "Иванов, И.И. Цифровая школа : электронный учебник [Электронный ресурс] / И.И. Иванов. – М., 2020. – ISBN 978-5-4461-1234-5."},
            {"title": "Уже проверено", "status": "confirmed",
             "verification": {"method": "catalog", "checked_at": "2026-09-01", "note": "РГБ"}},
        ])
        proc = run_script("verify_sources.py", "--queries", registry)
        self.assertEqual(0, proc.returncode, proc.stderr)
        self.assertNotIn("Traceback", proc.stderr)
        self.assertIn("[#1]", proc.stdout)
        self.assertIn('"Лингвистика и поэтика"', proc.stdout)
        self.assertIn('Якобсон Р. "Лингвистика и поэтика"', proc.stdout)
        self.assertIn('"Цифровая школа"', proc.stdout)
        self.assertNotIn("Уже проверено", proc.stdout)
        payload = json.loads(run_script("verify_sources.py", "--queries", registry, "--json").stdout)
        self.assertEqual([1, 2], [item["index"] for item in payload["items"]])
        second = [query["query"] for query in payload["items"][1]["queries"]]
        self.assertIn("ISBN 978-5-4461-1234-5", second)
        self.assertIn('"Цифровая школа" 2020', second)
        for query in second:
            self.assertNotIn("Электронный ресурс", query)
            self.assertNotIn(" / ", query)
            self.assertNotIn("электронный учебник", query)

    def test_mark_writes_confirmed_with_verification(self):
        registry = self.write_json("sources.json", [
            {"id": "vygotsky1934", "type": "book", "title": "Мышление и речь", "status": "pending"},
            {"id": 2, "raw_text": "Фейк, Ф. Выдумка. – М., 2021.", "status": "pending"},
        ], bom=True)
        proc = run_script("verify_sources.py", "--mark", registry, "--id", "VYGOTSKY1934", "--status", "verified",
                          "--method", "catalog", "--url", "https://search.rsl.ru/", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        data = json.loads(registry.read_text(encoding="utf-8"))
        self.assertEqual("confirmed", data[0]["status"])
        self.assertEqual([], vs.verification_missing(data[0]))
        self.assertEqual("catalog", data[0]["verification"]["method"])
        proc = run_script("verify_sources.py", "--mark", registry, "--id", "2", "--status", "rejected")
        self.assertEqual(2, proc.returncode)
        self.assertIn("--note", proc.stderr)
        self.assertEqual(2, run_script("verify_sources.py", "--mark", registry, "--id", "vygotsky1934",
                                       "--status", "confirmed", "--method", "catalog").returncode)
        self.assertEqual(2, run_script("verify_sources.py", "--mark", registry, "--id", "404", "--status", "pending").returncode)
        proc = run_script("verify_sources.py", "--mark", registry, "--index", "2", "--status", "rejected", "--note", "не найдено")
        self.assertEqual(0, proc.returncode, proc.stderr)
        report = json.loads(run_script("verify_sources.py", "--report", registry, "--json").stdout)
        self.assertEqual(1, report["counts"]["confirmed"])
        self.assertEqual(0, report["counts"]["confirmed_without_verification"])
        self.assertEqual(["не найдено"], [item["note"] for item in report["rejected"]])
        run_script("verify_sources.py", "--mark", registry, "--id", "vygotsky1934", "--status", "pending")
        self.assertNotIn("verification", json.loads(registry.read_text(encoding="utf-8"))[0])

    def test_edited_fields_after_mark_make_verification_stale(self):
        # A9: --mark хранит хеш реквизитов; правка confirmed-записи видна в --report и --queries.
        registry = self.write_json("sources.json", [
            {"id": "vygotsky1934", "type": "book", "authors": ["Выготский Л.С."], "title": "Мышление и речь",
             "place": "М.", "publisher": "Соцэкгиз", "year": 1934, "status": "pending", "notes": "из черновика"},
        ])
        proc = run_script("verify_sources.py", "--mark", registry, "--id", "vygotsky1934", "--status", "confirmed",
                          "--method", "catalog", "--url", "https://search.rsl.ru/", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        data = json.loads(registry.read_text(encoding="utf-8"))
        self.assertEqual(vs.bibliographic_fields_sha256(data[0]), data[0]["verification"]["fields_sha256"])
        self.assertFalse(vs.verification_stale(data[0]))
        report = json.loads(run_script("verify_sources.py", "--report", registry, "--json").stdout)
        self.assertEqual((0, 0), (report["exit_code"], report["counts"]["confirmed_stale"]))
        # служебные поля и статус хеш не меняют
        data[0]["notes"] = "сверено с каталогом"
        self.assertFalse(vs.verification_stale(data[0]))
        data[0]["year"] = 1956
        registry.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        self.assertTrue(vs.verification_stale(data[0]))
        proc = run_script("verify_sources.py", "--report", registry, "--json")
        self.assertEqual(1, proc.returncode, proc.stdout)
        report = json.loads(proc.stdout)
        self.assertEqual(1, report["counts"]["confirmed_stale"])
        self.assertEqual(["vygotsky1934"], [item["id"] for item in report["confirmed_stale"]])
        text = run_script("verify_sources.py", "--report", registry).stdout
        self.assertIn("УСТАРЕВШЕЙ ПРОВЕРКОЙ", text)
        queries = json.loads(run_script("verify_sources.py", "--queries", registry, "--json").stdout)
        self.assertEqual(["vygotsky1934"], [item["id"] for item in queries["items"]])
        self.assertIn("реквизиты изменены", queries["items"][0]["reason"])
        # старые записи без fields_sha256 устаревшими не считаются
        legacy = {"id": 1, "title": "Педагогика", "status": "confirmed",
                  "verification": {"method": "catalog", "checked_at": "2026-09-01T10:00:00Z", "note": "РГБ"}}
        self.assertFalse(vs.verification_stale(legacy))

    def test_outputs_inside_skill_dir_are_refused(self):
        docx_path = build_docx(self.dir / "vkr.docx", full_document())
        target_dir = SKILL / "wp5-test-output-must-not-exist"
        target = target_dir / "out.json"
        self.assertFalse(target_dir.exists())
        try:
            proc = run_script("verify_sources.py", "--extract", docx_path, "-o", target, "--json")
            self.assertEqual(2, proc.returncode, proc.stdout)
            self.assertEqual("OUTPUT_INSIDE_SKILL", json.loads(proc.stdout)["error"]["code"])
            for name in ("ai_detection_heuristic.py", "content_ownership_check.py"):
                proc = run_script(name, docx_path, "--json", "-o", target)
                self.assertEqual(2, proc.returncode, name + proc.stderr)
                self.assertIn("каталога скилла", proc.stderr)
            self.assertFalse(target_dir.exists())
        finally:
            if target_dir.exists():
                shutil.rmtree(target_dir)

    def test_scripts_without_python_docx(self):
        blocker = self.dir / "blocker"
        (blocker / "docx").mkdir(parents=True)
        (blocker / "docx" / "__init__.py").write_text("raise ImportError('python-docx заблокирован тестом')\n", encoding="utf-8")
        env = {"PYTHONPATH": str(blocker)}
        docx_path = build_docx(self.dir / "vkr.docx", full_document())
        for name in ("verify_sources.py", "ai_detection_heuristic.py", "content_ownership_check.py"):
            proc = run_script(name, "--help", env_extra=env)
            self.assertEqual(0, proc.returncode, name + proc.stderr)
        proc = run_script("verify_sources.py", "--extract", docx_path, "-o", self.dir / "x.json", env_extra=env)
        self.assertEqual(2, proc.returncode)
        self.assertIn("python-docx", proc.stderr)
        self.assertNotIn("Traceback", proc.stderr)
        for name in ("ai_detection_heuristic.py", "content_ownership_check.py"):
            proc = run_script(name, docx_path, "--json", env_extra=env)
            self.assertEqual(2, proc.returncode, name)
            self.assertEqual("DEPENDENCY_MISSING", json.loads(proc.stdout)["error"]["code"])
        registry = self.write_json("sources.json", [{"id": 1, "title": "Книга", "status": "pending"}])
        self.assertEqual(1, run_script("verify_sources.py", "--report", registry, env_extra=env).returncode)


# ---------------------------------------------------------------------------
# Анализаторы: границы разделов
# ---------------------------------------------------------------------------


class AnalyzerSectionsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def doc(self, items, name="doc.docx"):
        return build_docx(self.dir / name, items)

    def test_literature_heading_is_not_swallowed_by_conclusion(self):
        path = self.doc(full_document(bib_heading="Литература"))
        full_text, paragraphs = ai.extract_section(Document(str(path)), "conclusion")
        self.assertEqual(3, len(paragraphs))
        self.assertNotIn("Рапай", full_text)
        body_paragraphs = ai.split_paragraphs_from_doc(Document(str(path)))
        joined = "\n".join(body_paragraphs)
        for foreign in ("Рапай", "Министерство просвещения", "Аннотация работы", "Глава I. Теоретические основы\t5"):
            self.assertNotIn(foreign, joined)
        self.assertEqual(0, ai.check_pronouns(joined)["total"])
        owned = own.split_paragraphs(Document(str(path)))
        self.assertFalse(any("Рапай" in text or "Абрамова" in text for text in owned))
        self.assertEqual(3, len(own.extract_section(Document(str(path)), "conclusion")))
        # Заголовок без стиля (обычный абзац «Литература») тоже закрывает заключение.
        plain = self.doc(body() + [("p", "Литература"), ("bib_text", ENTRIES), ("last_sheet", False)], "plain.docx")
        conclusion, paragraphs = ai.extract_section(Document(str(plain)), "conclusion")
        self.assertEqual(3, len(paragraphs))
        self.assertNotIn("Рапай", conclusion)
        self.assertEqual(3, len(own.extract_section(Document(str(plain)), "conclusion")))

    def test_paragraphs_starting_with_service_words_do_not_break_chapter(self):
        extra = [
            ("p", "Список требований к тренажёру сформирован по итогам интервью с учителями геометрии."),
            ("p", "Приложение разработано на Unity 6 и запускается на автономном шлеме Pico 4 без компьютера."),
            ("p", "Библиографический анализ показал, что тема тренажёров по стереометрии изучена слабо."),
            ("list", ["список заданий с автоматической проверкой;", "Приложение для учителя с журналом результатов."]),
        ]
        path = self.doc(body(chapter_extra=extra))
        document = Document(str(path))
        full_text, _ = ai.extract_section(document, "chapter1")
        self.assertIn("Первая глава абзац последний", full_text)
        self.assertIn("Приложение разработано на Unity 6", full_text)
        self.assertIn("список заданий", full_text)
        self.assertIn("Первая глава абзац последний", "\n".join(own.extract_section(document, "chapter1")))
        self.assertIn("Первая глава абзац последний", "\n".join(ai.split_paragraphs_from_doc(document)))
        structure = ai.read_structure(document)
        self.assertEqual(["introduction", "chapter1", "chapter2", "conclusion"], structure.section_names())

    def test_chapter_two_is_not_chapter_one_and_long_titles(self):
        long_one = "Глава I. " + "Теоретические основы применения технологий виртуальной реальности " * 4
        long_two = "Глава II. " + "Анализ целевой аудитории, конкурентов и требований к образовательному продукту " * 3
        self.assertGreater(len(long_one), 180)
        items = [
            ("h1", "Введение"), ("p", prose("Введение", 4)),
            ("h1", long_one.strip()), ("p", prose("Текст первой главы", 4)),
            ("h1", long_two.strip()), ("p", prose("Текст второй главы", 4)),
            ("h1", "Глава III"), ("p", prose("Текст третьей главы", 4)),
            ("h1", "Заключение"), ("p", prose("Заключение", 4)),
        ]
        document = Document(str(self.doc(items)))
        first, _ = ai.extract_section(document, "chapter1")
        second, _ = ai.extract_section(document, "chapter2")
        third, _ = ai.extract_section(document, "chapter3")
        self.assertIn("Текст первой главы", first)
        self.assertNotIn("Текст второй главы", first)
        self.assertIn("Текст второй главы", second)
        self.assertNotIn("Текст третьей главы", second)
        self.assertIn("Текст третьей главы", third)
        self.assertIn("Текст второй главы", own.extract_section(document, "chapter2")[0])

    def test_heading_fallback_without_styles_and_numbered_chapters(self):
        plain = [
            ("p", "Введение"), ("p", prose("Введение без стилей", 4)),
            ("p", "Глава 1. Теория"), ("p", prose("Теория без стилей", 4)),
            ("p", "Глава 2 посвящена анализу, а глава 3 — разработке и пилотированию тренажёра по стереометрии."),
            ("p", "Заключение"), ("p", prose("Заключение без стилей", 4)),
            ("p", "Список литературы"), ("bib_text", ENTRIES),
        ]
        structure = ai.read_structure(self.doc(plain, "plain.docx"))
        self.assertEqual(["introduction", "chapter1", "conclusion"], structure.section_names())
        self.assertIn("Глава 2 посвящена анализу", "\n".join(structure.section_texts("chapter1")))
        numbered = [
            ("h1", "Введение"), ("p", prose("Введение", 4)),
            ("h1", "1. Теоретические основы"), ("h2", "1.1. Понятия"), ("p", prose("Первая", 4)),
            ("h1", "2. Анализ"), ("p", prose("Вторая", 4)),
            ("h1", "Заключение"), ("p", prose("Итог", 4)),
        ]
        structure = ai.read_structure(self.doc(numbered, "numbered.docx"))
        self.assertEqual(["introduction", "chapter1", "chapter2", "conclusion"], structure.section_names())
        self.assertNotIn("Первая", "\n".join(structure.section_texts("introduction")))

    def test_toc_code_captions_and_title_are_excluded(self):
        extra = [
            ("code", ["public bool Check(int[] answers) {", "    return answers.Length > 0; // стоит отметить", "}"]),
            ("caption", "Листинг 1 — Проверка ответа ученика"),
            ("p", "Рисунок 2 — Экран результатов пилотирования"),
            ("p", "Источник: составлено автором"),
        ]
        items = title_page() + [("p", "Содержание"), ("toc_field", TOC_LINES, "toc 1"),
                                ("toc_sdt", ["Введение", "Заключение"])] + body(chapter_extra=extra)
        document = Document(str(self.doc(items)))
        intro, _ = ai.extract_section(document, "introduction")
        self.assertTrue(intro.startswith("Абзац введения один"), intro[:80])
        chapter, _ = ai.extract_section(document, "chapter1")
        for foreign in ("public bool", "Листинг 1", "Рисунок 2", "Источник:", "1.1. Пространственное мышление"):
            self.assertNotIn(foreign, chapter)
        joined = "\n".join(ai.split_paragraphs_from_doc(document))
        self.assertNotIn("\t", joined)
        self.assertNotIn("Москва – 2026", joined)

    def test_word_saved_fixture_sections(self):
        structure = ai.read_structure(WORD_FIXTURE)
        self.assertEqual(["introduction", "chapter1", "chapter2", "chapter3", "conclusion"], structure.section_names())
        intro = "\n".join(structure.section_texts("introduction"))
        self.assertTrue(intro.startswith("Актуальность исследования"), intro[:60])
        joined = "\n".join(text for _, text in structure.body_items())
        for foreign in ("Абрамова, Т.Н.", "def check", "public bool", "Анкета для учащихся", "Лебедева Анна", "Выводы по главе"):
            self.assertNotIn(foreign, joined)
        questions = own.split_paragraphs(Document(str(WORD_FIXTURE)))
        self.assertFalse(any(text.startswith(("1.", "Абрамова")) for text in questions))


class AnalyzerCliTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_empty_document_is_insufficient_data_with_exit_1(self):
        path = self.dir / "empty.docx"
        Document().save(str(path))
        proc = run_script("ai_detection_heuristic.py", path, "--json")
        self.assertEqual(1, proc.returncode, proc.stdout + proc.stderr)
        self.assertEqual("INSUFFICIENT_DATA", json.loads(proc.stdout)["risk"]["risk_level"])
        proc = run_script("ai_detection_heuristic.py", path)
        self.assertEqual(1, proc.returncode)
        self.assertIn("INSUFFICIENT_DATA", proc.stderr)
        proc = run_script("ai_detection_heuristic.py", path, "--per-section", "--json")
        self.assertEqual(1, proc.returncode)
        payload = json.loads(proc.stdout)
        self.assertEqual("insufficient_data", payload["status"])
        self.assertEqual(["introduction", "chapter1", "conclusion"], payload["missing_sections"])
        proc = run_script("content_ownership_check.py", path, "--json")
        self.assertEqual(1, proc.returncode)
        self.assertEqual("INSUFFICIENT_DATA", json.loads(proc.stdout)["risk_level"])
        short = build_docx(self.dir / "short.docx", [("h1", "Введение"), ("p", "Короткий текст введения.")])
        proc = run_script("ai_detection_heuristic.py", short, "--section", "introduction", "--json")
        self.assertEqual(1, proc.returncode)
        self.assertEqual("INSUFFICIENT_DATA", json.loads(proc.stdout)["risk"]["risk_level"])

    def test_per_section_reports_missing_and_insufficient(self):
        def section(tag, count=3):
            return [("p", prose(f"{tag} {i}", 4, i)) for i in range(count)]

        items = ([("h1", "Введение")] + section("Введение") + [("h1", "Глава I. Теория")] + section("Теория")
                 + [("h1", "Глава III. Разработка"), ("p", "Короткий текст главы.")]
                 + [("h1", "Заключение")] + section("Итог"))
        path = build_docx(self.dir / "gaps.docx", items)
        proc = run_script("ai_detection_heuristic.py", path, "--per-section", "--json")
        self.assertEqual(1, proc.returncode, proc.stdout + proc.stderr)
        payload = json.loads(proc.stdout)
        self.assertEqual("incomplete", payload["status"])
        self.assertEqual(["chapter2"], payload["missing_sections"])
        self.assertEqual(["chapter3"], payload["insufficient_sections"])
        self.assertEqual("INSUFFICIENT_DATA", payload["sections"]["chapter3"]["risk"]["risk_level"])
        self.assertNotEqual("INSUFFICIENT_DATA", payload["sections"]["chapter1"]["risk"]["risk_level"])
        text = run_script("ai_detection_heuristic.py", path, "--per-section")
        self.assertEqual(1, text.returncode)
        self.assertIn("НЕ НАЙДЕНЫ РАЗДЕЛЫ: chapter2", text.stdout)
        self.assertIn("chapter3", text.stdout)

        complete = ([("h1", "Введение")] + section("Введение") + [("h1", "Глава I. Теория")] + section("Теория")
                    + [("h1", "Глава II. Анализ")] + section("Анализ") + [("h1", "Заключение")] + section("Итог"))
        path = build_docx(self.dir / "complete.docx", complete)
        output = self.dir / "out" / "per-section.json"
        proc = run_script("ai_detection_heuristic.py", path, "--per-section", "-o", output)
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        payload = json.loads(output.read_text(encoding="utf-8"))
        self.assertEqual("ok", payload["status"])
        self.assertEqual(["introduction", "chapter1", "chapter2", "conclusion"], list(payload["sections"]))

        proc = run_script("ai_detection_heuristic.py", path, "--section", "chapter3", "--json")
        self.assertEqual(1, proc.returncode)
        self.assertEqual("section_not_found", json.loads(proc.stdout)["status"])
        self.assertEqual(1, run_script("content_ownership_check.py", path, "--section", "chapter3").returncode)
        self.assertEqual(2, run_script("ai_detection_heuristic.py", path, "--section", "chapterX").returncode)
        self.assertEqual(2, run_script("content_ownership_check.py", self.dir / "missing.docx").returncode)

    def test_cliche_allowlist_forms_and_project_boundary(self):
        variants = {
            "indented": 'cliche_allowlist:\n  - "стоит отметить"\n  - \'в современном мире\'\nuser_seed: "x"\n',
            "flush": 'cliche_allowlist:\n- "стоит отметить"\n- в современном мире  # комментарий\n\nother: 1\n',
            "inline": 'cliche_allowlist: ["стоит отметить", \'в современном мире\']\n',
            "inline_multiline": 'cliche_allowlist: [\n  "стоит отметить",\n  "в современном мире"\n]\n',
        }
        for name, text in variants.items():
            with self.subTest(form=name):
                self.assertEqual(["стоит отметить", "в современном мире"], ai.parse_cliche_allowlist("# State\n" + text))
        self.assertEqual([], ai.parse_cliche_allowlist('cliche_allowlist: []\n# cliche_allowlist:\n#   - "стоит отметить"\n'))

        cliche_text = [
            ("h1", "Введение"),
            ("p", "Стоит отметить, что в современном мире тренажёры применяются всё шире. " + prose("Введение", 6)),
            ("p", "Стоит отметить значение наглядности для учащихся основной школы. " + prose("Второй", 6, 2)),
            ("h1", "Глава I. Теория"),
            ("p", "В современном мире цифровая среда меняет урок. Стоит отметить роль учителя. " + prose("Глава", 8, 4)),
            ("p", prose("Глава продолжение", 8, 1)),
            ("h1", "Заключение"),
            ("p", "Стоит отметить, что задачи решены. " + prose("Заключение", 6, 3)),
        ]
        project = self.dir / "project"
        docx_path = build_docx(project / "final" / "vkr.docx", cliche_text)

        def strong_hits():
            proc = run_script("ai_detection_heuristic.py", docx_path, "--json")
            self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
            payload = json.loads(proc.stdout)
            return {phrase for phrase, _ in payload["ngrams"]["top_strong"]}, payload["meta"]["cliche_allowlist"]

        hits, info = strong_hits()
        self.assertIn("стоит отметить", hits)
        self.assertIsNone(info["state_file"])
        # vkr-state.md выше корня проекта (каталог с vkr-project.json) не читается
        (project / "vkr-project.json").write_text("{}", encoding="utf-8")
        (self.dir / "vkr-state.md").write_text(variants["flush"], encoding="utf-8")
        hits, info = strong_hits()
        self.assertIn("стоит отметить", hits)
        for name in ("flush", "indented"):
            with self.subTest(state=name):
                (project / "vkr-state.md").write_text("---\n" + variants[name] + "---\n", encoding="utf-8")
                hits, info = strong_hits()
                self.assertNotIn("стоит отметить", hits)
                self.assertNotIn("в современном мире", hits)
                self.assertEqual(["стоит отметить", "в современном мире"], info["phrases"])
                proc = run_script("content_ownership_check.py", docx_path, "--all", "--json")
                self.assertEqual(0, proc.returncode, proc.stderr)
                self.assertEqual(info["phrases"], json.loads(proc.stdout)["cliche_allowlist"]["phrases"])
        marker = "Следует отметить, что выборка небольшая."
        self.assertIn("author-position", [p["category"] for p in own.detect_applicable_patterns(marker)])
        self.assertNotIn("author-position",
                         [p["category"] for p in own.detect_applicable_patterns(marker, allowlist=["следует отметить"])])

    PROFILE_STATE = (
        "# Состояние проекта ВКР\n\n## Стилевой профиль\n\n"
        "### Личный стилистический профиль (personal_style_profile)\n\n"
        "**Любимые обороты пользователя (из её README):**\n- «любимое словечко»\n\n"
        "**Якорьки для ВКР (3 штуки):**\n"
        "- «если смотреть глазами ученика» — плановое число использований: 3–4\n"
        "- «Ошибка на отдельном шаге» — плановое число использований: 3–4\n"
        "- [«якорёк 3»] — плановое число использований: 3–5\n\n"
        "**Оборот для критики/сомнения:** «это верно с оговоркой»\n\n"
        "**Оборот для технического вывода:** [«…»]\n\n"
        "### Запретные обороты\n\n- «на наш взгляд»\n\n"
        "## Источники\n\n**Якорьки для ВКР:**\n- «не из профиля»\n"
    )

    def test_personal_anchors_are_parsed_from_state_profile(self):  # B14 (F18)
        self.assertEqual(
            ["если смотреть глазами ученика", "ошибка на отдельном шаге", "это верно с оговоркой"],
            ai.parse_personal_anchors(self.PROFILE_STATE),
        )
        template = (SKILL / "references" / "state-file-pattern.md").read_text(encoding="utf-8")
        self.assertEqual([], ai.parse_personal_anchors(template))  # заглушки шаблона [«якорёк 1»] не якорьки
        self.assertEqual([], ai.parse_personal_anchors("# State\n\n**Якорьки для ВКР:**\n- «вне профиля»\n"))

    def test_author_voice_counts_personal_anchors_from_project_state(self):  # B14 (F18)
        def section(tag, anchors):
            paragraphs = [prose(f"{tag} абзац {i}", 8, i) for i in range(14)]
            for index, anchor in enumerate(anchors):
                paragraphs[2 + index * 3] = anchor + " " + paragraphs[2 + index * 3]
            return [("p", text) for text in paragraphs]

        items = ([("h1", "Введение")] + section("Введение", ["Если смотреть глазами ученика, подсказка должна быть короткой."])
                 + [("h1", "Глава I. Теория")] + section("Теория", ["Ошибка на отдельном шаге видна сразу.", "Если смотреть глазами ученика, шаги должны быть видимыми."])
                 + [("h1", "Заключение")] + section("Итог", ["Ошибка на отдельном шаге больше не прячется в итоговом ответе."]))
        project = self.dir / "project"
        docx_path = build_docx(project / "final" / "vkr.docx", items)
        (project / "vkr-project.json").write_text("{}", encoding="utf-8")

        def report():
            proc = run_script("ai_detection_heuristic.py", docx_path, "--json")
            self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
            return json.loads(proc.stdout)

        without = report()
        self.assertGreater(without["meta"]["words"], 3000)
        self.assertEqual("no author voice (sounds AI-generic)", without["anchors"]["verdict"])
        self.assertIn("No author voice (sounds AI-generic)", without["risk"]["flags"])
        self.assertEqual("builtin", without["anchors"]["anchor_source"])

        (project / "vkr-state.md").write_text(self.PROFILE_STATE, encoding="utf-8")
        payload = report()
        self.assertEqual(4, payload["anchors"]["total_anchors"])
        self.assertEqual("balanced", payload["anchors"]["verdict"])
        self.assertNotIn("No author voice (sounds AI-generic)", payload["risk"]["flags"])
        self.assertEqual("personal_style_profile+builtin", payload["anchors"]["anchor_source"])
        self.assertEqual(ai.parse_personal_anchors(self.PROFILE_STATE), payload["meta"]["personal_anchors"]["phrases"])
        sections = json.loads(run_script("ai_detection_heuristic.py", docx_path, "--per-section", "--json").stdout)
        self.assertEqual(2, sections["sections"]["chapter1"]["anchors"]["total_anchors"])

    def test_ownership_sampling_no_placeholder_and_no_bibliography(self):
        items = body() + [("h1", "Библиографический список"), ("bib_text", ENTRIES), ("last_sheet",)]
        extra = [("p", prose(f"Дополнительный абзац главы {i}", 3, i)) for i in range(30)]
        items = items[:8] + extra + items[8:]
        path = build_docx(self.dir / "many.docx", items)
        proc = run_script("content_ownership_check.py", path, "--all", "--json")
        self.assertEqual(0, proc.returncode, proc.stderr)
        payload = json.loads(proc.stdout)
        self.assertEqual("all", payload["sampling"])
        self.assertEqual(payload["available_paragraphs"], payload["total_paragraphs"])
        questions = [question["question"] for item in payload["items"] for question in item["questions"]]
        self.assertTrue(questions)
        self.assertFalse([question for question in questions if "{" in question or "}" in question])
        texts = [item["full_text"] for item in payload["items"]]
        self.assertFalse(any("Абрамова" in text or "Рапай" in text or "Smith" in text for text in texts))

        sample = run_script("content_ownership_check.py", path, "--json", "--seed", "7")
        again = run_script("content_ownership_check.py", path, "--json", "--seed", "7")
        other = run_script("content_ownership_check.py", path, "--json", "--seed", "8")
        self.assertEqual(sample.stdout, again.stdout)
        first, second = json.loads(sample.stdout), json.loads(other.stdout)
        self.assertEqual("sample", first["sampling"])
        self.assertEqual(25, first["total_paragraphs"])
        self.assertIn("ВЫБОРКА", first["sampling_note"])
        self.assertNotEqual([i["paragraph_id"] for i in first["items"]], [i["paragraph_id"] for i in second["items"]])
        text = run_script("content_ownership_check.py", path)
        self.assertIn("ВЫБОРКА: 25 из", text.stdout)
        self.assertIn("--all", text.stdout)
        self.assertNotIn("Прогони через эти вопросы КАЖДЫЙ абзац своей работы", text.stdout)


class RegexNormalizationTest(unittest.TestCase):
    def test_passive_voice_ignores_yo_and_false_auxiliary(self):
        plain = ai.check_passive_voice_ratio("Анализ был проведен. Эксперимент проведен успешно. Метод определен. Результат отражен в таблице.")
        yo = ai.check_passive_voice_ratio("Анализ был проведён. Эксперимент проведён успешно. Метод определён. Результат отражён в таблице.")
        self.assertEqual(plain["passive_markers_count"], yo["passive_markers_count"])
        self.assertEqual(4, yo["passive_markers_count"])
        false = ai.check_passive_voice_ratio("Работа была интересной. Он был в школе. Будет интересно посмотреть. Страна была большой.")
        self.assertEqual(0, false["passive_markers_count"])

    def test_pronouns_include_possessive_forms(self):
        self.assertGreater(ai.check_pronouns("На наш взгляд, наши результаты важны.")["total"], 0)
        self.assertEqual(0, ai.check_pronouns("Методика И.Я. Лернера была использована.")["total"])

    def test_ownership_markers_ignore_yo(self):
        for text in ("Для реализации был применен движок Unity.", "Для реализации был применён движок Unity."):
            self.assertIn("technology-choice", [p["category"] for p in own.detect_applicable_patterns(text)])
        self.assertEqual(
            ai.compute_ngram_density("Всё большую актуальность приобретает тема.")["strong_cliches_count"],
            ai.compute_ngram_density("Все большую актуальность приобретает тема.")["strong_cliches_count"],
        )


if __name__ == "__main__":
    unittest.main()
