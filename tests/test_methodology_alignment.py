"""Сверка справочников скилла с документом методички assets/methodology-09-03-02-2024.docx.

Запуск: ``python -m unittest discover -s tests -t .`` или
``python tests/test_methodology_alignment.py``.
"""

from __future__ import annotations

import hashlib
import os
import re
import sys
import unittest
from pathlib import Path

os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
sys.dont_write_bytecode = True

try:
    from docx import Document
except ImportError:  # pragma: no cover - окружение без python-docx
    Document = None


ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "vkr-mpgu"
METHODOLOGY = SKILL / "assets" / "methodology-09-03-02-2024.docx"
METHODOLOGY_SHA256 = "8d890c1474c0621a8e5259b514126935c92d94393127c2d27094fa80cb447eae"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def docx_text(path: Path) -> str:
    doc = Document(path)
    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return re.sub(r"[  ]+", " ", "\n".join(chunks)).lower()


def reference(name: str) -> str:
    return (SKILL / "references" / name).read_text(encoding="utf-8").lower()


@unittest.skipIf(Document is None, "python-docx не установлен")
class MethodologySourceTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.source = docx_text(METHODOLOGY)

    def test_source_document_is_unchanged(self) -> None:
        self.assertEqual(METHODOLOGY_SHA256, sha256(METHODOLOGY))

    def test_source_contains_facts_the_skill_relies_on(self) -> None:
        for fragment in (
            "от 50 до 60 стр",
            "от 40 до 50 стр",
            "работа выполняется на русском языке",
            "глава 1. теоретическая часть",
            "глава 2. аналитическая часть",
            "глава 3. собственно проектная часть",
            "образовательного продукта",
            "этап пилотирования",
            "не менее 5 человек",
            "левого поля – 35 мм",
            "три пустых файла-вкладыша",
            "список составляется в алфавитном порядке",
            "проверяется на антиплагиат один раз",
            "принципу конструктивной согласованности",
            "новизну работы",
        ):
            self.assertIn(fragment, self.source, fragment)

    def test_originality_scale_matches_source(self) -> None:
        for fragment in ("80% и выше – “отлично”", "70–79% – “хорошо”", "50–69% – “удовлетворительно”",
                         "90% и выше – “отлично”", "80–89% – “хорошо”", "70–79% – “удовлетворительно”"):
            self.assertIn(fragment, self.source, fragment)

    def test_chapter_numbering_ambiguity_is_real_and_documented(self) -> None:
        # Правило в обоих разделах — римские цифры; один пример арабский, другой римский.
        self.assertIn("нумерация глав – римскими цифрами", self.source)
        self.assertIn("например, глава 1.", self.source)
        self.assertIn("например, глава i.", self.source)
        methodology_ref = reference("methodology-mpgu.md")
        self.assertIn("### глава i. теоретическая часть", methodology_ref)
        self.assertIn("толкование", methodology_ref)

    def test_origin_statements_match_source(self) -> None:
        # Справочники говорят: кода направления в тексте нет, форма справки называет Институт международного образования.
        self.assertNotIn("09.03.02", self.source)
        self.assertIn("института международного образования", self.source)
        for name in ("methodology-mpgu.md", "requirements-precedence.md"):
            text = " ".join(reference(name).split())
            self.assertIn("кода направления в тексте нет", text, name)
            self.assertIn("институт международного образования", text, name)

    def test_presentation_requirements_are_complete_in_defense(self) -> None:
        defense = reference("defense.md").replace("\n", " ")
        for fragment in ("конструктивной согласованности", "целевая аудитория", "новизн", "пилотный запуск",
                         "рекомендации", "методики и выборк", "пример задания", "не задаёт число слайдов"):
            self.assertIn(fragment, defense, fragment)


class ReferenceAlignmentTest(unittest.TestCase):
    def test_methodology_reference(self) -> None:
        methodology_ref = reference("methodology-mpgu.md")
        for fragment in ("40–50", "50–60", "образовательный", "пилот", "минимум 5", "левое поле",
                         "80%+", "50–69%", "90%+", "70–79%", "три главы"):
            self.assertIn(fragment, methodology_ref, fragment)
        self.assertNotIn("три главы приведены в методичке как пример", methodology_ref)

    def test_requirements_precedence(self) -> None:
        precedence = reference("requirements-precedence.md")
        for fragment in ("mpgu-09-project", "три главы", "частной практике", "не устанавливает минимум источников",
                         "справк", "vkr-project.json"):
            self.assertIn(fragment, precedence, fragment)

    def test_structure_project(self) -> None:
        structure = reference("structure-project.md")
        for fragment in ("глава i. теоретическая", "глава ii. аналитическая", "глава iii. собственно проектная",
                         "profile", "mpgu-09-project", "2 страницы"):
            self.assertIn(fragment, structure, fragment)
        self.assertNotIn("выводы по каждой главе есть и пронумерованы", structure)

    def test_originality_targets_are_consistent(self) -> None:
        for name in ("originality-techniques.md", "failure-recovery.md", "writing-workflow-express.md"):
            text = reference(name)
            self.assertNotIn("подозрительно", text, name)
            self.assertNotIn("удовл-хорошо", text, name)
        failure = reference("failure-recovery.md")
        self.assertIn("один раз", failure)
        self.assertNotIn("меньше 80%, проектная — меньше 90%", failure)
        skill = (SKILL / "SKILL.md").read_text(encoding="utf-8")
        self.assertIn("85–93%", skill)
        self.assertIn("90–94%", skill)
        self.assertIn("50–69%", skill)

    def test_validator_and_generator_profile(self) -> None:
        validator = (SKILL / "scripts" / "validate_vkr.py").read_text(encoding="utf-8")
        generator = (SKILL / "scripts" / "create_vkr_docx.py").read_text(encoding="utf-8")
        for fragment in ('"page_width_mm": 210', '"page_height_mm": 297', '"mpgu-09-project"', "def check_mpgu_09_profile"):
            self.assertIn(fragment, validator, fragment)
        self.assertIn("section.page_width = Mm(210)", generator)
        self.assertIn("section.page_height = Mm(297)", generator)


if __name__ == "__main__":
    unittest.main(verbosity=2)
