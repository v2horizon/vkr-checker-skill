"""Тесты validate_vkr 6.33 (WP3).

Фикстуры в tests/fixtures/validator/ сохранены Microsoft Word через COM (см.
manifest.json и build/); Word для запуска тестов не нужен. Остальные случаи
собираются python-docx в памяти или правкой XML Word-фикстур во временном каталоге.

Запуск: python -m unittest discover -s tests -t .   или   python tests/test_validator_633.py
"""
from __future__ import annotations

import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "vkr-mpgu" / "scripts"
VALIDATOR_PATH = SCRIPTS / "validate_vkr.py"
FIXTURES = ROOT / "tests" / "fixtures" / "validator"
PROFILES = ("generic", "mpgu-09-regular", "mpgu-09-project")


def load_validator():
    spec = importlib.util.spec_from_file_location("validate_vkr_633_under_test", VALIDATOR_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


V = load_validator()


def failures(results, severity=None):
    return [
        item["message"] for item in results
        if not item["ok"] and (severity is None or item["severity"] == severity)
    ]


def report_messages(report, severity="error"):
    return [
        item["message"]
        for items in report["checks"].values()
        for item in items
        if not item["ok"] and item["severity"] == severity
    ]


def fixture(name):
    path = FIXTURES / name
    if not path.is_file():
        raise AssertionError(f"нет фикстуры {path}")
    return path


def new_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    return doc


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def rewrite_part(src, dst, part, replace):
    """Копия DOCX с заменой текста в XML-части (replace: [(old, new)])."""
    with zipfile.ZipFile(src) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in entries:
            if info.filename == part:
                text = data.decode("utf-8")
                for old, new in replace:
                    if old not in text:
                        raise AssertionError(f"{old!r} не найдено в {part}")
                    text = text.replace(old, new)
                data = text.encode("utf-8")
            archive.writestr(info, data)


def run_cli(args, block_docx=False, cwd=None):
    env = dict(os.environ)
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env.pop("PYTHONIOENCODING", None)
    if block_docx:
        code = (
            "import runpy, sys; sys.modules['docx'] = None; "
            f"sys.argv = [{str(VALIDATOR_PATH)!r}] + {list(args)!r}; "
            f"runpy.run_path({str(VALIDATOR_PATH)!r}, run_name='__main__')"
        )
        command = [sys.executable, "-c", code]
    else:
        command = [sys.executable, str(VALIDATOR_PATH)] + list(args)
    result = subprocess.run(command, capture_output=True, env=env, cwd=cwd, timeout=300)
    return result.returncode, result.stdout.decode("utf-8"), result.stderr.decode("utf-8")


# ---------------------------------------------------------------------------
# CLI и API (A5, A-632-R4, D13, R3-06)
# ---------------------------------------------------------------------------

class CliContractTests(unittest.TestCase):
    def test_api_report_has_version_and_status(self):
        report = V.validate_vkr(str(fixture("gen_project_word.docx")), profile="mpgu-09-project")
        self.assertEqual("6.33", report["validator_version"])
        self.assertEqual("ok", report["status"])
        self.assertEqual(64, len(report["document_sha256"]))
        self.assertEqual(0, report["summary"]["errors"], report_messages(report))

    def test_api_errors(self):
        missing = V.validate_vkr("нет-такого-файла.docx")
        self.assertEqual(("error", "FILE_NOT_FOUND"), (missing["status"], missing["error_code"]))
        bad_profile = V.validate_vkr(str(fixture("gen_project_word.docx")), profile="bogus")
        self.assertEqual("UNKNOWN_PROFILE", bad_profile["error_code"])
        with tempfile.TemporaryDirectory() as tmp:
            fake = Path(tmp) / "fake.docx"
            fake.write_bytes(b"not a zip")
            broken = V.validate_vkr(str(fake))
        self.assertEqual("INVALID_DOCX", broken["error_code"])
        self.assertEqual("6.33", broken["validator_version"])

    def test_exit_codes_and_json(self):
        code, out, _ = run_cli([str(fixture("gen_project_word.docx")), "--profile", "mpgu-09-project", "--json"])
        self.assertEqual(0, code)
        self.assertEqual(0, json.loads(out)["summary"]["errors"])
        code, out, _ = run_cli([str(fixture("bad_margin_left30.docx")), "--profile", "mpgu-09-project", "--json"])
        self.assertEqual(1, code)
        self.assertGreater(json.loads(out)["summary"]["errors"], 0)
        code, out, _ = run_cli(["нет-файла.docx", "--json"])
        self.assertEqual(2, code)
        self.assertEqual("FILE_NOT_FOUND", json.loads(out)["error_code"])
        code, out, _ = run_cli([str(fixture("gen_project_word.docx")), "--profile", "bogus", "--json"])
        self.assertEqual(2, code)
        self.assertEqual("UNKNOWN_PROFILE", json.loads(out)["error_code"])

    def test_output_file_is_utf8_without_bom_and_dirs_created(self):
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "audit" / "nested" / "automated-validation.json"
            code, _out, _err = run_cli([str(fixture("ok_title_section.docx")), "--profile", "mpgu-09-project", "-o", str(target)])
            self.assertEqual(0, code)
            raw = target.read_bytes()
            self.assertFalse(raw.startswith(b"\xef\xbb\xbf"))
            self.assertEqual("6.33", json.loads(raw.decode("utf-8"))["validator_version"])

    def test_without_python_docx(self):
        code, out, err = run_cli(["--help"], block_docx=True)
        self.assertEqual(0, code)
        self.assertIn("--profile", out)
        code, out, err = run_cli([str(fixture("gen_project_word.docx")), "--json"], block_docx=True)
        self.assertEqual(2, code)
        payload = json.loads(out)
        self.assertEqual("PYTHON_DOCX_MISSING", payload["error_code"])
        self.assertIn("python-docx", payload["error"])
        code, out, err = run_cli([str(fixture("gen_project_word.docx"))], block_docx=True)
        self.assertEqual(2, code)
        self.assertEqual("", out)
        self.assertIn("python-docx не установлен", err)
        with tempfile.TemporaryDirectory() as tmp:
            target = Path(tmp) / "report.json"
            code, _out, _err = run_cli([str(fixture("gen_project_word.docx")), "-o", str(target)], block_docx=True)
            self.assertEqual(2, code)
            self.assertEqual("PYTHON_DOCX_MISSING", json.loads(target.read_text(encoding="utf-8"))["error_code"])


# ---------------------------------------------------------------------------
# Word-фикстуры: правильные → 0 ошибок, испорченные → нужная ошибка
# ---------------------------------------------------------------------------

OK_FIXTURES = (
    "gen_project_word.docx",
    "ok_pagenum_building_block.docx",
    "ok_toc_sdt.docx",
    "ok_title_section.docx",
    "ok_headings_recolored_auto.docx",
    "ok_spec8_styles.docx",
    "ok_docdefaults_spacing15.docx",
    "ok_docdefaults_font_tnr14.docx",
    "ok_docdefaults_indent125.docx",
)

# файл → фрагменты ожидаемых ОШИБОК (профиль mpgu-09-project); других ошибок быть не должно
BAD_FIXTURES = {
    "bad_margin_left30.docx": ["левое поле: 30.0 мм"],
    "bad_normal_font_arial.docx": ["Основной шрифт: Arial", "Заголовки набраны шрифтом Arial"],
    "bad_normal_size12.docx": ["Основной размер шрифта: 12 pt"],
    "bad_normal_spacing_single.docx": ["междустрочный интервал — «1.0 (single)»"],
    "bad_normal_indent0.docx": ["Отступ 1,25 см НЕ соблюдён"],
    "bad_docdefaults_spacing_single.docx": ["междустрочный интервал — «1.0 (single)»"],
    "bad_docdefaults_font_arial12.docx": ["Основной шрифт: Arial", "Заголовки набраны шрифтом Arial", "Основной размер шрифта: 12 pt"],
    "bad_docdefaults_indent0.docx": ["Отступ 1,25 см НЕ соблюдён"],
    "bad_body_align_left.docx": ["Выравнивание основного текста влево"],
    "bad_no_page_field.docx": ["Нет поля PAGE в секциях: 1"],
    "bad_title_page_number.docx": ["На титульном листе выводится номер страницы"],
    "bad_letter.docx": ["215.9 × 279.4 мм"],
    "bad_marker_in_header.docx": ["[ТРЕБУЕТ УТОЧНЕНИЯ] (верхний колонтитул)"],
    "bad_second_section_margins.docx": ["Секция 2, левое поле: 15.0 мм", "Секция 2, правое поле: 25.0 мм"],
    "bad_arabic_chapter.docx": ["не римскими цифрами"],
    "bad_arabic_chapter_normal_bold.docx": ["не римскими цифрами"],
    "bad_heading_trailing_dot.docx": ["Точка в конце заголовка: 3.1. Пилотное тестирование."],
    "bad_citation_after_period.docx": ["Ссылка стоит после точки"],
    "bad_chapter_duplicate_prefix.docx": ["Префикс главы повторяется"],
    "bad_bibliography_unsorted.docx": ["не в алфавитном порядке"],
    "bad_trailing_empty_page.docx": ["пустая страница с номером"],
    "bad_headings_blue.docx": ["не-чёрный цвет #4F81BD"],
}


class WordFixtureTests(unittest.TestCase):
    def test_correct_word_documents_have_zero_errors_in_all_profiles(self):
        for name in OK_FIXTURES:
            for profile in PROFILES:
                with self.subTest(name=name, profile=profile):
                    report = V.validate_vkr(str(fixture(name)), profile=profile)
                    self.assertEqual(0, report["summary"]["errors"], report_messages(report))

    def test_corrupted_copies_give_exactly_the_expected_errors(self):
        for name, expected in BAD_FIXTURES.items():
            with self.subTest(name=name):
                report = V.validate_vkr(str(fixture(name)), profile="mpgu-09-project")
                errors = report_messages(report)
                for fragment in expected:
                    self.assertTrue(any(fragment in message for message in errors), (fragment, errors))
                self.assertEqual(len(expected), len(errors), errors)

    def test_generic_profile_softens_methodology_specific_rules(self):
        for name, fragment in (
            ("bad_bibliography_unsorted.docx", "не в алфавитном порядке"),
            ("bad_arabic_chapter.docx", "не римскими цифрами"),
        ):
            with self.subTest(name=name):
                report = V.validate_vkr(str(fixture(name)), profile="generic")
                self.assertEqual(0, report["summary"]["errors"], report_messages(report))
                self.assertTrue(any(fragment in m for m in report_messages(report, "warning")))

    def test_sdt_structures_are_really_present_in_fixtures(self):
        with zipfile.ZipFile(fixture("ok_pagenum_building_block.docx")) as archive:
            footers = "".join(archive.read(n).decode("utf-8") for n in archive.namelist() if n.startswith("word/footer"))
        self.assertIn("Page Numbers (Bottom of Page)", footers)
        self.assertIn("<w:sdt>", footers)
        with zipfile.ZipFile(fixture("ok_toc_sdt.docx")) as archive:
            body = archive.read("word/document.xml").decode("utf-8")
        self.assertIn('w:docPartGallery w:val="Table of Contents"', body)


# ---------------------------------------------------------------------------
# Номера страниц (REG-D-01, REG-D-02, R-F-10, R2-05, находка скептика F)
# ---------------------------------------------------------------------------

class PageNumberTests(unittest.TestCase):
    def mechanics(self, doc_or_path, profile="mpgu-09-project"):
        doc = Document(str(doc_or_path)) if isinstance(doc_or_path, Path) else doc_or_path
        return V.check_document_mechanics(doc, profile)

    def test_page_field_inside_sdt_gallery_block_is_found(self):
        results = self.mechanics(fixture("ok_pagenum_building_block.docx"))
        self.assertFalse(any("PAGE" in message for message in failures(results, "error")), failures(results))

    def test_fldsimple_page_is_found(self):
        doc = Document(str(fixture("bad_no_page_field.docx")))
        paragraph = doc.sections[0].footer.paragraphs[0]
        simple = OxmlElement("w:fldSimple")
        simple.set(qn("w:instr"), " PAGE ")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "2"
        run.append(text)
        simple.append(run)
        paragraph._p.append(simple)
        self.assertEqual([], failures(self.mechanics(doc), "error"))

    def test_numpages_alone_is_not_a_page_number(self):
        doc = Document(str(fixture("bad_no_page_field.docx")))
        paragraph = doc.sections[0].footer.paragraphs[0]
        simple = OxmlElement("w:fldSimple")
        simple.set(qn("w:instr"), " NUMPAGES ")
        paragraph._p.append(simple)
        self.assertTrue(any("Нет поля PAGE" in m for m in failures(self.mechanics(doc), "error")))

    def test_title_page_as_separate_section_without_footer(self):
        results = self.mechanics(fixture("ok_title_section.docx"))
        self.assertEqual([], failures(results, "error"))
        self.assertTrue(any(item["ok"] and "титульном листе не выводится" in item["message"] for item in results))

    def test_title_page_message_is_not_inverted(self):
        results = self.mechanics(fixture("bad_title_page_number.docx"))
        self.assertTrue(any("На титульном листе выводится номер" in m for m in failures(results, "error")))
        doc = Document(str(fixture("gen_project_word.docx")))
        doc.sections[0].different_first_page_header_footer = False
        messages = failures(self.mechanics(doc), "error")
        self.assertTrue(any("На титульном листе выводится номер" in m for m in messages), messages)
        self.assertFalse(any("включён особый первый колонтитул" in m for m in messages))

    def test_restart_after_title_section_is_error_for_mpgu(self):
        doc = Document(str(fixture("ok_title_section.docx")))
        sect = doc.sections[1]._sectPr
        pg = sect.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sect.insert(0, pg)
        pg.set(qn("w:start"), "1")
        self.assertTrue(any("нумерация начинается заново с 1" in m for m in failures(self.mechanics(doc), "error")))


# ---------------------------------------------------------------------------
# Эффективное форматирование (D1, D2, R2-04, D9)
# ---------------------------------------------------------------------------

class EffectiveFormattingTests(unittest.TestCase):
    def body_doc(self, count=6):
        doc = new_doc()
        add_heading(doc, "Введение")
        for index in range(count):
            doc.add_paragraph(
                f"Абзац {index} основного текста работы описывает исследование и содержит достаточно "
                f"длинное предложение, чтобы проверка выравнивания и отступа учитывала его как основной текст."
            )
        return doc

    def test_spacing_from_custom_style_based_on_normal(self):
        doc = self.body_doc()
        doc.styles["Normal"].paragraph_format.line_spacing = 1.0
        from docx.enum.style import WD_STYLE_TYPE
        style = doc.styles.add_style("Основной текст ВКР", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        style.paragraph_format.line_spacing = 1.5
        for paragraph in doc.paragraphs[1:]:
            paragraph.style = style
        self.assertEqual([], failures(V.check_line_spacing(doc), "error"))
        doc.styles["Основной текст ВКР"].paragraph_format.line_spacing = 1.0
        self.assertTrue(failures(V.check_line_spacing(doc), "error"))

    def test_spacing_and_indent_from_docdefaults(self):
        doc = self.body_doc()
        normal = doc.styles["Normal"].element
        ppr = normal.find(qn("w:pPr"))
        if ppr is not None:
            normal.remove(ppr)
        defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:pPrDefault")).find(qn("w:pPr"))
        spacing = defaults.find(qn("w:spacing"))
        spacing.set(qn("w:line"), "360")
        spacing.set(qn("w:lineRule"), "auto")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:firstLine"), "709")
        defaults.append(ind)
        self.assertEqual([], failures(V.check_line_spacing(doc), "error"))
        self.assertEqual([], failures(V.check_first_line_indent(doc)))
        ind.set(qn("w:firstLine"), "0")
        self.assertTrue(failures(V.check_first_line_indent(doc), "error"))

    def test_normal_style_indent_zero_is_error_not_pass(self):
        doc = self.body_doc()
        doc.styles["Normal"].paragraph_format.first_line_indent = Cm(0)
        self.assertTrue(failures(V.check_first_line_indent(doc), "error"))

    def test_left_alignment_from_style_is_error(self):
        doc = self.body_doc()
        doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        messages = failures(V.check_document_mechanics(doc, "generic"), "error")
        self.assertTrue(any("Выравнивание основного текста влево" in m for m in messages), messages)
        doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        messages = failures(V.check_document_mechanics(doc, "generic"))
        self.assertFalse(any("Выравнивание" in m for m in messages), messages)

    def test_font_size_and_color_from_styles(self):
        doc = self.body_doc()
        doc.styles["Normal"].font.name = "Times New Roman"
        doc.styles["Normal"].font.size = Pt(12)
        self.assertTrue(any("12 pt" in m for m in failures(V.check_fonts(doc), "error")))
        doc.styles["Normal"].font.size = Pt(14)
        heading = doc.styles["Heading 1"]
        heading.font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor(0x2F, 0x54, 0x96)
        self.assertTrue(failures(V.check_font_color(doc), "error"))
        for run in doc.paragraphs[0].runs:
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "auto")
            run._r.get_or_add_rPr().append(color)
        self.assertEqual([], failures(V.check_font_color(doc), "error"))

    def test_theme_text_color_is_black(self):
        doc = self.body_doc()
        doc.paragraphs[0].style = doc.styles["Normal"]  # у Heading 1 шаблона python-docx синий цвет темы
        run = doc.paragraphs[1].runs[0]
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        color.set(qn("w:themeColor"), "text1")
        run._r.get_or_add_rPr().append(color)
        self.assertEqual([], failures(V.check_font_color(doc), "error"))

    def test_hansi_slot_used_for_cyrillic(self):
        doc = self.body_doc()
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                fonts = OxmlElement("w:rFonts")
                fonts.set(qn("w:ascii"), "Times New Roman")
                fonts.set(qn("w:hAnsi"), "Arial")
                run._r.get_or_add_rPr().append(fonts)
        self.assertTrue(any("Arial" in m for m in failures(V.check_fonts(doc), "error")))

    def test_page_break_before_in_style_counts_as_new_page(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Текст введения " * 10)
        doc.styles["Heading 1"].paragraph_format.page_break_before = True
        add_heading(doc, "Глава I. Теория")
        doc.add_paragraph("Текст главы " * 10)
        messages = failures(V.check_document_mechanics(doc, "generic"), "warning")
        self.assertFalse(any("Нет разрыва страницы" in m for m in messages), messages)


# ---------------------------------------------------------------------------
# Исключения из метрик основного текста (REG-D-04, R-F-09, R2-15, D3)
# ---------------------------------------------------------------------------

class ExclusionTests(unittest.TestCase):
    def test_vkr_title_on_title_page_and_last_page(self):
        """Генератор 6.33 (WP2) ставит стиль VKR Title и на титул, и на последний лист."""
        from docx.enum.style import WD_STYLE_TYPE

        doc = new_doc()
        doc.styles.add_style("VKR Title", WD_STYLE_TYPE.PARAGRAPH)
        doc.styles.add_style("VKR Appendix Label", WD_STYLE_TYPE.PARAGRAPH)
        for text in ("Лебедева Анна Сергеевна", "Научный руководитель – доцент Кузнецов"):
            doc.add_paragraph(text, style="VKR Title")
        doc.add_paragraph("Аннотация")
        doc.add_paragraph("Работа изложена на 45 страницах.")
        add_heading(doc, "Введение")
        doc.add_paragraph("Педагогическое пилотирование описано в работе [1].")
        for number in ("I", "II", "III"):
            add_heading(doc, f"Глава {number}. Раздел")
            doc.add_paragraph("Текст главы с апробацией и ссылкой [1].")
        add_heading(doc, "Заключение")
        doc.add_paragraph("Итоги работы.")
        add_heading(doc, "Список использованной литературы")
        doc.add_paragraph("1. Андреев А.А. Книга. – М., 2020. – 100 с.")
        doc.add_paragraph("Приложение 1", style="VKR Appendix Label")
        doc.add_paragraph("Выпускная квалификационная работа", style="VKR Title")
        doc.add_paragraph("(Ф.И.О.)        (подпись)        (дата)", style="VKR Title")
        self.assertTrue(all(item["ok"] for item in V.check_mpgu_09_profile(doc, "mpgu-09-project")))
        self.assertEqual([], failures(V.check_cross_references(doc, "mpgu-09-project")))
        self.assertEqual([], failures(V.check_surnames_format(doc)))

    def test_lists_code_captions_bibliography_title_do_not_break_metrics(self):
        report = V.validate_vkr(str(fixture("ok_spec8_styles.docx")), profile="mpgu-09-project")
        for group in ("Отступ первой строки", "Междустрочный интервал", "Нумерация и формальные признаки", "Шрифт и размер"):
            self.assertEqual([], failures(report["checks"][group]), group)

    def test_generator_list_items_are_not_body_text(self):
        doc = new_doc()
        doc.styles["Normal"].paragraph_format.first_line_indent = Cm(1.25)
        doc.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_heading(doc, "Введение")
        doc.add_paragraph("Основной текст " * 12)
        for _ in range(12):
            item = doc.add_paragraph("пункт списка, который длиннее восьмидесяти знаков и выровнен влево без абзацного отступа;", style="List Bullet")
            item.alignment = WD_ALIGN_PARAGRAPH.LEFT
            item.paragraph_format.first_line_indent = Cm(0)
        self.assertEqual([], failures(V.check_first_line_indent(doc)))
        self.assertFalse(any("Выравнивание" in m for m in failures(V.check_document_mechanics(doc, "generic"))))


# ---------------------------------------------------------------------------
# Ссылки на источники (R2-01, REG-D-03, R-F-03, R2-02, R2-03)
# ---------------------------------------------------------------------------

class CitationTests(unittest.TestCase):
    def test_page_numbers_are_not_source_numbers(self):
        cases = {
            "[15, с. 23–25]": [15], "[15, с. 23—25]": [15], "[15, с. 23-25]": [15], "[15, p. 23]": [15],
            "[15, S. 23]": [15], "[15, С. 23]": [15], "[15, т. 2, с. 45]": [15], "[15, ч. 3, с. 12]": [15],
            "[4, с. 112–114]": [4], "[5, p. 12]": [5], "[15; 17]": [15, 17], "[15, 17, 19]": [15, 17, 19],
            "[Цит. по: 27, с. 235]": [27], "[См.: 12; 13]": [12, 13], "[15]": [15], "[1; 5, с. 7]": [1, 5],
            "arr[25]": [], "matrix[n][30]": [], "[x]": [], "[N]": [], "[Электронный ресурс]": [], "[2019]": [],
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(expected, V._extract_citation_numbers(text))

    def test_citation_after_abbreviations_is_correct(self):
        good = (
            "Сходные данные приводят А.В. Хуторской и др. [19, с. 12].",
            "Нормы опубликованы в 2021 г. [11].",
            "Исследования 2020–2023 гг. [5] подтвердили вывод.",
            "Подход описан в другом издании (см. [19, с. 40]).",
            "Типовые ошибки связаны с индексами, т. е. [11, с. 36] с пошаговым исполнением.",
            "Описание приведено в XIX в. [3], а также в XX вв. [4].",
            "Цитата заканчивается многоточием… [2].",
            "Цитата заканчивается тремя точками... [2].",
            "Работы J. Smith et al. [7] и Brown etc. [8] рассмотрены.",
            "Метод описан в разделе 2.1. [3] настоящей работы.",
            "Мысль высказана Л.С. [2] и развита позже.",
        )
        doc = new_doc()
        add_heading(doc, "Введение")
        for sentence in good:
            doc.add_paragraph(sentence)
        messages = failures(V.check_document_mechanics(doc, "mpgu-09-project"))
        self.assertFalse(any("после точки" in m for m in messages), messages)

    def test_real_citation_after_period_is_error(self):
        for sentence in ("Вывод подтверждён литературой. [2, с. 12]", "Вывод подтверждён литературой.[2, с. 12]", "Итог «кавычки». [3]"):
            with self.subTest(sentence=sentence):
                doc = new_doc()
                add_heading(doc, "Введение")
                doc.add_paragraph(sentence)
                self.assertTrue(any("после точки" in m for m in failures(V.check_document_mechanics(doc, "generic"), "error")))

    def test_code_is_not_a_citation_or_marker(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Результаты сравнения приведены в работе [1].")
        for line in ("history[x] = matrix[n][30]", "<span>{{ user.name }}</span>", "# TODO: учитывать порядок", "data[25] * weight[n]"):
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
        add_heading(doc, "Список литературы")
        doc.add_paragraph("1. Андреев А.А. Книга. – М., 2020. – 100 с.")
        self.assertEqual([], failures(V.check_unresolved_placeholders(doc)))
        self.assertEqual([], failures(V.check_cross_references(doc, "generic"), "error"))
        mechanics = failures(V.check_document_mechanics(doc, "generic"))
        self.assertFalse(any("после точки" in m for m in mechanics), mechanics)

    def test_inline_code_fragment_is_skipped_but_prose_marker_is_caught(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        paragraph = doc.add_paragraph("Функция ")
        code = paragraph.add_run("items[x] = {{ value }}")
        code.font.name = "Courier New"
        paragraph.add_run(" возвращает [N] элементов.")
        messages = failures(V.check_unresolved_placeholders(doc))
        self.assertEqual(1, len(messages))
        self.assertIn("[N]", messages[0])
        self.assertNotIn("{{", messages[0])

    def test_lowercase_math_and_todo_word_are_not_markers(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Отрезок [x] и множество [n] рассмотрены в учебном ToDo-сервисе; XXX Олимпиада.")
        self.assertEqual([], failures(V.check_unresolved_placeholders(doc)))


# ---------------------------------------------------------------------------
# Сверка ссылок со списком литературы (REG-D-12, R3-02, R-F-06, R2-09, #36, скептик R2)
# ---------------------------------------------------------------------------

class CrossReferenceTests(unittest.TestCase):
    def doc_with_bibliography(self, heading, entries, text="Текст со ссылкой [99, с. 3].", style=None):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph(text)
        add_heading(doc, heading)
        for entry in entries:
            doc.add_paragraph(entry, style=style) if style else doc.add_paragraph(entry)
        return doc

    def test_heading_variants(self):
        for heading in (
            "Список использованной литературы", "Список использованных источников", "Список литературы",
            "Библиографический список", "Литература", "Список использованной литературы и источников",
            "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
        ):
            with self.subTest(heading=heading):
                doc = self.doc_with_bibliography(heading, ["1. Андреев А.А. Книга. – М., 2020. – 100 с."])
                self.assertTrue(any("[99]" in m for m in failures(V.check_cross_references(doc), "error")))
                self.assertTrue(all(item["ok"] for item in V.check_structure(doc) if "Список литературы" in item["message"]))

    def test_autonumbered_word_list(self):
        doc = self.doc_with_bibliography("Список литературы", ["Иванов И.И. Книга. – М., 2020. – 100 с."], style="List Number")
        self.assertTrue(any("[99]" in m for m in failures(V.check_cross_references(doc), "error")))

    def test_after_toc_update_body_paragraph_is_not_a_bibliography_entry(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Document(str(fixture("gen_project_word.docx")))
            anchor = next(p for p in source.paragraphs if p.text.startswith("Сравнение существующих решений"))
            new = anchor.insert_paragraph_before("25. Пятым этапом становится анализ журнала ошибок [25, с. 7].")
            new.style = anchor.style
            path = Path(tmp) / "xr.docx"
            source.save(path)
            report = V.validate_vkr(str(path), profile="mpgu-09-project")
            self.assertTrue(any("[25]" in m for m in report_messages(report)), report_messages(report))

    def test_citation_to_missing_source_after_word_toc_update(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "badref.docx"
            rewrite_part(fixture("gen_project_word.docx"), path, "word/document.xml", [("[6]", "[99]")])
            report = V.validate_vkr(str(path), profile="mpgu-09-project")
            errors = report_messages(report)
            self.assertTrue(any("[99]" in m and "НЕТ в списке" in m for m in errors), errors)

    def test_citation_in_footnote_counts(self):
        report = V.validate_vkr(str(fixture("ok_spec8_styles.docx")), profile="mpgu-09-project")
        warnings = report_messages(report, "warning")
        self.assertFalse(any("НЕ упоминаются" in m for m in warnings), warnings)

    def test_marker_in_footnote_is_found(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "fn.docx"
            rewrite_part(fixture("ok_spec8_styles.docx"), path, "word/footnotes.xml", [("Описание API", "[ПРОВЕРИТЬ] Описание API")])
            messages = failures(V.check_unresolved_placeholders(Document(str(path))))
            self.assertTrue(any("(сноска)" in m for m in messages), messages)

    def test_alphabetical_order(self):
        unsorted = ["1. Яковлев А.А. Книга. – М., 2020.", "2. Андреев Б.Б. Книга. – М., 2019.", "3. Fowler M. Refactoring. – Boston, 2018."]
        doc = self.doc_with_bibliography("Список литературы", unsorted, text="Ссылки [1], [2], [3].")
        self.assertTrue(failures(V.check_bibliography_order(doc, "mpgu-09-project"), "error"))
        self.assertEqual([], failures(V.check_bibliography_order(doc, "generic"), "error"))
        self.assertTrue(failures(V.check_bibliography_order(doc, "generic"), "warning"))
        ordered = ["1. Андреев Б.Б. Книга. – М., 2019.", "2. Ёлкин В.В. Книга. – М., 2019.", "3. Об образовании: федер. закон. – 2012.", "4. Fowler M. Refactoring. – Boston, 2018."]
        doc = self.doc_with_bibliography("Список литературы", ordered, text="Ссылки [1], [2], [3], [4].")
        self.assertEqual([], failures(V.check_bibliography_order(doc, "mpgu-09-project")))

    def test_duplicate_and_gap_numbers(self):
        doc = self.doc_with_bibliography("Список литературы", ["1. Андреев А.А. Книга. – М., 2020.", "3. Борисов Б.Б. Книга. – М., 2020.", "3. Васильев В.В. Книга. – М., 2020."], text="[1], [3].")
        messages = failures(V.check_bibliography_order(doc, "mpgu-09-project"), "error")
        self.assertTrue(any("повторяются номера [3]" in m for m in messages), messages)
        doc = self.doc_with_bibliography("Список литературы", ["1. Андреев А.А. Книга. – М., 2020.", "3. Борисов Б.Б. Книга. – М., 2020."], text="[1], [3].")
        messages = failures(V.check_bibliography_order(doc, "mpgu-09-project"), "error")
        self.assertTrue(any("не сквозная" in m for m in messages), messages)

    def test_typed_number_inside_autonumbered_list_duplicates_word_numbers(self):
        doc = self.doc_with_bibliography("Список литературы", ["Андреев А.А. Книга. – М., 2020.", "Борисов Б.Б. Книга. – М., 2020."], text="[1], [2], [3].", style="List Number")
        doc.add_paragraph("3. Васильев В.В. Книга. – М., 2020.")
        doc.add_paragraph("Григорьев Г.Г. Книга. – М., 2020.", style="List Number")
        messages = failures(V.check_bibliography_order(doc, "mpgu-09-project"), "error")
        self.assertTrue(any("повторяются номера [3]" in m for m in messages), messages)

    def test_subgroup_heading_inside_bibliography(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Ссылки [1] и [2].")
        add_heading(doc, "Список литературы")
        doc.add_paragraph("Нормативные правовые акты")
        doc.add_paragraph("1. Об образовании в Российской Федерации: федер. закон от 29.12.2012 № 273-ФЗ.")
        doc.add_paragraph("Литература")
        doc.add_paragraph("2. Андреев А.А. Дидактика. – М., 2020. – 214 с.")
        self.assertEqual([], failures(V.check_bibliography_order(doc, "mpgu-09-project")))
        self.assertEqual([], failures(V.check_cross_references(doc, "mpgu-09-project")))

    def test_diacritics_follow_generator_sort_key(self):
        doc = self.doc_with_bibliography("Список литературы", ["1. Mayer R. Multimedia Learning. – 2009.", "2. Müller K. Didaktik. – 2015."], text="[1], [2].")
        self.assertEqual([], failures(V.check_bibliography_order(doc, "mpgu-09-project")))


# ---------------------------------------------------------------------------
# Оглавление (R-F-04, REG-D-07, F21/P0013)
# ---------------------------------------------------------------------------

class TocTests(unittest.TestCase):
    def test_sdt_toc_is_recognized(self):
        doc = Document(str(fixture("ok_toc_sdt.docx")))
        self.assertTrue(all(item["ok"] for item in V.check_structure(doc)), V.check_structure(doc))
        results = V.check_toc_filled(doc, "mpgu-09-project")
        self.assertTrue(any(item["ok"] and "Оглавление заполнено" in item["message"] for item in results), results)

    def test_right_click_phrase_outside_toc_is_not_placeholder(self):
        doc = Document(str(fixture("ok_spec8_styles.docx")))
        self.assertTrue(any("щёлкните правой" in p.text for p in doc.paragraphs))
        self.assertEqual([], failures(V.check_toc_filled(doc, "mpgu-09-project")))

    def test_generator_placeholder_is_still_error(self):
        doc = new_doc()
        doc.add_paragraph("Содержание")
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        for kind, text in (("begin", None), (None, r'TOC \o "1-3" \h \z \u'), ("separate", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText")
                el.text = text
            run._r.append(el)
        placeholder = paragraph.add_run("[Для обновления содержания: щёлкните правой кнопкой → Обновить поле]")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        placeholder._r.append(end)
        add_heading(doc, "Введение")
        self.assertTrue(any("НЕ обновлено" in m for m in failures(V.check_toc_filled(doc), "error")))
        self.assertTrue(any("Для обновления содержания" in m for m in failures(V.check_unresolved_placeholders(doc), "error")))

    def test_annotation_in_toc_is_reported(self):
        doc = new_doc()
        doc.add_paragraph("Содержание")
        for line in ("Аннотация\t2", "Введение\t4", "Глава I. Теория\t5", "Заключение\t20"):
            doc.add_paragraph(line, style="TOC 1" if "TOC 1" in [s.name for s in doc.styles] else None)
        add_heading(doc, "Введение")
        results = V.check_toc_filled(doc, "mpgu-09-project")
        self.assertTrue(any("внесены: Аннотация" in m for m in failures(results, "error")), results)


# ---------------------------------------------------------------------------
# Главы и заголовки (REG-D-08, R2-10, R2-11, R2-05, R2-16)
# ---------------------------------------------------------------------------

class ChapterTests(unittest.TestCase):
    def mechanics_errors(self, doc, profile="mpgu-09-project"):
        return failures(V.check_document_mechanics(doc, profile), "error")

    def test_roman_numbering_variants_are_correct(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        add_heading(doc, "Глава I Теоретические основы")
        heading = add_heading(doc, "ГЛАВА II")
        heading.runs[0].add_break()
        heading.add_run("Анализ условий внедрения")
        add_heading(doc, "Глава III.")
        add_heading(doc, "Проектирование", level=2)
        messages = self.mechanics_errors(doc)
        self.assertFalse(any("римскими" in m or "Точка в конце" in m for m in messages), messages)
        profile = V.check_mpgu_09_profile(doc, "mpgu-09-project")
        self.assertTrue(any(item["ok"] and "глав: 3" in item["message"] for item in profile), profile)

    def test_short_body_paragraph_starting_with_chapter_is_not_a_chapter(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Глава II посвящена анализу условий внедрения.")
        for number in ("I", "II", "III"):
            add_heading(doc, f"Глава {number}. Раздел")
        doc.add_paragraph("Пилотное тестирование проведено.")
        results = V.check_mpgu_09_profile(doc, "mpgu-09-project")
        self.assertTrue(all(item["ok"] for item in results), results)

    def test_arabic_chapter_in_bold_normal_paragraph_is_error_for_mpgu(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        paragraph = doc.add_paragraph()
        paragraph.add_run("Глава 1. Теоретическая часть").bold = True
        self.assertTrue(any("не римскими" in m for m in self.mechanics_errors(doc)))
        self.assertFalse(any("не римскими" in m for m in self.mechanics_errors(doc, "generic")))

    def test_heading_ending_with_abbreviation(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        add_heading(doc, "2.2. Анализ решений: тренажёры, онлайн-курсы и т.д.", level=2)
        add_heading(doc, "1.1. Визуальные среды: Scratch, Blockly и др.", level=2)
        add_heading(doc, "1.2. Что такое тренажёр?", level=2)
        add_heading(doc, "1.3. Развитие идей…", level=2)
        self.assertFalse(any("Точка в конце" in m for m in self.mechanics_errors(doc)))
        add_heading(doc, "1.4. Обычный заголовок.", level=2)
        self.assertTrue(any("Точка в конце" in m for m in self.mechanics_errors(doc)))

    def test_structural_word_as_body_text_is_not_a_heading(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Текст введения.")
        add_heading(doc, "Заключение")
        body = doc.add_paragraph("Заключение.")
        body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self.assertFalse(any("Точка в конце" in m for m in self.mechanics_errors(doc)))

    def test_duplicate_chapter_prefix(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        add_heading(doc, "Глава I. Глава I. Теоретические основы")
        self.assertTrue(any("Префикс главы повторяется" in m for m in self.mechanics_errors(doc, "generic")))


# ---------------------------------------------------------------------------
# Мелкие ложные срабатывания (REG-D-09, REG-D-14, R-F-07, R-F-08, F8, F9, #17)
# ---------------------------------------------------------------------------

class SmallFalsePositiveTests(unittest.TestCase):
    def test_landscape_a4_section(self):
        doc = Document(str(fixture("ok_spec8_styles.docx")))
        self.assertGreater(doc.sections[-1].page_width, doc.sections[-1].page_height)
        self.assertEqual([], failures(V.check_page_size(doc)))
        self.assertEqual([], failures(V.check_margins(doc)))

    def test_sentence_starting_with_table_is_a_mention(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Таблица 1 показывает, что разработанный продукт выигрывает у аналогов.")
        doc.add_paragraph("Таблица 1")
        doc.add_table(rows=1, cols=2)
        self.assertEqual([], failures(V.check_figure_table_references(doc)))
        tables = V.check_tables(doc)
        self.assertTrue(any(item["ok"] and "1 таблиц и 1 меток" in item["message"] for item in tables), tables)

    def test_full_name_on_title_page_is_not_a_surname_warning(self):
        report = V.validate_vkr(str(fixture("gen_project_word.docx")), profile="mpgu-09-project")
        self.assertEqual([], failures(report["checks"]["Оформление фамилий"]))
        doc = Document(str(fixture("gen_project_word.docx")))
        self.assertTrue(any("Иванова Мария Петровна" in p.text for p in doc.paragraphs))
        self.assertTrue(any("Петров Сергей Викторович" in p.text for p in doc.paragraphs))

    def test_annotation_goal_forms_and_keywords_message(self):
        doc = new_doc()
        doc.add_paragraph("Аннотация")
        doc.add_paragraph(
            "Работа «Тренажёр» изложена на 45 страницах, список включает 20 источников. Объектом исследования "
            "является обучение. Целью работы является разработка тренажёра. Актуальность связана с практикой."
        )
        add_heading(doc, "Введение")
        messages = failures(V.check_annotation_content(doc))
        self.assertFalse(any("«цель»" in m for m in messages), messages)
        self.assertTrue(any("перечня ключевых слов" in m for m in messages), messages)
        self.assertFalse(any("5-7" in m for m in messages))

    def test_surnames_with_yo_and_initials_after(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Подход развивали в работах, где Воробьёв описал модель, а Королёва А.А. уточнила её.")
        messages = failures(V.check_surnames_format(doc))
        self.assertTrue(messages)
        self.assertIn("Воробьёв", messages[0])
        self.assertIn("Королёва + инициалы после фамилии", messages[0])

    def test_capitalized_word_at_sentence_start_is_not_a_surname(self):  # B7 (F15)
        doc = new_doc()
        add_heading(doc, "Глава III. Разработка")
        doc.add_paragraph("Рисунок 1 — Архитектура веб-тренажёра «Шаг сортировки»")
        doc.add_paragraph("Клиентская часть — одностраничное приложение на React и TypeScript, серверная часть — API на Node.js.")
        doc.add_paragraph("Модуль проверяет каждый шаг. Клиентская часть получает ответ за 200 мс! Российская школа…")
        doc.add_paragraph("Итоги: «Клиентская часть» и серверная часть разделены.")
        self.assertEqual([], failures(V.check_surnames_format(doc)))

    def test_sentence_start_word_used_as_surname_elsewhere_is_checked(self):  # B7 (F15)
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Модель поэтапного формирования описана в работе И.И. Петрова. Петров выделяет четыре этапа.")
        doc.add_paragraph("Смирнов предложил шкалу оценки, а позже работы Смирнова развили её.")
        messages = failures(V.check_surnames_format(doc))
        self.assertTrue(messages)
        self.assertIn("Петров", messages[0])
        self.assertIn("Смирнов", messages[0])  # в начале абзаца, но в середине предложения — «работы Смирнова»

    def test_pronouns_in_bibliography_titles_are_ignored(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        doc.add_paragraph("Проведён анализ источников [1].")
        add_heading(doc, "Список литературы")
        doc.add_paragraph("1. Наша новая школа: национальная инициатива. – М., 2010.")
        self.assertEqual([], failures(V.check_personal_pronouns(doc)))

    def test_trailing_empty_page(self):
        results = V.check_document_mechanics(Document(str(fixture("bad_trailing_empty_page.docx"))), "generic")
        self.assertTrue(any("пустая страница" in m for m in failures(results, "error")))
        results = V.check_document_mechanics(Document(str(fixture("gen_project_word.docx"))), "generic")
        self.assertFalse(any("пустая страница" in m for m in failures(results)))

    def test_paragraph_gaps_warning(self):
        doc = new_doc()
        add_heading(doc, "Введение")
        for _ in range(5):
            doc.add_paragraph("Абзац основного текста с интервалом после. " * 3).paragraph_format.space_after = Pt(10)
        self.assertTrue(any("Лишние пробелы" in m for m in failures(V.check_line_spacing(doc), "warning")))


# ---------------------------------------------------------------------------
# cliche_allowlist (B6, находка скептика B, E-verify: YAML без отступа)
# ---------------------------------------------------------------------------

class AllowlistTests(unittest.TestCase):
    def test_project_root_boundary_and_yaml_without_indent(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "vkr-state.md").write_text("cliche_allowlist:\n  - чужая фраза\n", encoding="utf-8")
            project = root / "project"
            (project / "final").mkdir(parents=True)
            (project / "vkr-project.json").write_text("{}", encoding="utf-8")
            docx_path = project / "final" / "vkr.docx"
            docx_path.write_bytes(b"")
            self.assertEqual([], V._load_cliche_allowlist(docx_path))
            (project / "vkr-state.md").write_text("cliche_allowlist:\n- \"стоит отметить\"\n- в современном мире\n", encoding="utf-8")
            self.assertEqual(["стоит отметить", "в современном мире"], V._load_cliche_allowlist(docx_path))

    def test_import_does_not_read_current_directory(self):
        self.assertEqual([], load_validator().CLICHE_ALLOWLIST)


# ---------------------------------------------------------------------------
# references/mpgu-formatting.md ↔ поведение валидатора по профилям (notes WP3 п. 4)
# ---------------------------------------------------------------------------

class FormattingReferenceTests(unittest.TestCase):
    DOC = ROOT / "vkr-mpgu" / "references" / "mpgu-formatting.md"
    SEVERITY = {"ошибка": "error", "предупреждение": "warning"}

    def severity_table(self):
        text = self.DOC.read_text(encoding="utf-8")
        section = text.split("Серьёзность проверок валидатора зависит от профиля", 1)[1]
        rows = []
        for line in section.splitlines():
            if not line.startswith("|"):
                if rows:
                    break
                continue
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if cells[0] == "Нарушение" or set(cells[0]) <= set("-: "):
                continue
            rows.append(cells)
        return rows

    def scenarios(self, folder):
        restart = Document(str(fixture("ok_title_section.docx")))
        sect = restart.sections[1]._sectPr
        pg = sect.find(qn("w:pgNumType"))
        if pg is None:
            pg = OxmlElement("w:pgNumType")
            sect.insert(0, pg)
        pg.set(qn("w:start"), "1")
        restart.save(str(folder / "restart.docx"))
        # Разрыв раздела, после которого Word оставил «Особый колонтитул для первой страницы».
        first = Document(str(fixture("ok_title_section.docx")))
        first.sections[1].different_first_page_header_footer = True
        first.save(str(folder / "first.docx"))
        return {
            "арабскими": (fixture("bad_arabic_chapter.docx"), "не римскими цифрами"),
            "алфавитном": (fixture("bad_bibliography_unsorted.docx"), "не в алфавитном порядке"),
            "заново": (folder / "restart.docx", "нумерация начинается заново"),
            "титульном": (fixture("bad_title_page_number.docx"), "На титульном листе выводится номер"),
            "без номера": (folder / "first.docx", "первая страница секции останется без номера"),
            "Times New Roman": (fixture("bad_normal_font_arial.docx"), "Заголовки набраны шрифтом"),
        }

    def test_profile_severity_table_matches_validator(self):
        rows = self.severity_table()
        self.assertEqual(6, len(rows), rows)
        with tempfile.TemporaryDirectory() as temporary:
            scenarios = self.scenarios(Path(temporary))
            for label, mpgu, generic in rows:
                keys = [key for key in scenarios if key in label]
                self.assertEqual(1, len(keys), label)
                path, fragment = scenarios.pop(keys[0])
                for profile, expected in (("mpgu-09-project", mpgu), ("generic", generic)):
                    with self.subTest(row=label, profile=profile):
                        report = V.validate_vkr(str(path), profile=profile)
                        found = {level for level in ("error", "warning") if any(fragment in m for m in report_messages(report, level))}
                        self.assertEqual({self.SEVERITY[expected]}, found)
            self.assertEqual({}, scenarios)

    def test_section_break_advice_uses_validator_wording(self):
        text = self.DOC.read_text(encoding="utf-8")
        advice = " ".join(text.split("**Разрыв раздела в Word**", 1)[1].split("\n\n", 1)[0].split())
        source = VALIDATOR_PATH.read_text(encoding="utf-8")
        for option in ("«Особый колонтитул для первой страницы»", "«Формат номеров страниц»", "«продолжить»"):
            self.assertIn(option, advice)
            self.assertIn(option, source)


if __name__ == "__main__":
    unittest.main(verbosity=2)
