#!/usr/bin/env python3
"""Входные DOCX для фикстур валидатора (до обработки в Word).

1. gen_project*.docx — генератор create_vkr_docx.py (снимок 6.32) из небольшого JSON.
2. kitchen.docx — python-docx документ со стилями SPEC 8 (VKR Title, VKR Listing,
   VKR Caption, VKR Bibliography, VKR Appendix Label), автонумерованным списком
   литературы и спорными, но правильными конструкциями из находок аудита.

Запуск: python make_inputs.py <каталог_генератора> <каталог_вывода>
Дальше word_fixtures.ps1 открывает их в Word, обновляет поля и сохраняет.
"""
from __future__ import annotations

import copy
import importlib.util
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


INTRO = (
    "Цифровые тренажёры закрепились в школьной практике как средство регулярной "
    "отработки умений, однако большинство доступных решений ориентировано на старшие "
    "классы и не учитывает типичные ошибки обучающихся 8 класса [2, с. 14]. "
    "Л.С. Выготский связывал развитие мышления с интериоризацией внешних действий [2, с. 112–114].\n\n"
    "Объект исследования — процесс обучения алгоритмизации в основной школе. Предмет "
    "исследования — цифровой тренажёр как средство отработки умений. Цель работы — "
    "разработка и пилотное тестирование тренажёра. Работа опирается на идеи "
    "В.В. Давыдова и др. [3], а также на нормы, принятые в 2012 г. [4].\n\n"
    "Глава II посвящена анализу условий внедрения."
)

CH1_TEXT = (
    "П.Я. Гальперин описал поэтапное формирование умственных действий, и эта модель "
    "объясняет последовательность заданий тренажёра (см. [2, с. 40]). Сходные данные "
    "приводят А.А. Андреев и др. [1, с. 12]. Результаты сравнения функций приведены в таблице 1.\n\n"
    "Типичные ошибки обучающихся связаны с индексами массивов и границами циклов, т. е. [5, p. 7] "
    "с теми элементами, которые требуют точного пошагового исполнения алгоритма."
)

LIST_ITEMS = [
    "закрепление: серии однотипных задач с постепенным усложнением условий и немедленной обратной связью по каждому шагу решения;",
    "диагностика: журнал ошибок по темам курса, который учитель просматривает перед уроком и использует для планирования повторения;",
    "мотивация: наглядная шкала прогресса без соревновательного рейтинга, чтобы слабые обучающиеся не теряли интерес к заданиям.",
]

LISTING = (
    "def update_history(history, matrix, n, x):\n"
    "    # TODO: учитывать порядок строк\n"
    "    history[x] = matrix[n][30]\n"
    "    template = \"<span>{{ user.name }}</span>\"\n"
    "    return data[25] * weight[n]"
)

CH2_TEXT = (
    "Целевая аудитория продукта — обучающиеся 8 класса и учителя информатики. "
    "Анкетирование 46 обучающихся показало, что большинство выполняет домашние задания "
    "с телефона [5, p. 12]. Учебный процесс в образовательной организации построен по "
    "методическим рекомендациям, опубликованным в 2021 г. [4].\n\n"
    "Сравнение существующих решений проводилось по критериям удобства, предложенным "
    "в работе [6], и по педагогическим критериям [3; 1]."
)

CH3_TEXT = (
    "Проектная часть включает разработку содержания тренажёра, пилотное тестирование "
    "и анализ полученных данных. Пилотное тестирование проведено с 12 обучающимися "
    "8 класса в течение двух недель [Цит. по: 1, с. 235].\n\n"
    "По итогам пилота доля верно решённых задач выросла с 54 до 71 %, а число "
    "обращений к подсказкам снизилось [5, с. 23–25]."
)

BIBLIOGRAPHY_SORTED = [
    "1. Андреев А.А. Дидактика цифрового обучения. – М.: Просвещение, 2020. – 214 с.",
    "2. Выготский Л.С. Мышление и речь. – М.: Лабиринт, 1999. – 352 с.",
    "3. Давыдов В.В. Теория развивающего обучения. – М.: Интор, 1996. – 544 с.",
    "4. Об образовании в Российской Федерации: федер. закон от 29.12.2012 № 273-ФЗ. – Доступ из справ.-правовой системы «КонсультантПлюс».",
    "5. Шмелёв А.Г. Психодиагностика в образовании. – М.: Юрайт, 2021. – 318 с.",
    "6. Nielsen J. Usability Engineering. – San Francisco: Morgan Kaufmann, 1993. – 362 p.",
]


def project_data(bibliography, skip_last_page=False):
    return {
        "title": "Разработка цифрового тренажёра по алгоритмизации для обучающихся 8 класса",
        "title_page": {
            "author": "Иванова Мария Петровна",
            "institute": "Институт математики и информатики",
            "direction_code": "09.03.02",
            "direction_name": "Информационные системы и технологии",
            "supervisor_position": "доцент",
            "supervisor_name": "Петров Сергей Викторович",
            "department": "информатики",
            "head_name": "Сидоров Олег Павлович",
            "year": 2026,
        },
        "annotation": (
            "Выпускная квалификационная работа «Разработка цифрового тренажёра по "
            "алгоритмизации для обучающихся 8 класса» изложена на 45 страницах. "
            "Список литературы включает 6 источников. Объект исследования — процесс "
            "обучения алгоритмизации. Целью работы является разработка тренажёра. "
            "Актуальность работы связана с нехваткой тренажёров для основной школы. "
            "Ключевые слова: цифровой тренажёр, алгоритмизация, пилотное тестирование."
        ),
        "introduction": INTRO,
        "chapters": [
            {
                "title": "Теоретические основы применения тренажёров в обучении",
                "paragraphs": [
                    {
                        "title": "1.1. Функции цифрового тренажёра, виды заданий и т.д.",
                        "blocks": [
                            {"type": "text", "content": CH1_TEXT},
                            {"type": "list", "ordered": False, "items": LIST_ITEMS},
                            {
                                "type": "table", "number": 1,
                                "title": "Дидактические функции цифрового тренажёра",
                                "headers": ["Функция", "Проявление в тренажёре"],
                                "rows": [["Закрепление", "Серии однотипных задач"],
                                         ["Диагностика", "Журнал ошибок по темам"]],
                                "source": "составлено автором",
                            },
                        ],
                    }
                ],
                "conclusion": "Цифровой тренажёр выполняет функции закрепления, диагностики и мотивации.",
            },
            {
                "title": "Анализ условий внедрения тренажёра в образовательной организации",
                "paragraphs": [
                    {
                        "title": "2.1. Характеристика целевой аудитории",
                        "blocks": [
                            {"type": "text", "content": CH2_TEXT},
                            {"type": "listing", "number": 1, "caption": "Обработка истории попыток", "code": LISTING},
                        ],
                    }
                ],
                "conclusion": "Анализ подтвердил потребность в мобильном тренажёре с журналом ошибок.",
            },
            {
                "title": "Проектирование и пилотное тестирование тренажёра",
                "paragraphs": [
                    {"title": "3.1. Пилотное тестирование", "blocks": [{"type": "text", "content": CH3_TEXT}]}
                ],
                "conclusion": "Пилотное тестирование подтвердило применимость тренажёра в учебном процессе.",
            },
        ],
        "conclusion": (
            "В работе решены поставленные задачи. Раскрыты дидактические функции цифрового "
            "тренажёра и условия его внедрения в образовательной организации.\n\n"
            "Разработан тренажёр и проведено пилотное тестирование с обучающимися 8 класса. "
            "Результаты пилота показали рост доли верно решённых задач."
        ),
        "bibliography": list(bibliography),
        "appendices": [
            {
                "label": "Приложение 1",
                "title": "Анкета для обучающихся",
                "content": (
                    "Каким устройством обучающийся обычно пользуется для выполнения домашних заданий?\n\n"
                    "Чтобы открыть меню тренажёра, щёлкните правой кнопкой мыши по значку модуля."
                ),
            }
        ],
        "skip_last_page": skip_last_page,
    }


def load_generator(gen_dir: Path):
    sys.path.insert(0, str(gen_dir))
    spec = importlib.util.spec_from_file_location("create_vkr_docx_632", gen_dir / "create_vkr_docx.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


# ---------- kitchen.docx: стили SPEC 8 и спорные правильные конструкции ----------

def paragraph_style(doc, name, base="Normal"):
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    return style


def set_fonts(target, name):
    target.font.name = name
    rpr = target.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(key), name)
    for key in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(key), None)


def add_numpr(paragraph, num_id, ilvl=0):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    lvl = OxmlElement("w:ilvl")
    lvl.set(qn("w:val"), str(ilvl))
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.append(lvl)
    numpr.append(nid)
    pstyle = ppr.find(qn("w:pStyle"))
    if pstyle is None:
        ppr.insert(0, numpr)
    else:
        pstyle.addnext(numpr)


def add_decimal_numbering(doc):
    """Добавляет абстрактную нумерацию «1.» и возвращает numId."""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids + [0]) + 1
    num_id = max(num_ids + [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:lvlText", "%1."), ("w:lvlJc", "left")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        lvl.append(el)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:firstLine"), "709")
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids + [0]) + 1
    num_id = max(num_ids + [0]) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "–"), ("w:lvlJc", "left")):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        lvl.append(el)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    lvl.append(ppr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def add_toc_field(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-2" \h \z \u '
    run._r.append(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    run2 = p.add_run("[Для обновления содержания: щёлкните правой кнопкой → Обновить поле]")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run2._r.append(end)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_break(__import__("docx.enum.text", fromlist=["WD_BREAK"]).WD_BREAK.PAGE)
    return p


def build_kitchen(path: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    section.left_margin, section.right_margin = Mm(35), Mm(10)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    set_fonts(normal, "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    for name in ("Heading 1", "Heading 2"):
        h = doc.styles[name]
        set_fonts(h, "Times New Roman")
        h.font.size = Pt(14)
        h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h.paragraph_format.space_before = Pt(0)
        h.paragraph_format.space_after = Pt(42)
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    doc.styles["Heading 2"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.styles["Heading 2"].paragraph_format.first_line_indent = Cm(1.25)

    title = paragraph_style(doc, "VKR Title")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    listing = paragraph_style(doc, "VKR Listing")
    set_fonts(listing, "Consolas")
    listing.font.size = Pt(11)
    listing.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    listing.paragraph_format.first_line_indent = Cm(0)
    listing.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    caption = paragraph_style(doc, "VKR Caption")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    bibliography = paragraph_style(doc, "VKR Bibliography")
    bibliography.paragraph_format.first_line_indent = Cm(1.25)
    appendix_label = paragraph_style(doc, "VKR Appendix Label")
    appendix_label.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    appendix_label.paragraph_format.first_line_indent = Cm(0)
    appendix_label.paragraph_format.page_break_before = True

    bullet_id = add_bullet_numbering(doc)
    decimal_id = add_decimal_numbering(doc)

    for text in (
        "Министерство просвещения Российской Федерации",
        "«Московский педагогический государственный университет»",
        "Иванова Мария Петровна",
        "Разработка цифрового тренажёра по алгоритмизации",
        "Выпускная квалификационная работа",
    ):
        doc.add_paragraph(text, style="VKR Title")
    p = doc.add_paragraph("Научный руководитель – доцент кафедры информатики, кандидат педагогических наук Петров Сергей Викторович", style="VKR Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph("Москва – 2026 год", style="VKR Title")
    page_break(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(42)
    p.add_run("Аннотация").bold = True
    doc.add_paragraph(
        "Выпускная квалификационная работа «Разработка цифрового тренажёра по алгоритмизации» "
        "изложена на 45 страницах, список литературы включает 6 источников. Объектом исследования "
        "является процесс обучения алгоритмизации. Целью работы является разработка тренажёра "
        "и его пилотное тестирование. Актуальность работы определяется нехваткой тренажёров "
        "для основной школы. Ключевые слова: тренажёр, алгоритмизация, пилот."
    )
    page_break(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(42)
    p.add_run("Содержание").bold = True
    add_toc_field(doc)

    doc.add_paragraph("Введение", style="Heading 1")
    doc.add_paragraph(
        "Цифровые тренажёры применяются в школьной практике для регулярной отработки умений "
        "[2, с. 14]. Л.С. Выготский связывал развитие мышления с интериоризацией внешних "
        "действий, что подтверждают и более поздние исследования, выполненные в 2019 г. [3]."
    )
    doc.add_paragraph("Глава II посвящена анализу условий внедрения.")
    doc.add_paragraph(
        "Объект исследования — процесс обучения алгоритмизации. Цель работы — разработка "
        "тренажёра для обучающихся 8 класса и проверка его в образовательной организации."
    )

    doc.add_paragraph("Глава I Теоретические основы применения тренажёров", style="Heading 1")
    doc.add_paragraph("1.1. Функции тренажёра, виды заданий и т.д.", style="Heading 2")
    doc.add_paragraph(
        "Сходные данные приводят А.Н. Леонтьев и др. [4, с. 12]. Модель поэтапного формирования "
        "действий объясняет последовательность заданий (см. [1, с. 5]). Ошибки связаны с индексами "
        "массивов, т. е. [6, p. 7] с элементами пошагового исполнения, и с границами циклов [5, с. 112–114]."
    )
    doc.add_paragraph(
        "Цитата «…и так далее…» [2] показывает, что отработка без обратной связи малоэффективна; "
        "эту позицию разделяют и другие авторы [1; 5], а также методисты [Цит. по: 3, с. 235]."
    )
    for item in (
        "закрепление: серии однотипных задач с постепенным усложнением условий и немедленной обратной связью;",
        "диагностика: журнал ошибок по темам курса, который учитель просматривает перед уроком;",
    ):
        lp = doc.add_paragraph(item, style="List Paragraph")
        add_numpr(lp, bullet_id)
        lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lp.paragraph_format.first_line_indent = Cm(0)
    doc.add_paragraph(
        "Таблица 1 показывает, что тренажёр закрывает функции закрепления и диагностики, "
        "которые не поддерживает ни одно из рассмотренных решений."
    )
    label = doc.add_paragraph("Таблица 1", style="VKR Caption")
    label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cap = doc.add_paragraph(style="VKR Caption")
    cap.add_run("Функции цифрового тренажёра").bold = True
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for row, values in zip(table.rows, (("Функция", "Проявление"), ("Диагностика", "Журнал ошибок"))):
        for cell, value in zip(row.cells, values):
            cell.text = value
    for line in (
        "def update_history(history, matrix, n, x):",
        "    # TODO: учитывать порядок строк",
        "    history[x] = matrix[n][30]",
        "    return \"<li>{{ task.title }}</li>\"",
    ):
        doc.add_paragraph(line, style="VKR Listing")
    doc.add_paragraph("Листинг 1 — Обработка истории попыток", style="VKR Caption")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Выводы по Главе I").bold = True
    doc.add_paragraph("Цифровой тренажёр выполняет функции закрепления, диагностики и мотивации обучающихся.")

    h = doc.add_paragraph(style="Heading 1")
    h.add_run("ГЛАВА II")
    h.add_run().add_break()
    h.add_run("Анализ условий внедрения тренажёра")
    doc.add_paragraph(
        "Целевая аудитория продукта — обучающиеся 8 класса и учителя информатики. Учебный процесс "
        "организован по методическим рекомендациям FOOTNOTE_HERE, принятым в образовательной организации."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Выводы по Главе II").bold = True
    doc.add_paragraph("Анализ подтвердил потребность в мобильном тренажёре с журналом ошибок и подсказками.")

    doc.add_paragraph("Глава III.", style="Heading 1")
    doc.add_paragraph("Проектирование и пилотное тестирование", style="Heading 2")
    doc.add_paragraph(
        "Педагогическая проверка включала пилотное тестирование с 12 обучающимися в течение двух "
        "недель; методическое сопровождение обеспечивал учитель информатики [5, с. 23–25]."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run("Выводы по Главе III").bold = True
    doc.add_paragraph("Пилотное тестирование подтвердило применимость тренажёра в учебном процессе.")

    doc.add_paragraph("Заключение", style="Heading 1")
    doc.add_paragraph(
        "В работе раскрыты функции цифрового тренажёра, проанализированы условия его внедрения "
        "и проведено пилотное тестирование, подтвердившее рост доли верно решённых задач."
    )

    doc.add_paragraph("Список использованных источников", style="Heading 1")
    for entry in (
        "Андреев А.А. Дидактика цифрового обучения. – М.: Просвещение, 2020. – 214 с.",
        "Выготский Л.С. Мышление и речь. – М.: Лабиринт, 1999. – 352 с.",
        "Давыдов В.В. Теория развивающего обучения. – М.: Интор, 1996. – 544 с.",
        "Леонтьев А.Н. Деятельность. Сознание. Личность. – М.: Смысл, 2005. – 352 с.",
        "Шмелёв А.Г. Психодиагностика в образовании. – М.: Юрайт, 2021. – 318 с.",
        "Nielsen J. Usability Engineering. – San Francisco: Morgan Kaufmann, 1993. – 362 p.",
    ):
        bp = doc.add_paragraph(entry, style="VKR Bibliography")
        add_numpr(bp, decimal_id)

    doc.add_paragraph("Приложение 1", style="VKR Appendix Label")
    doc.add_paragraph("Анкета для обучающихся", style="Heading 2").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Чтобы открыть меню тренажёра, щёлкните правой кнопкой мыши по значку модуля и выберите пункт «Настройки»."
    )
    doc.add_paragraph("Приложение 2", style="VKR Appendix Label")
    doc.add_paragraph("Фрагмент шаблона страницы", style="Heading 2").alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in ("<ul>", "  <li v-for=\"task in tasks\">{{ task.title }} [x]</li>", "</ul>"):
        doc.add_paragraph(line, style="VKR Listing")
    doc.save(path)


KEEP_STYLE_NAMES = {
    "normal", "heading 1", "heading 2", "heading 3", "toc 1", "toc 2", "toc 3", "toc heading",
    "caption", "list paragraph", "list bullet", "list number", "hyperlink", "footnote text",
    "footnote reference", "header", "footer", "table grid", "title", "default paragraph font",
    "normal table", "no list",
}


def prune_styles(path: Path):
    """Удаляет неиспользуемые стили шаблона python-docx, чтобы фикстуры после Word были небольшими."""
    import re
    import zipfile

    from lxml import etree

    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as archive:
        entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
    used = set()
    for info, data in entries:
        if info.filename.startswith("word/") and info.filename.endswith(".xml") and "styles" not in info.filename:
            used.update(re.findall(rb'w:(?:pStyle|rStyle|tblStyle|numStyleLink|styleLink) w:val="([^"]+)"', data))
    used = {u.decode("utf-8") for u in used}
    styles_data = dict((info.filename, data) for info, data in entries)["word/styles.xml"]
    root = etree.fromstring(styles_data)
    by_id = {el.get(w + "styleId"): el for el in root.findall(w + "style")}
    keep = set(used)
    for sid, el in by_id.items():
        name = el.find(w + "name")
        if el.get(w + "default") in ("1", "true") or (name is not None and name.get(w + "val", "").casefold() in KEEP_STYLE_NAMES):
            keep.add(sid)
    changed = True
    while changed:
        changed = False
        for sid in list(keep):
            el = by_id.get(sid)
            if el is None:
                continue
            for tag in ("basedOn", "link", "next"):
                ref = el.find(w + tag)
                if ref is not None and ref.get(w + "val") not in keep:
                    keep.add(ref.get(w + "val"))
                    changed = True
    for sid, el in by_id.items():
        if sid not in keep:
            root.remove(el)
    new_styles = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in entries:
            if info.filename in ("word/stylesWithEffects.xml", "docProps/thumbnail.jpeg"):
                continue
            if info.filename == "word/styles.xml":
                data = new_styles
            if info.filename == "[Content_Types].xml":
                data = re.sub(rb'<Override[^>]+PartName="/(?:word/stylesWithEffects\.xml|docProps/thumbnail\.jpeg)"[^>]*/>', b"", data)
            if info.filename.endswith(".rels"):
                data = re.sub(rb'<Relationship[^>]+Target="(?:stylesWithEffects\.xml|docProps/thumbnail\.jpeg)"[^>]*/>', b"", data)
            archive.writestr(info.filename, data)


def main():
    gen_dir = Path(sys.argv[1]).resolve()
    out = Path(sys.argv[2]).resolve()
    out.mkdir(parents=True, exist_ok=True)
    generator = load_generator(gen_dir)
    generator.build_vkr_docx(project_data(BIBLIOGRAPHY_SORTED), out / "gen_project.docx")
    unsorted = [BIBLIOGRAPHY_SORTED[4], BIBLIOGRAPHY_SORTED[1], BIBLIOGRAPHY_SORTED[0],
                BIBLIOGRAPHY_SORTED[5], BIBLIOGRAPHY_SORTED[2], BIBLIOGRAPHY_SORTED[3]]
    # номера остаются 1..6 по порядку, меняются только тексты записей → порядок не алфавитный
    unsorted = [f"{i}. {entry.split('. ', 1)[1]}" for i, entry in enumerate(unsorted, 1)]
    data = project_data(BIBLIOGRAPHY_SORTED)
    data["bibliography"] = unsorted
    generator.build_vkr_docx(data, out / "gen_project_unsorted_bib.docx")
    generator.build_vkr_docx(project_data(BIBLIOGRAPHY_SORTED, skip_last_page=True), out / "gen_project_skiplast.docx")
    build_kitchen(out / "kitchen.docx")
    for name in ("gen_project.docx", "gen_project_unsorted_bib.docx", "gen_project_skiplast.docx", "kitchen.docx"):
        prune_styles(out / name)
    (out / "project_small.json").write_text(
        json.dumps(project_data(BIBLIOGRAPHY_SORTED), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("ok", sorted(p.name for p in out.glob("*.docx")))


if __name__ == "__main__":
    main()
