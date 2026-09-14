"""Регрессии 6.32, сохранённые в контракте 6.33.

Класс ``DoctorRegressionTests`` 6.32 строил ``audit/manifest.json`` и отчёты
аудиторов вручную. В 6.33 такие записи doctor показывает как
``AUDIT_LEGACY_RUNS`` и не засчитывает, а служебные записи пишет
``vkr_audit.py``. Те же сценарии (устаревший снимок, одна финальная волна,
повтор аудиторов и execution_id, открытые BLOCKER/MAJOR, подложенные отчёты,
отчёт валидатора от другого DOCX, профиль и т. д.) покрыты через инструмент в
``tests/test_core_audit_633.py`` (классы ``FinalGateTest``, ``AuditFlowTest``,
``DoctorStageTest``). Здесь остаются проверки скриптов, не зависящие от формата
manifest, и одна проверка doctor на проекте, созданном ``init_vkr_project``.

Запуск: ``python -m unittest discover -s tests -t .`` или
``python tests/test_audit_fixes_632.py``.
"""

from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
import importlib.util
from pathlib import Path

os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
sys.dont_write_bytecode = True

from docx import Document  # noqa: E402


ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "vkr-mpgu"
SCRIPTS = SKILL / "scripts"
WORD_SAVED_FIXTURE = ROOT / "tests" / "fixtures" / "build" / "vkr-word-saved.docx"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))


def load_module(name: str):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


doctor = load_module("vkr_project_doctor")
init_project = load_module("init_vkr_project")
generator = load_module("create_vkr_docx")
bibliography = load_module("format_bibliography")
source_verifier = load_module("verify_sources")
validator = load_module("validate_vkr")
ownership = load_module("content_ownership_check")
cleaner = load_module("clean_docx_metadata")
memory = load_module("vkr_memory")


def write_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


class DoctorOnInitializedProjectTests(unittest.TestCase):
    """Doctor 6.33 на проекте из init: повреждённая группа evidence не роняет проверку."""

    def test_malformed_evidence_group_is_reported_without_crash(self):
        with tempfile.TemporaryDirectory() as td:
            base = Path(td)
            config_path = base / "intake.json"
            write_json(config_path, {
                "topic": "Информационная система сопровождения практики",
                "profile": "mpgu-09-project",
                "title_page": {"author": "Иванов Иван Иванович", "program_code": "09.03.02"},
            })
            project = base / "project"
            result = init_project.initialize(project, init_project.normalize_config(init_project.load_config(config_path)))
            self.assertEqual("initialized", result["status"])
            evidence_path = project / "evidence" / "index.json"
            evidence = json.loads(evidence_path.read_text(encoding="utf-8"))
            evidence["pilot"] = None
            write_json(evidence_path, evidence)
            report = doctor.diagnose(project, "prefinal")
            codes = {item["code"] for item in report["findings"]}
            self.assertEqual("FAIL", report["status"])
            self.assertIn("EVIDENCE_LIST_SHAPE", codes)
            self.assertNotIn("INTERNAL_ERROR", codes)


class ScriptContractTests(unittest.TestCase):
    def test_init_config_bom_boolean_and_validation(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "intake.json"
            path.write_text("\ufeff" + json.dumps({"collective": False, "mode": "express-4d"}), encoding="utf-8")
            config = init_project.normalize_config(init_project.load_config(path))
            self.assertIs(config["collective"], False)
            self.assertEqual("express-4d", config["mode"])
            self.assertEqual("balanced", config["audit_intensity"])
            # 6.33: collective — только JSON bool, строка "false" — ошибка (SPEC 2).
            path.write_text(json.dumps({"collective": "false", "mode": "express-4d"}), encoding="utf-8")
            with self.assertRaises(ValueError):
                init_project.normalize_config(init_project.load_config(path))
            path.write_text(json.dumps({"typo_profile": "x"}), encoding="utf-8")
            with self.assertRaises(ValueError):
                init_project.normalize_config(init_project.load_config(path))

    def test_bibliography_handles_numeric_pages_and_preserves_ids(self):
        sources = bibliography.normalize_input_data([
            {"id": 7, "type": "article", "authors": [{"last": "Иванов"}], "title": "Статья", "journal": "Журнал", "pages": 12},
            {"id": 2, "type": "electronic", "title": "English resource", "doi": "10.1/example", "pages": 4},
        ])
        text, mapping = bibliography.format_bibliography(sources, preserve_ids=True)
        self.assertIn("7. Иванов", text)
        # 6.33: маркер электронного ресурса всегда по-русски (gost-citations.md, «Общие правила»).
        self.assertIn("2. English resource [Электронный ресурс]", text)
        self.assertEqual({}, mapping)
        collection = bibliography.format_source({
            "type": "conference",
            "authors": [{"last": "Иванов", "initials": "И.И."}],
            "title": "Доклад",
            "container_title": "Материалы конференции",
            "place": "Москва",
            "publisher": "МПГУ",
            "year": 2024,
            "pages": 25,
            "doi": "10.1000/test",
        })
        self.assertIn("Материалы конференции", collection)
        self.assertIn("DOI: 10.1000/test", collection)

    def test_bibliography_docx_is_a4(self):
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "bibliography.docx"
            bibliography.save_as_docx("1. Источник.", output)
            section = Document(output).sections[0]
            self.assertAlmostEqual(210, section.page_width.mm, delta=0.2)
            self.assertAlmostEqual(297, section.page_height.mm, delta=0.2)

    def test_source_extraction_supports_autonumbered_list(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "source.docx"
            doc = Document()
            doc.add_paragraph("Содержаниеание")
            # A blank python-docx document has no built-in TOC styles.  A plain
            # TOC-looking row still verifies that extraction starts at the
            # bibliography heading found from the end of the document.
            doc.add_paragraph("Введение\t3")
            doc.add_paragraph("Библиографический список", style="Heading 1")
            doc.add_paragraph("Иванов И.И. Проверяемый источник. Москва, 2024.", style="List Number")
            doc.add_paragraph("Петров П.П. Второй источник. Москва, 2023.", style="List Number")
            doc.add_paragraph("Приложение 1", style="Heading 1")
            doc.save(path)
            sources = source_verifier.extract_sources_from_docx(path)
            self.assertEqual([1, 2], [item["id"] for item in sources])
            self.assertEqual(2, len(sources))

    def test_normative_source_hints_keep_real_title(self):
        order = source_verifier.parse_source_hints(
            "О федеральных государственных образовательных стандартах высшего "
            "образования [Текст]: Приказ Минобрнауки России от 12.08.2020 № 976. "
            "– Доступ из справ.-правовой системы «КонсультантПлюс»."
        )
        self.assertEqual(
            "О федеральных государственных образовательных стандартах высшего образования",
            order["title"],
        )
        self.assertEqual("normative", order["type"])

        law = source_verifier.parse_source_hints(
            "Об образовании в Российской Федерации [Текст]: федер. закон "
            "от 29.12.2012 № 273-ФЗ. – Доступ из справ.-правовой системы."
        )
        self.assertEqual("Об образовании в Российской Федерации", law["title"])
        self.assertEqual("normative", law["type"])

    def test_generator_fields_ids_appendix_and_no_service_warning(self):
        with tempfile.TemporaryDirectory() as td:
            output = Path(td) / "vkr.docx"
            data = {
                "title": "Тестовая работа",
                "title_page": {"author": "И.И. Иванов"},
                "annotation": "Аннотация " * 30,
                "introduction": "Введение " * 80,
                "chapters": [{"title": "Проверка", "paragraphs": [{"title": "1.1. Раздел", "text": "Основной текст " * 100}]}],
                "conclusion": "Заключение " * 70,
                "sources": [
                    {"id": 2, "type": "book", "title": "Бета", "year": 2024},
                    {"id": 1, "type": "book", "title": "Альфа", "year": 2023},
                ],
                "appendices": [{"title": "Материалы", "content": "Текст приложения"}],
            }
            generator.build_vkr_docx(data, output)
            doc = Document(output)
            texts = [p.text for p in doc.paragraphs]
            self.assertTrue(any(text.startswith("2. Бета") for text in texts))
            self.assertTrue(any(text.startswith("1. Альфа") for text in texts))
            self.assertNotIn("⚠️", "\n".join(texts))
            appendix_title = next(p for p in doc.paragraphs if p.text == "Материалы")
            # 6.33: заголовок приложения — стиль VKR Appendix Title, в оглавление — через поле TC.
            self.assertEqual("VKR Appendix Title", appendix_title.style.name)

            with zipfile.ZipFile(output) as archive:
                root = ET.fromstring(archive.read("word/document.xml"))
                toc_paragraph = next(p for p in root.iter() if any(n.tag.endswith("}instrText") and "TOC" in (n.text or "") for n in p.iter()))
                nodes = list(toc_paragraph.iter())
                placeholder_index = next(i for i, n in enumerate(nodes) if n.tag.endswith("}t") and "Для обновления" in (n.text or ""))
                end_index = next(i for i, n in enumerate(nodes) if n.tag.endswith("}fldChar") and n.attrib.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType") == "end")
                self.assertLess(placeholder_index, end_index)
                sect_pr = next(node for node in root.iter() if node.tag.endswith("}sectPr"))
                child_names = [node.tag.rsplit("}", 1)[-1] for node in sect_pr]
                self.assertLess(child_names.index("pgNumType"), child_names.index("cols"))

    def test_clean_metadata_in_place(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "document.docx"
            doc = Document()
            doc.core_properties.author = "Author"
            doc.add_paragraph("Text")
            doc.save(path)
            cleaner.clean_metadata(path, path, seed="test")
            cleaned = Document(path)
            self.assertEqual("", cleaned.core_properties.author)

    def test_ai_detector_empty_document_returns_nonzero(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "empty.docx"
            Document().save(path)
            result = subprocess.run(
                [sys.executable, str(SCRIPTS / "ai_detection_heuristic.py"), str(path), "--json"],
                capture_output=True, text=True, encoding="utf-8",
            )
            # 6.33: INSUFFICIENT_DATA — код 1 (analyzers-cli.md), код 2 — только ошибки ввода.
            self.assertEqual(1, result.returncode)
            self.assertEqual("INSUFFICIENT_DATA", json.loads(result.stdout)["risk"]["risk_level"])

    def test_section_matching_does_not_confuse_chapter_i_and_ii(self):
        doc = Document()
        doc.add_paragraph("Глава I. Первая", style="Heading 1")
        doc.add_paragraph("Первая глава " * 20)
        doc.add_paragraph("Глава II. Вторая", style="Heading 1")
        doc.add_paragraph("Вторая глава " * 20)
        self.assertIn("Первая глава", ownership.extract_section(doc, "chapter1")[0])
        self.assertIn("Вторая глава", ownership.extract_section(doc, "chapter2")[0])

    def test_surname_order_and_annotation_length_are_reported(self):
        doc = Document()
        doc.add_paragraph("Аннотация", style="Heading 1")
        doc.add_paragraph(
            "Работа описывает цель и объект исследования. Ключевые слова: один, два. "
            + "Подробное описание " * 90
        )
        doc.add_paragraph("Содержание", style="Heading 1")
        doc.add_paragraph("В тексте Смирнов А.А. сформулировал подход.")
        annotation = validator.check_annotation_content(doc)
        self.assertTrue(
            any(not item["ok"] and "слишком длинная" in item["message"] for item in annotation)
        )
        surnames = validator.check_surnames_format(doc)
        self.assertTrue(
            any(not item["ok"] and "после фамилии" in item["message"] for item in surnames)
        )

    def test_table_caption_does_not_count_as_a_reference(self):
        doc = Document()
        doc.add_paragraph("Таблица 1 — Результаты испытаний")
        missing = validator.check_figure_table_references(doc)
        self.assertTrue(
            any(not item["ok"] and "Таблицы ['1']" in item["message"] for item in missing)
        )

        doc.add_paragraph("Результаты представлены в таблице 1.")
        referenced = validator.check_figure_table_references(doc)
        self.assertTrue(any(item["ok"] and "Все 1" in item["message"] for item in referenced))

    def test_memory_detects_untracked_and_supports_removed_files(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td) / "project"
            for relative in (
                "memory", "logs", "audit", "drafts", "final",
                "sources", "evidence", "exports",
            ):
                (root / relative).mkdir(parents=True, exist_ok=True)
            (root / "vkr-state.md").write_text("state", encoding="utf-8")
            write_json(root / "memory/handoff.json", {})
            write_json(
                root / "memory/artifact-index.json",
                {"schema_version": memory.SCHEMA_VERSION, "artifacts": {}},
            )
            write_json(root / "audit/manifest.json", {"runs": []})
            (root / "logs/activity.jsonl").write_text("", encoding="utf-8")
            (root / "logs/tool-runs.jsonl").write_text("", encoding="utf-8")
            (root / "memory/decisions.jsonl").write_text("", encoding="utf-8")
            new_file = root / "drafts/new.md"
            new_file.write_text("new", encoding="utf-8")
            report = memory.status(root, 10)
            self.assertEqual("external_changes_detected", report["status"])
            untracked = {item["logical_id"] for item in report["artifact_changes"]["untracked"]}
            # 6.33: корневые файлы (vkr-state.md) тоже отслеживаются.
            self.assertEqual({"drafts/new.md", "vkr-state.md"}, untracked)

            # Файл события лежит вне проекта: корневые файлы проекта отслеживаются памятью.
            event = Path(td) / "event.json"
            event.write_text("\ufeff" + json.dumps({
                "event_type": "checkpoint",
                "summary": "record",
                "files": ["drafts/new.md", "vkr-state.md"],
            }), encoding="utf-8")
            memory.record(root, event)
            new_file.unlink()
            event.write_text(json.dumps({
                "event_type": "checkpoint",
                "summary": "remove",
                "removed_files": ["drafts/new.md"],
            }), encoding="utf-8")
            memory.record(root, event)
            self.assertEqual("ok", memory.status(root, 10)["status"])

    def test_memory_legacy_index_needs_rebaseline(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            for relative in ("memory", "logs", "audit", "drafts"):
                (root / relative).mkdir(parents=True, exist_ok=True)
            (root / "vkr-state.md").write_text("state", encoding="utf-8")
            write_json(root / "memory/handoff.json", {})
            write_json(root / "memory/artifact-index.json", {"schema_version": 1, "artifacts": {}})
            write_json(root / "audit/manifest.json", {"runs": []})
            for relative in ("logs/activity.jsonl", "logs/tool-runs.jsonl", "memory/decisions.jsonl"):
                (root / relative).write_text("", encoding="utf-8")
            (root / "drafts/new.md").write_text("new", encoding="utf-8")
            # 6.33: индекс старой схемы — needs_rebaseline, а не external_changes_detected (SPEC 9.6).
            self.assertEqual("needs_rebaseline", memory.status(root, 10)["status"])
            memory.record(root, None, rebaseline=True)
            self.assertEqual("ok", memory.status(root, 10)["status"])

    def test_assets_are_a4_and_template_has_page_field(self):
        for path in (SKILL / "assets").glob("*.docx"):
            doc = Document(path)
            for section in doc.sections:
                self.assertAlmostEqual(210, section.page_width.mm, delta=0.2, msg=path.name)
                self.assertAlmostEqual(297, section.page_height.mm, delta=0.2, msg=path.name)
        template = SKILL / "assets" / "template-mpgu.docx"
        with zipfile.ZipFile(template) as archive:
            footer_text = "".join(archive.read(name).decode("utf-8") for name in archive.namelist() if name.startswith("word/footer") and name.endswith(".xml"))
        self.assertIn("PAGE", footer_text)

    def test_word_saved_fixture_has_no_font_or_color_false_error(self):
        # Прежний blocker 6.31: ложные ошибки шрифта и цвета после сохранения в Word.
        # Фикстура сохранена Word и закоммичена; подробные сценарии — tests/test_validator_633.py
        # (WordFixtureTests, EffectiveFormattingTests).
        self.assertTrue(WORD_SAVED_FIXTURE.is_file(), WORD_SAVED_FIXTURE)
        doc = Document(WORD_SAVED_FIXTURE)
        for result in validator.check_fonts(doc) + validator.check_font_color(doc):
            self.assertFalse(not result["ok"] and result["severity"] == "error", result)


if __name__ == "__main__":
    unittest.main(verbosity=2)
