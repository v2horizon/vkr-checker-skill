"""Прежние проверки валидатора (6.32), переведены на unittest в 6.33 (WP3).

Утверждения сохранены; запуск: python -m unittest tests.test_validator_enhancements
или python tests/test_validator_enhancements.py.
"""
from __future__ import annotations

import importlib.util
import unittest
from pathlib import Path

from docx import Document
from docx.shared import Mm


ROOT = Path(__file__).resolve().parents[1]
VALIDATOR = ROOT / "vkr-mpgu" / "scripts" / "validate_vkr.py"
spec = importlib.util.spec_from_file_location("validate_vkr_enhancements", VALIDATOR)
module = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(module)


def make_doc(*paragraphs: str) -> Document:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    return doc


class ValidatorEnhancementTests(unittest.TestCase):
    def test_placeholders_are_errors(self):
        placeholder_doc = make_doc(
            "Результаты пилота: [N] участников, средняя оценка [X].",
            "По итогам внесены [список изменений].",
        )
        placeholder = module.check_unresolved_placeholders(placeholder_doc)
        self.assertEqual("error", placeholder[0]["severity"])
        self.assertIn("[N]", placeholder[0]["message"])

    def test_private_practice_pilot_minimum(self):
        private_doc = make_doc("Пилотирование в частной практике репетиторства проведено с 3 участниками.")
        self.assertEqual("error", module.check_pilot_sample_size(private_doc)[0]["severity"])

    def test_engineering_pilot_has_no_universal_minimum(self):
        engineering_doc = make_doc("Пилотное тестирование серверного прототипа проведено с 3 пользователями.")
        engineering = module.check_pilot_sample_size(engineering_doc)
        self.assertIs(True, engineering[0]["ok"])
        self.assertEqual("info", engineering[0]["severity"])

    def test_pedagogical_frame_depends_on_profile(self):
        non_educational_doc = make_doc(
            "Разработан сервис сбора телеметрии. Нагрузочное тестирование показало 1200 запросов в секунду."
        )
        context = module.check_pedagogical_frame(non_educational_doc)
        self.assertIs(True, context[0]["ok"])
        self.assertEqual("info", context[0]["severity"])
        project_context = module.check_pedagogical_frame(non_educational_doc, "mpgu-09-project")
        self.assertIs(False, project_context[0]["ok"])
        self.assertEqual("error", project_context[0]["severity"])

    def test_project_profile_three_chapters_and_pilot(self):
        project_doc = Document()
        project_doc.add_paragraph("Глава I. Теоретическая часть", style="Heading 1")
        project_doc.add_paragraph("Образовательный продукт предназначен для студентов.")
        project_doc.add_paragraph("Глава II. Аналитическая часть", style="Heading 1")
        project_doc.add_paragraph("Описаны целевая аудитория и методические основания.")
        project_doc.add_paragraph("Глава III. Проектная часть", style="Heading 1")
        project_doc.add_paragraph("Пилотирование образовательного приложения проведено с 8 студентами.")
        self.assertTrue(all(item["ok"] for item in module.check_mpgu_09_profile(project_doc, "mpgu-09-project")))
        self.assertIs(True, module.check_pilot_sample_size(project_doc, "mpgu-09-project")[0]["ok"])

        broken_project = Document()
        for number in ("I", "II", "III", "IV"):
            broken_project.add_paragraph(f"Глава {number}. Раздел", style="Heading 1")
        broken_checks = module.check_mpgu_09_profile(broken_project, "mpgu-09-project")
        self.assertEqual(2, sum(not item["ok"] for item in broken_checks))

    def test_page_size(self):
        a4_doc = Document()
        a4_doc.sections[0].page_width = Mm(210)
        a4_doc.sections[0].page_height = Mm(297)
        self.assertTrue(all(item["ok"] for item in module.check_page_size(a4_doc)))
        letter_doc = Document()
        self.assertTrue(any(not item["ok"] for item in module.check_page_size(letter_doc)))

    def test_personal_pronouns(self):
        pronoun_doc = make_doc("Я реализовал прототип, а затем мы провели тестирование.")
        pronouns = module.check_personal_pronouns(pronoun_doc)
        self.assertTrue(any(item["severity"] == "error" for item in pronouns))


if __name__ == "__main__":
    unittest.main(verbosity=2)
