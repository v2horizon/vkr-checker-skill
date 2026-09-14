"""Сборка DOCX из черновиков, рендерер и импорт (vkr-mpgu 6.33).

Фикстура: tests/fixtures/build/project — проект со всеми конструкциями формата
черновиков. tests/fixtures/build/vkr-word-saved.docx — DOCX этой фикстуры после
сборки, обновления полей и сохранения в Word 16 (Word COM, скрипт
tests/fixtures/build/make_word_fixture.ps1). Если фикстурные черновики меняются,
этот файл нужно пересоздать тем же скриптом.

Запуск: python -m unittest discover -s tests -t .   или   python tests/test_build_import_633.py
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

os.environ.setdefault("PYTHONDONTWRITEBYTECODE", "1")
sys.dont_write_bytecode = True

from docx import Document  # noqa: E402
from docx.shared import Mm  # noqa: E402
from lxml import etree  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "vkr-mpgu"
SCRIPTS = SKILL / "scripts"
FIXTURE = ROOT / "tests" / "fixtures" / "build" / "project"
WORD_FIXTURE = ROOT / "tests" / "fixtures" / "build" / "vkr-word-saved.docx"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))


def load_module(name: str):
    spec = importlib.util.spec_from_file_location(f"{name}_under_test_build", SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


generator = load_module("create_vkr_docx")
fb = load_module("format_bibliography")
build = load_module("build_vkr")


def run_script(name: str, *args):
    env = dict(os.environ)
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    return subprocess.run([sys.executable, str(SCRIPTS / name), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8", env=env)


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def copy_fixture(folder: Path) -> Path:
    target = folder / "project"
    shutil.copytree(str(FIXTURE), str(target))
    return target


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def document_root(path: Path):
    with zipfile.ZipFile(str(path)) as archive:
        return etree.fromstring(archive.read("word/document.xml")), etree.fromstring(archive.read("word/numbering.xml"))


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace("`", "")).strip()


def new_project(folder: Path) -> Path:
    """Проект сразу после init: конфиг, шаблонные черновики, пустой реестр."""
    root = folder / "imported"
    (root / "drafts").mkdir(parents=True)
    (root / "evidence").mkdir()
    shutil.copy(str(FIXTURE / "vkr-project.json"), str(root / "vkr-project.json"))
    (root / "sources.json").write_text("[]\n", encoding="utf-8")
    (root / "evidence" / "index.json").write_text(json.dumps(
        {"methodology": [], "sources": [], "product": [], "pilot": [], "figures": [], "defense": []}), encoding="utf-8")
    templates = {
        "annotation.md": "# Аннотация\n\n<!-- Заполняется после основного текста. -->\n",
        "introduction.md": "# Введение\n\n<!-- Черновик создаёт основной агент после утверждения плана. -->\n",
        "chapter-1.md": "# Глава I. [Название]\n\n<!-- Черновик -->\n",
        "chapter-2.md": "# Глава II. [Название]\n\n<!-- Черновик -->\n",
        "chapter-3.md": "# Глава III. [Название]\n\n<!-- Черновик проектной главы -->\n",
        "conclusion.md": "# Заключение\n\n<!-- Черновик -->\n",
    }
    for name, text in templates.items():
        (root / "drafts" / name).write_text(text, encoding="utf-8")
    return root


def draft_summary(root: Path) -> dict:
    """Главы, названия параграфов и число ссылок по черновикам."""
    summary = {"chapters": [], "citations": {}}
    for path in sorted((root / "drafts").glob("*.md")):
        blocks, _ = build.parse_markdown(path.read_text(encoding="utf-8"))
        count = sum(len(fb.citation_refs(text)) for block in blocks for text in build.block_texts(block))
        summary["citations"][path.name] = count
        if path.name.startswith("chapter-"):
            report = build.Report()
            chapter = build.split_chapter(blocks, path.name, int(re.findall(r"\d+", path.name)[0]), report)
            summary["chapters"].append((chapter["title"], [p["title"] for p in chapter["paragraphs"]], bool(chapter["conclusion"])))
    return summary


# ---------------------------------------------------------------------------
# Рендерер: схема и документация
# ---------------------------------------------------------------------------

class GeneratorSchemaTest(unittest.TestCase):
    TITLE_PAGE = {
        "university": "Университет-TP", "institute": "Институт-TP", "department": "кафедра Кафедра-TP",
        "program_code": "09.03.02", "program_name": "Направление-TP", "program_profile": "Профиль-TP",
        "work_type": "Выпускная квалификационная работа (магистерская диссертация)", "author": "Автор-TP Иван Иванович",
        "group": "ГР-TP", "supervisor": "Руководитель-TP", "supervisor_title": "Должность-TP",
        "head_of_department": "Заведующий-TP", "head_title": "Степень-TP", "city": "Город-TP", "year": "2031",
    }

    @staticmethod
    def doc_tables(markdown: str) -> dict:
        sections, current = {}, None
        for line in markdown.splitlines():
            if line.startswith("## "):
                current = line[3:].strip()
                sections[current] = []
            elif current and line.startswith("|") and not re.match(r"^\|[-\s|]+\|$", line):
                sections[current].append([cell.strip() for cell in line.strip().strip("|").split("|")])
        return sections

    def test_documented_keys_match_generator(self):
        doc = (SKILL / "references" / "docx-input-schema.md").read_text(encoding="utf-8")
        tables = self.doc_tables(doc)

        def keys(section: str, column: int = 0, skip_header: bool = True, stop_at_header: str = "") -> set:
            rows = tables[section][1:] if skip_header else tables[section]
            found = set()
            for row in rows:
                if stop_at_header and row[0] == stop_at_header:
                    break
                found.update(re.findall(r"`([a-z_]+)`", row[column]))
            return found

        self.assertEqual(set(generator.TOP_LEVEL_KEYS), keys("Ключи верхнего уровня"))
        title_rows = tables["title_page"]
        split = next(i for i, row in enumerate(title_rows) if row[0] == "Устаревший ключ")
        documented_tp = {re.findall(r"`([a-z_]+)`", row[0])[0] for row in title_rows[1:split]}
        documented_legacy = {re.findall(r"`([a-z_]+)`", row[0])[0] for row in title_rows[split + 1:]}
        self.assertEqual(set(generator.TITLE_PAGE_KEYS), documented_tp)
        self.assertEqual(set(generator.LEGACY_TITLE_PAGE_KEYS), documented_legacy)
        self.assertEqual(set(generator.CHAPTER_KEYS), keys("Глава"))
        self.assertEqual(set(generator.PARAGRAPH_KEYS), keys("Параграф"))
        self.assertEqual(set(generator.APPENDIX_KEYS), keys("Приложение"))
        blocks = {re.findall(r"`([a-z_]+)`", row[0])[0]: set(re.findall(r"`([a-z_]+)`", row[1])) for row in tables["Блоки"][1:]}
        self.assertEqual({k: set(v) for k, v in generator.BLOCK_KEYS.items()}, blocks)
        self.assertIn("final/vkr.docx", doc)
        self.assertIn("build_vkr.py", doc)
        example = json.loads(re.search(r"## Пример\n\n```json\n(.*?)\n```", doc, re.S).group(1))
        content, _warnings = generator.normalize_content(example)
        self.assertEqual(1, len(content["chapters"]))
        drafts_doc = (SKILL / "references" / "drafts-format.md").read_text(encoding="utf-8")
        self.assertIn("build_vkr.py", drafts_doc)
        self.assertIn("import_docx.py", drafts_doc)

    def build_minimal(self, folder: Path, **extra) -> tuple:
        data = {
            "title": "Тема-TP",
            "title_page": dict(self.TITLE_PAGE),
            "annotation": "Аннотация.",
            "introduction": "Введение.",
            "chapters": [{"title": "Глава I. Теория", "paragraphs": [{"title": "1.1. Раздел", "text": "Текст."}]}],
            "conclusion": "Заключение.",
            "bibliography": ["Иванов, А.Б. Книга [Текст] / А.Б. Иванов. – М.: Наука, 2020."],
        }
        data.update(extra)
        output = folder / "out.docx"
        warnings = []
        generator.build_vkr_docx(data, output, warnings=warnings)
        return Document(str(output)), warnings, output

    def title_texts(self, doc) -> list:
        texts = []
        for paragraph in doc.paragraphs:
            if paragraph.text == "Аннотация":
                break
            texts.append(paragraph.text)
        return texts

    def test_every_title_page_key_is_rendered(self):
        with tempfile.TemporaryDirectory() as td:
            doc, warnings, _ = self.build_minimal(Path(td))
            title = "\n".join(self.title_texts(doc))
            for key, value in self.TITLE_PAGE.items():
                expected = {"department": "Кафедра-TP", "work_type": "(магистерская диссертация)",
                            "university": "«Университет-TP»", "group": "Группа ГР-TP"}.get(key, value)
                with self.subTest(key=key):
                    self.assertIn(expected, title)
            for fixed in ("Выпускная квалификационная работа", "«Допустить к защите»", "Заведующий кафедрой",
                          "Научный руководитель –", "Тема-TP", "Код и направление подготовки: 09.03.02. Направление-TP",
                          "Город-TP – 2031 год"):
                self.assertIn(fixed, title)
            styles = {p.style.name for p in doc.paragraphs[:len(self.title_texts(doc))]}
            self.assertEqual({"VKR Title"}, styles)
            self.assertFalse([w for w in warnings if "титульный лист" in w])

    def test_title_page_missing_fields_become_markers_and_unknown_keys_fail(self):
        with tempfile.TemporaryDirectory() as td:
            doc, warnings, _ = self.build_minimal(Path(td), title_page={"author": "Иванов Иван Иванович"})
            title = "\n".join(self.title_texts(doc))
            self.assertIn("[ЗАПОЛНИТЬ: И.О. Фамилия научного руководителя]", title)
            self.assertTrue(any("supervisor" in w for w in warnings))
            with self.assertRaises(generator.InputError) as ctx:
                self.build_minimal(Path(td), title_page={"author": "И", "universty": "опечатка"})
            self.assertIn("universty", str(ctx.exception))
            with self.assertRaises(generator.InputError):
                self.build_minimal(Path(td), title_page={"supervisor": "без автора"})
            with self.assertRaises(generator.InputError):
                self.build_minimal(Path(td), unknown_top="x")
            legacy = {"author": "Иванов Иван Иванович", "direction_code": "09.03.02", "direction_name": "ИСТ-легаси",
                      "supervisor_name": "П.С. Легаси", "supervisor_position": "доцент", "supervisor_degree": "канд. пед. наук",
                      "head_name": "С.И. Легаси", "head_degree": "профессор", "profile": "Профиль-легаси"}
            doc, warnings, _ = self.build_minimal(Path(td), title_page=legacy)
            title = "\n".join(self.title_texts(doc))
            for text in ("ИСТ-легаси", "П.С. Легаси", "доцент, канд. пед. наук", "С.И. Легаси", "профессор", "Профиль-легаси."):
                self.assertIn(text, title)
            self.assertTrue(any("устаревшие ключи" in w for w in warnings))

    def test_collective_authors_and_chapter_prefix(self):
        with tempfile.TemporaryDirectory() as td:
            tp = dict(self.TITLE_PAGE)
            tp.pop("author")
            tp["authors"] = ["Первый Участник", "Второй Участник"]
            doc, _, _ = self.build_minimal(Path(td), title_page=tp)
            texts = [p.text for p in doc.paragraphs]
            self.assertIn("Первый Участник", texts)
            self.assertIn("Второй Участник", texts)
            self.assertIn("Глава I. Теория", texts)
            self.assertFalse(any("Глава I. Глава" in t for t in texts))
            self.assertIn("1.1. Раздел", texts)

    def test_skip_last_page_leaves_no_trailing_break(self):
        with tempfile.TemporaryDirectory() as td:
            _, _, output = self.build_minimal(Path(td), skip_last_page=True,
                                              appendices=[{"title": "Материалы", "content": "Текст приложения."}])
            root, _ = document_root(output)
            body = root.find(f"{W}body")
            self.assertEqual([], [br for br in body.iter(f"{W}br") if br.get(f"{W}type") == "page"])
            last = [child for child in body if child.tag == f"{W}p"][-1]
            self.assertEqual("Текст приложения.", "".join(t.text or "" for t in last.iter(f"{W}t")))

    def test_cli_bom_and_exit_codes(self):
        with tempfile.TemporaryDirectory() as td:
            folder = Path(td)
            data = {"title": "Тема", "title_page": {"author": "Иванов Иван Иванович"},
                    "chapters": [{"title": "Теория", "paragraphs": [{"title": "Раздел", "text": "Текст [1]."}]}],
                    "sources": [{"id": 1, "type": "book", "title": "Книга", "place": "М.", "publisher": "Н", "year": 2020}]}
            path = folder / "input.json"
            path.write_bytes(b"\xef\xbb\xbf" + json.dumps(data, ensure_ascii=False).encode("utf-8"))
            proc = run_script("create_vkr_docx.py", path, "-o", folder / "ok.docx", "--json")
            self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
            self.assertEqual("ok", json.loads(proc.stdout)["status"])
            broken = dict(data, chapters=[])
            path.write_text(json.dumps(broken, ensure_ascii=False), encoding="utf-8")
            proc = run_script("create_vkr_docx.py", path, "-o", folder / "bad.docx")
            self.assertEqual(2, proc.returncode)
            self.assertIn("chapters", proc.stderr)
            self.assertNotIn("Traceback", proc.stderr)
            wrong_ref = dict(data)
            wrong_ref["chapters"] = [{"title": "Теория", "paragraphs": [{"title": "Раздел", "text": "Текст [7]."}]}]
            path.write_text(json.dumps(wrong_ref, ensure_ascii=False), encoding="utf-8")
            proc = run_script("create_vkr_docx.py", path, "-o", folder / "bad.docx")
            self.assertEqual(1, proc.returncode)
            self.assertIn("[7]", proc.stderr)

    def test_direct_generator_sorts_sources_and_rewrites_references(self):
        with tempfile.TemporaryDirectory() as td:
            data = {
                "title": "Тема", "title_page": {"author": "Иванов Иван Иванович"},
                "chapters": [{"title": "Теория", "paragraphs": [{"title": "Раздел", "blocks": [
                    {"type": "text", "content": "Сначала [1, с. 5], затем [@abr] и `arr[2]`."},
                    {"type": "listing", "caption": "Код", "code": "x = arr[2]"}]}]}],
                "sources": [
                    {"id": 1, "type": "book", "authors": ["Яковлев А.А."], "title": "Яблоки", "place": "М.", "publisher": "Н", "year": 2020},
                    {"id": "abr", "type": "book", "authors": ["Абрамов Б.Б."], "title": "Арбузы", "place": "М.", "publisher": "Н", "year": 2021},
                    {"id": 2, "type": "book", "authors": ["Борисов В.В."], "title": "Бананы", "place": "М.", "publisher": "Н", "year": 2022},
                ],
            }
            data["introduction"] = "Первая строка абзаца\nпродолжается здесь.\n\nВторой абзац."
            output = Path(td) / "direct.docx"
            generator.build_vkr_docx(data, output)
            texts = [p.text for p in Document(str(output)).paragraphs]
            self.assertIn("Первая строка абзаца продолжается здесь.", texts)  # D16: без разрыва строки
            self.assertIn("Второй абзац.", texts)
            self.assertIn("Сначала [3, с. 5], затем [1] и arr[2].", texts)
            self.assertIn("x = arr[2]", texts)
            bib = [t for t in texts if re.match(r"^\d+\. ", t)]
            self.assertEqual(["1. Абрамов", "2. Борисов", "3. Яковлев"], [t[:10].rstrip(",") for t in bib])
            generator.build_vkr_docx(data, output, keep_source_order=False)
            data_keep = dict(data, keep_source_order=True)
            with self.assertRaises(generator.ContentError):
                generator.build_vkr_docx(data_keep, output)  # строковый id при сохранении порядка


class CliContractTest(unittest.TestCase):
    """SPEC 10: --help без python-docx, понятная ошибка и код 2 при его отсутствии."""

    def test_scripts_without_python_docx(self):
        with tempfile.TemporaryDirectory() as td:
            folder = Path(td)
            blocker = folder / "blocker"
            (blocker / "docx").mkdir(parents=True)
            (blocker / "docx" / "__init__.py").write_text("raise ImportError('python-docx заблокирован тестом')\n", encoding="utf-8")
            env = dict(os.environ)
            env.update(PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8", PYTHONPATH=str(blocker))

            def run(name, *args):
                return subprocess.run([sys.executable, str(SCRIPTS / name), *map(str, args)],
                                      capture_output=True, text=True, encoding="utf-8", env=env)

            for name in ("build_vkr.py", "import_docx.py", "create_vkr_docx.py", "format_bibliography.py"):
                proc = run(name, "--help")
                self.assertEqual(0, proc.returncode, name + proc.stderr)
            root = copy_fixture(folder)
            proc = run("build_vkr.py", root, "--json")
            self.assertEqual(2, proc.returncode, proc.stdout)
            self.assertIn("DEPENDENCY_MISSING", proc.stdout)
            self.assertFalse((root / "final" / "vkr.docx").exists())
            proc = run("import_docx.py", root, "--docx", WORD_FIXTURE)
            self.assertEqual(2, proc.returncode)
            self.assertIn("python-docx", proc.stderr)
            data = folder / "content.json"
            data.write_text(json.dumps({"title": "Т", "title_page": {"author": "А"}, "chapters": [{"title": "Г"}]}), encoding="utf-8")
            proc = run("create_vkr_docx.py", data, "-o", folder / "x.docx")
            self.assertEqual(2, proc.returncode)
            self.assertIn("python-docx", proc.stderr)
            self.assertNotIn("Traceback", proc.stderr)


# ---------------------------------------------------------------------------
# Сборка фикстуры
# ---------------------------------------------------------------------------

class BuildFixtureTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.root = copy_fixture(Path(cls.tmp.name))
        cls.proc = run_script("build_vkr.py", cls.root, "--json")
        cls.report = json.loads(cls.proc.stdout)
        cls.docx = cls.root / "final" / "vkr.docx"
        cls.doc = Document(str(cls.docx)) if cls.docx.is_file() else None
        cls.paragraphs = [(p.style.name, p.text) for p in cls.doc.paragraphs] if cls.doc else []

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_build_manifest_records_warnings(self):
        # Предупреждения сборки повторяются в манифесте, иначе новый чат их не увидит.
        manifest = json.loads((self.root / "exports" / "build-manifest.json").read_text(encoding="utf-8"))
        self.assertEqual([w["code"] for w in self.report["warnings"]], [w["code"] for w in manifest["warnings"]])
        self.assertTrue(any(w["code"] == "RENDER" for w in manifest["warnings"]), manifest["warnings"])

    def test_build_succeeds_and_writes_outputs(self):
        self.assertEqual(0, self.proc.returncode, self.proc.stdout + self.proc.stderr)
        self.assertEqual("ok", self.report["status"])
        for rel in ("exports/vkr-content.json", "exports/citation-map.json", "final/vkr.docx"):
            self.assertTrue((self.root / rel).is_file(), rel)
        self.assertEqual({"vkr_type": "project", "chapters": 3, "paragraphs": 6, "appendices": 2, "tables": 2, "figures": 2,
                          "listings": 2, "lists": 4, "sources_in_list": 9, "sources_cited": 9}, self.report["stats"])
        self.assertEqual([], self.report["errors"])
        codes = {w["code"] for w in self.report["warnings"]}
        self.assertIn("SOURCES_NOT_CONFIRMED", codes)
        self.assertNotIn("SOURCE_NOT_CITED", codes)
        verification = next(w["message"] for w in self.report["warnings"] if w["code"] == "SOURCE_VERIFICATION_MISSING")
        self.assertIn("egorov2020", verification)  # confirmed без verification — как проверяет doctor на final
        self.assertNotRegex(verification, r"(?:: |, )3(?:,|$)")  # у id 3 verification полный
        self.assertNotIn("zhukov2023", verification)  # pending проверяется отдельно

    def test_report_names_canonical_next_steps(self):
        # 6.33: после final/vkr.docx — update_docx_fields → clean_docx_metadata --in-place → vkr_audit validate.
        steps = self.report["next_steps"]
        self.assertEqual(["update_fields", "clean_metadata", "validate"], [step["step"] for step in steps])
        docx = str((self.root / "final" / "vkr.docx").resolve())
        project = str(self.root.resolve())
        self.assertRegex(steps[0]["command"], r'update_docx_fields\.py"')
        self.assertTrue(steps[0]["command"].endswith(f'"{docx}"'), steps[0]["command"])
        self.assertIn("F9", steps[0]["manual"])
        self.assertRegex(steps[1]["command"], r'clean_docx_metadata\.py"')
        self.assertTrue(steps[1]["command"].endswith(f'"{docx}" --in-place'), steps[1]["command"])
        self.assertRegex(steps[2]["command"], r'vkr_audit\.py"')
        self.assertTrue(steps[2]["command"].endswith(f'"{project}" validate'), steps[2]["command"])
        for step in steps:
            script = Path(re.match(r'^python "([^"]+)"', step["command"]).group(1))
            self.assertEqual(SCRIPTS.resolve(), script.parent)
            self.assertTrue(script.is_file(), script)
        self.assertNotIn("validate_vkr", json.dumps(self.report, ensure_ascii=False))
        self.assertEqual("exports/build-manifest.json", self.report["outputs"]["build_manifest"])
        self.assertTrue((self.root / "exports" / "build-manifest.json").is_file())

    def test_bibliography_is_alphabetical(self):
        entries = [text for style, text in self.paragraphs if style == "VKR Bibliography"]
        heads = ["Абрамова", "Егоров", "Ёлкин", "Жуков", "Об образовании", "Об утверждении", "Яковлев", "Smith", "Unity Manual"]
        self.assertEqual(len(heads), len(entries))
        for number, (head, entry) in enumerate(zip(heads, entries), 1):
            self.assertTrue(entry.startswith(f"{number}. {head}"), entry)

    def test_each_reference_points_to_the_right_entry(self):
        numbers = read_json(self.root / "exports" / "citation-map.json")["numbers"]
        registry = {str(src["id"]): src for src in read_json(self.root / "sources.json")}
        entries = [text for style, text in self.paragraphs if style == "VKR Bibliography"]
        for number, ident in numbers.items():
            expected = f"{number}. " + fb.format_source(fb.normalize_source(registry[str(ident)], 1))
            self.assertEqual(expected, entries[int(number) - 1])
        by_id = {str(ident): int(n) for n, ident in numbers.items()}
        int_ids = {int(str(ident)): int(n) for n, ident in numbers.items() if str(ident).isdigit()}
        resolve = lambda kind, ref: by_id.get(str(ref)) if kind == "id" else int_ids.get(int(ref))
        docx_texts = {norm(text) for _, text in self.paragraphs}
        for table in self.doc.tables:
            for row in table.rows:
                docx_texts.update(norm(cell.text) for cell in row.cells)
        checked = 0
        for path in sorted((self.root / "drafts").glob("*.md")):
            blocks, _ = build.parse_markdown(path.read_text(encoding="utf-8"))
            for block in blocks:
                if block["kind"] == "listing":
                    continue
                for text in build.block_texts(block):
                    if not fb.citation_refs(text):
                        continue
                    expected = norm(build.sort_citation_numbers(fb.rewrite_citations(text, resolve)))
                    if block["kind"] == "table" and text == block.get("source"):
                        expected = "Источник: " + expected
                    with self.subTest(file=path.name, text=text[:50]):
                        self.assertIn(expected, docx_texts)
                        checked += 1
        self.assertEqual(9, checked)  # 9 текстовых фрагментов со ссылками, включая ячейку таблицы и «Источник»
        self.assertEqual(7, int_ids[3])  # legacy [3] (Яковлев, id 3) стал [7]
        self.assertIn(norm("Анализ литературы показал, что тренажёр должен сочетать наглядность и обратную связь [7, с. 17]."), docx_texts)

    def test_code_is_not_renumbered(self):
        code = [text for style, text in self.paragraphs if style == "VKR Listing"]
        joined = "\n".join(code)
        for fragment in ("cells[3];", "TEMPLATE[3]", "индекс [x]", '"{{ var }}"'):
            self.assertIn(fragment, joined)

    def test_numbered_lists_restart_from_one(self):
        root, numbering = document_root(self.docx)
        nums = {num.get(f"{W}numId"): num for num in numbering.findall(f"{W}num")}
        abstract_fmt = {a.get(f"{W}abstractNumId"): a.find(f"{W}lvl/{W}numFmt").get(f"{W}val") for a in numbering.findall(f"{W}abstractNum")}
        groups, previous = [], None
        for p in root.find(f"{W}body"):
            num_id = p.find(f"{W}pPr/{W}numPr/{W}numId") if p.tag == f"{W}p" else None
            value = num_id.get(f"{W}val") if num_id is not None else None
            if value and value != previous:
                groups.append(value)
            previous = value
        self.assertEqual(4, len(groups))
        self.assertEqual(len(groups), len(set(groups)), "у каждого списка свой w:num")
        formats = []
        for value in groups:
            num = nums[value]
            override = num.find(f"{W}lvlOverride/{W}startOverride")
            self.assertIsNotNone(override)
            self.assertEqual("1", override.get(f"{W}val"))
            formats.append(abstract_fmt[num.find(f"{W}abstractNumId").get(f"{W}val")])
        self.assertEqual(["decimal", "decimal", "decimal", "bullet"], formats)
        for style, text in self.paragraphs:
            if style == "List Paragraph":
                self.assertFalse(text.startswith(("1.", "-", "–")), text)

    def test_page_breaks_no_empty_pages(self):
        root, _ = document_root(self.docx)
        body = root.find(f"{W}body")
        self.assertFalse([br for br in body.iter(f"{W}br") if br.get(f"{W}type") == "page"], "разрывы только через pageBreakBefore")
        starts = []
        for p in body.iter(f"{W}p"):
            if p.find(f"{W}pPr/{W}pageBreakBefore") is not None:
                text = "".join(t.text or "" for t in p.iter(f"{W}t"))
                self.assertTrue(text.strip(), "абзац с новой страницы не пустой")
                starts.append(text)
        self.assertEqual(["Аннотация", "Содержание", "Введение", "Глава I. Теоретические основы применения VR в обучении стереометрии",
                          "Глава II. Анализ платформ и требований к тренажёру", "Глава III. Разработка и пилотирование тренажёра",
                          "Заключение", "Список использованной литературы", "Приложение 1", "Приложение 2",
                          "Выпускная квалификационная работа"], starts)
        last = [child for child in body if child.tag == f"{W}p"][-1]
        self.assertTrue("".join(t.text or "" for t in last.iter(f"{W}t")).strip())

    def test_word_saved_fixture_title_on_first_page_and_no_empty_pages(self):
        root, _ = document_root(WORD_FIXTURE)
        body = root.find(f"{W}body")
        paragraphs = [p for p in body if p.tag == f"{W}p"]
        texts = ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in paragraphs]
        first_break = next(i for i, p in enumerate(paragraphs) if list(p.iter(f"{W}lastRenderedPageBreak")))
        self.assertEqual("Аннотация", texts[first_break], "титул занимает ровно одну страницу")
        breaks = [i for i, p in enumerate(paragraphs) if list(p.iter(f"{W}lastRenderedPageBreak"))]
        self.assertEqual(13, len(breaks))
        for index in breaks:
            self.assertTrue(texts[index].strip(), f"страница начинается с пустого абзаца {index}")

    def test_title_page_contains_project_fields(self):
        title_page = read_json(self.root / "vkr-project.json")["title_page"]
        texts = []
        for style, text in self.paragraphs:
            if text == "Аннотация":
                break
            self.assertEqual("VKR Title", style)
            texts.append(text)
        joined = "\n".join(texts)
        for key, value in title_page.items():
            if key == "department":
                value = "информатики и методики обучения информатике"
            with self.subTest(key=key):
                self.assertIn(str(value), joined)
        self.assertIn("«Допустить к защите»", joined)
        self.assertIn("Группа ИСТ-41", joined)
        self.assertLessEqual(len(texts), 40)

    def test_styles_and_fonts(self):
        styles = dict((text, style) for style, text in self.paragraphs)
        self.assertEqual("Heading 1", styles["Введение"])
        self.assertEqual("Heading 1", styles["Список использованной литературы"])
        self.assertEqual("Heading 2", styles["1.1. Пространственное мышление учащихся"])
        self.assertEqual("Heading 2", styles["1.2. Опыт использования виртуальной реальности в школе"])
        self.assertEqual("VKR Caption", styles["Таблица 1"])
        self.assertEqual("VKR Caption", styles["Рисунок 1 — Архитектура тренажёра"])
        self.assertEqual("VKR Caption", styles["Листинг 1 — Проверка ответа ученика"])
        self.assertEqual("VKR Caption", styles["Источник: составлено автором по [9]"])
        self.assertEqual("VKR Appendix Label", styles["Приложение 1"])
        self.assertEqual("VKR Appendix Title", styles["Исходный код модуля проверки"])
        self.assertFalse(any("Глава I. Глава" in text for _, text in self.paragraphs))
        with zipfile.ZipFile(str(self.docx)) as archive:
            styles_xml = etree.fromstring(archive.read("word/styles.xml"))
        for style in styles_xml.findall(f"{W}style"):
            name = style.find(f"{W}name").get(f"{W}val")
            if name in ("heading 1", "heading 2", "VKR Listing", "Normal"):
                fonts = style.find(f"{W}rPr/{W}rFonts")
                self.assertIsNotNone(fonts, name)
                self.assertIsNone(fonts.get(f"{W}asciiTheme"), name)
                expected = "Consolas" if name == "VKR Listing" else "Times New Roman"
                self.assertEqual(expected, fonts.get(f"{W}ascii"), name)
            if name in ("heading 1", "heading 2"):
                self.assertEqual("000000", style.find(f"{W}rPr/{W}color").get(f"{W}val"))
            if name in ("toc 1", "toc 2"):
                self.assertEqual("0", style.find(f"{W}pPr/{W}ind").get(f"{W}firstLine"))

    def test_toc_field_placeholder_and_appendix_tc(self):
        root, _ = document_root(self.docx)
        toc = next(p for p in root.iter(f"{W}p") if any("TOC" in (i.text or "") for i in p.iter(f"{W}instrText")))
        nodes = list(toc.iter())
        placeholder = next(i for i, n in enumerate(nodes) if n.tag == f"{W}t" and "Для обновления" in (n.text or ""))
        end = next(i for i, n in enumerate(nodes) if n.tag == f"{W}fldChar" and n.get(f"{W}fldCharType") == "end")
        self.assertLess(placeholder, end)
        instr = "".join(i.text for i in toc.iter(f"{W}instrText"))
        self.assertIn("\\f", instr)
        tc = [i.text for i in root.iter(f"{W}instrText") if (i.text or "").strip().startswith("TC")]
        self.assertEqual(['TC "Приложение 1. Исходный код модуля проверки" \\l 1', 'TC "Приложение 2. Анкета для учащихся" \\l 1'],
                         [t.strip() for t in tc])
        sect = next(n for n in root.iter(f"{W}sectPr"))
        names = [child.tag.rsplit("}", 1)[-1] for child in sect]
        self.assertLess(names.index("pgNumType"), names.index("cols"))
        # w:start не пишется: иначе Word перезапускает нумерацию в каждой новой секции.
        self.assertIsNone(sect.find(f"{W}pgNumType").get(f"{W}start"))

    def test_figure_file_inserted_within_text_width_and_placeholder(self):
        root, _ = document_root(self.docx)
        blips = list(root.iter(f"{A}blip"))
        self.assertEqual(1, len(blips))
        extent = next(root.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"))
        self.assertLessEqual(int(extent.get("cx")), int(Mm(165)))
        self.assertIn("[ Здесь вставить рисунок вручную в Word ]", [text for _, text in self.paragraphs])

    def test_rebuild_backs_up_previous_docx(self):
        with tempfile.TemporaryDirectory() as td:
            root = copy_fixture(Path(td))
            proc = run_script("build_vkr.py", root)
            self.assertEqual(0, proc.returncode, proc.stderr)
            # Итоговое сообщение называет те же шаги по порядку, без прежнего «F9, затем validate_vkr.py».
            needles = ("update_docx_fields.py", "Ctrl+A, F9", "clean_docx_metadata.py", "--in-place", "vkr_audit.py")
            positions = [proc.stdout.find(needle) for needle in needles]
            self.assertTrue(all(position >= 0 for position in positions), proc.stdout)
            self.assertEqual(sorted(positions), positions, proc.stdout)
            self.assertRegex(proc.stdout, r'vkr_audit\.py" "[^"]+" validate')
            self.assertNotIn("validate_vkr.py", proc.stdout)
            first = sha(root / "final" / "vkr.docx")
            proc = run_script("build_vkr.py", root, "--json", "-o", root / "audit" / "build-report.json")
            self.assertEqual(0, proc.returncode)
            backup = root / json.loads(proc.stdout)["outputs"]["backup"]
            self.assertTrue(re.match(r"^\d{8}T\d{6}Z-build", backup.parent.name), backup)
            self.assertEqual(first, sha(backup))
            self.assertTrue((root / "audit" / "build-report.json").is_file())
            content = read_json(root / "exports" / "vkr-content.json")
            proc = run_script("create_vkr_docx.py", root / "exports" / "vkr-content.json", "-o", Path(td) / "again.docx")
            self.assertEqual(0, proc.returncode, proc.stderr)
            self.assertNotIn("алфавитном", proc.stderr)
            self.assertEqual([t for s, t in self.paragraphs if s == "VKR Bibliography"],
                             [p.text for p in Document(str(Path(td) / "again.docx")).paragraphs if p.style.name == "VKR Bibliography"])
            self.assertTrue(content["bibliography"][0].startswith("1. Абрамова"))


class BuildErrorsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = copy_fixture(Path(self.tmp.name))
        (self.root / "final").mkdir()
        (self.root / "final" / "vkr.docx").write_bytes(b"previous build")

    def tearDown(self):
        self.tmp.cleanup()

    def build_json(self):
        proc = run_script("build_vkr.py", self.root, "--json")
        return proc.returncode, json.loads(proc.stdout)

    def assert_failed(self, code: str):
        rc, report = self.build_json()
        self.assertEqual(1, rc, report)
        self.assertIn(code, {e["code"] for e in report["errors"]})
        self.assertEqual(b"previous build", (self.root / "final" / "vkr.docx").read_bytes())
        self.assertFalse((self.root / "exports" / "vkr-content.json").exists())
        return report

    def edit_sources(self, fn):
        sources = read_json(self.root / "sources.json")
        fn(sources)
        (self.root / "sources.json").write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding="utf-8")

    def test_unknown_id(self):
        path = self.root / "drafts" / "chapter-2.md"
        path.write_text(path.read_text(encoding="utf-8") + "\nНовая мысль [@nope2024, с. 1].\n", encoding="utf-8")
        report = self.assert_failed("SOURCE_UNKNOWN")
        self.assertTrue(any("nope2024" in e["message"] and e.get("file") == "drafts/chapter-2.md" for e in report["errors"]))

    def test_unknown_legacy_number(self):
        path = self.root / "drafts" / "conclusion.md"
        path.write_text(path.read_text(encoding="utf-8") + "\nИтог [12].\n", encoding="utf-8")
        self.assert_failed("SOURCE_UNKNOWN")

    def test_rejected_source(self):
        self.edit_sources(lambda s: next(x for x in s if x["id"] == "egorov2020").update(status="rejected"))
        self.assert_failed("SOURCE_REJECTED")

    def test_unstructured_source(self):
        def unstructure(sources):
            index = next(i for i, x in enumerate(sources) if x["id"] == "smith2019")
            sources[index] = {"id": "smith2019", "raw_text": "Smith J. Immersive learning environments. 2019.", "status": "pending"}
        self.edit_sources(unstructure)
        self.assert_failed("SOURCE_UNSTRUCTURED")

    def test_duplicate_id(self):
        self.edit_sources(lambda s: s.append(dict(s[0])))
        self.assert_failed("SOURCE_ID_DUPLICATE")

    def test_required_drafts_by_vkr_type(self):
        (self.root / "drafts" / "chapter-3.md").unlink()
        self.assert_failed("DRAFT_REQUIRED_MISSING")
        config = read_json(self.root / "vkr-project.json")
        config["vkr_type"], config["profile"] = "regular", "mpgu-09-regular"
        (self.root / "vkr-project.json").write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
        (self.root / "drafts" / "chapter-2.md").write_text("# Пустая глава\n\n<!-- Черновик -->\n", encoding="utf-8")
        self.assert_failed("DRAFT_EMPTY")
        shutil.copy(str(FIXTURE / "drafts" / "chapter-2.md"), str(self.root / "drafts" / "chapter-2.md"))
        rc, report = self.build_json()
        self.assertEqual(0, rc, report["errors"])
        self.assertEqual(2, report["stats"]["chapters"])

    def test_missing_figure_file(self):
        (self.root / "evidence" / "figures" / "architecture.png").unlink()
        self.assert_failed("FIGURE_NOT_FOUND")

    def test_warnings_do_not_block(self):
        self.edit_sources(lambda s: s.append({"id": "extra", "type": "book", "authors": ["Борисов К.К."], "title": "Лишняя книга",
                                              "place": "М.", "publisher": "Наука", "year": 2020, "status": "confirmed"}))
        path = self.root / "drafts" / "introduction.md"
        path.write_text(path.read_text(encoding="utf-8") + "\nЦифра [ПРОВЕРИТЬ: число участников].\n", encoding="utf-8")
        rc, report = self.build_json()
        self.assertEqual(0, rc, report["errors"])
        codes = {w["code"] for w in report["warnings"]}
        self.assertTrue({"SOURCE_NOT_CITED", "MARKER", "RENDER"} <= codes, codes)
        self.assertEqual(10, report["stats"]["sources_in_list"])

    def test_content_only_and_output_option(self):
        rc, report = self.build_json()
        self.assertEqual(0, rc, report["errors"])
        (self.root / "final" / "vkr.docx").write_bytes(b"previous build")
        proc = run_script("build_vkr.py", self.root, "--content-only", "--json")
        self.assertEqual(0, proc.returncode)
        self.assertNotIn("docx", json.loads(proc.stdout)["outputs"])
        self.assertEqual([], json.loads(proc.stdout)["next_steps"])
        self.assertEqual(b"previous build", (self.root / "final" / "vkr.docx").read_bytes())
        proc = run_script("build_vkr.py", self.root, "--output", "exports/preview.docx", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        self.assertEqual("exports/preview.docx", json.loads(proc.stdout)["outputs"]["docx"])
        self.assertEqual([], json.loads(proc.stdout)["next_steps"])  # промежуточный DOCX: шаги сдачи не к нему
        self.assertTrue((self.root / "exports" / "preview.docx").is_file())
        self.assertEqual(b"previous build", (self.root / "final" / "vkr.docx").read_bytes())
        content = read_json(self.root / "exports" / "vkr-content.json")
        intro_texts = [b.get("content") for b in content["introduction"] if b["type"] == "text"]
        self.assertIn("Актуальность исследования определяется требованиями к цифровой образовательной "
                      "среде [5] и к предметным результатам по геометрии [6, с. 4].", intro_texts)

    def test_usage_errors(self):
        proc = run_script("build_vkr.py", self.root / "missing")
        self.assertEqual(2, proc.returncode)
        (self.root / "sources.json").write_text("[{", encoding="utf-8")
        proc = run_script("build_vkr.py", self.root)
        self.assertEqual(2, proc.returncode)
        self.assertNotIn("Traceback", proc.stderr)
        self.assertEqual(0, run_script("build_vkr.py", "--help").returncode)


# ---------------------------------------------------------------------------
# Импорт
# ---------------------------------------------------------------------------

class ImportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def built_fixture(self) -> Path:
        root = copy_fixture(self.folder)
        proc = run_script("build_vkr.py", root)
        self.assertEqual(0, proc.returncode, proc.stderr)
        return root

    def test_import_built_docx_into_new_project_and_rebuild(self):
        original = self.built_fixture()
        target = new_project(self.folder)
        proc = run_script("import_docx.py", target, "--docx", original / "final" / "vkr.docx", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        report = json.loads(proc.stdout)
        self.assertTrue((target / "audit" / "import-report.json").is_file())
        self.assertIn("последний лист (клятва)", report["skipped"])
        before, after = draft_summary(original), draft_summary(target)
        self.assertEqual(before["chapters"], after["chapters"])
        self.assertEqual(before["citations"], after["citations"])
        self.assertEqual(sorted(p.name for p in (original / "drafts").glob("*.md")), sorted(p.name for p in (target / "drafts").glob("*.md")))
        chapter3 = (target / "drafts" / "chapter-3.md").read_text(encoding="utf-8")
        self.assertIn("`XRGrabInteractable`", chapter3)
        self.assertIn("Листинг: Проверка ответа ученика\n\n```", chapter3)
        self.assertIn("    var y = cells[3];   // [3] в коде не перенумеровывается", chapter3)
        self.assertIn("![Экран результатов пилотирования]()", chapter3)
        chapter2 = (target / "drafts" / "chapter-2.md").read_text(encoding="utf-8")
        self.assertIn("Таблица: Сравнение VR-платформ для школы", chapter2)
        self.assertIn("Источник: составлено автором по [9]", chapter2)
        figure = re.search(r"!\[Архитектура тренажёра\]\((evidence/figures/[^)]+)\)", chapter2).group(1)
        self.assertEqual(sha(original / "evidence" / "figures" / "architecture.png"), sha(target / figure))
        index = read_json(target / "evidence" / "index.json")
        self.assertEqual([{"id": "figure-import-1", "path": figure, "kind": "figure"}], index["figures"])

        registry = read_json(target / "sources.json")
        self.assertEqual(9, len(registry))
        self.assertEqual(list(range(1, 10)), [entry["id"] for entry in registry])
        self.assertEqual({"pending"}, {entry["status"] for entry in registry})
        # Агент разбирает raw_text на поля: здесь поля берутся из исходного реестра по совпадению текста записи.
        known = {fb.format_source(fb.normalize_source(src, 1)): src for src in read_json(original / "sources.json")}
        structured = []
        for entry in registry:
            source = dict(known[entry["raw_text"]])
            source.update(id=entry["id"], status="confirmed")
            structured.append(source)
        (target / "sources.json").write_text(json.dumps(structured, ensure_ascii=False, indent=2), encoding="utf-8")
        proc = run_script("build_vkr.py", target, "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        original_list = [p.text for p in Document(str(original / "final" / "vkr.docx")).paragraphs if p.style.name == "VKR Bibliography"]
        rebuilt_list = [p.text for p in Document(str(target / "final" / "vkr.docx")).paragraphs if p.style.name == "VKR Bibliography"]
        self.assertEqual(original_list, rebuilt_list)

    def test_import_word_saved_docx_skips_toc(self):
        target = new_project(self.folder)
        proc = run_script("import_docx.py", target, "--docx", WORD_FIXTURE, "--json")
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        summary = draft_summary(target)
        expected = draft_summary(FIXTURE)
        self.assertEqual(expected["chapters"], summary["chapters"])
        all_text = "\n".join(p.read_text(encoding="utf-8") for p in (target / "drafts").glob("*.md"))
        self.assertNotIn("Содержание", all_text)
        self.assertNotRegex(all_text, r"\t\d+\n")
        self.assertNotIn("TC \"", all_text)

    def test_update_transfers_paragraph_edit_and_restores_ids(self):
        root = self.built_fixture()
        drafts_before = {p.name: p.read_bytes() for p in (root / "drafts").glob("*.md")}
        sources_before = (root / "sources.json").read_bytes()
        doc = Document(str(root / "final" / "vkr.docx"))
        paragraph = next(p for p in doc.paragraphs if p.text.startswith("Глава описывает"))
        for run in paragraph.runs[1:]:
            run._element.getparent().remove(run._element)
        paragraph.runs[0].text = "Глава описывает психологические и педагогические основания тренажёров [2, с. 7]."
        doc.save(str(root / "final" / "vkr.docx"))
        proc = run_script("import_docx.py", root, "--docx", root / "final" / "vkr.docx", "--update", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        report = json.loads(proc.stdout)
        self.assertEqual(["drafts/chapter-1.md"], report["files"]["updated"])
        chapter = (root / "drafts" / "chapter-1.md").read_text(encoding="utf-8")
        self.assertIn("Глава описывает психологические и педагогические основания тренажёров [@egorov2020, с. 7].", chapter)
        self.assertIn("## 1.1. Пространственное мышление учащихся", chapter)
        self.assertIn("[@yolkin2021; @abramova2022]", chapter)
        self.assertIn("Пространственное мышление рассматривается как способность оперировать образами\nобъектов", chapter)
        for name, content in drafts_before.items():
            if name != "chapter-1.md":
                self.assertEqual(content, (root / "drafts" / name).read_bytes(), name)
        self.assertEqual(sources_before, (root / "sources.json").read_bytes())
        backup = root / report["backup"] / "drafts" / "chapter-1.md"
        self.assertEqual(drafts_before["chapter-1.md"], backup.read_bytes())

    def test_update_with_word_saved_docx_changes_nothing(self):
        root = self.built_fixture()
        drafts_before = {p.name: p.read_bytes() for p in (root / "drafts").glob("*.md")}
        proc = run_script("import_docx.py", root, "--docx", WORD_FIXTURE, "--update", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        report = json.loads(proc.stdout)
        self.assertEqual([], report["files"]["updated"], report)
        self.assertEqual([], report["files"]["created"])
        for name, content in drafts_before.items():
            self.assertEqual(content, (root / "drafts" / name).read_bytes(), name)

    def test_new_import_requires_force_for_filled_project(self):
        root = self.built_fixture()
        proc = run_script("import_docx.py", root, "--docx", WORD_FIXTURE)
        self.assertEqual(2, proc.returncode)
        self.assertIn("--force", proc.stderr)
        self.assertIn("--replace-sources", proc.stderr)
        before = (root / "drafts" / "introduction.md").read_bytes()
        registry_before = (root / "sources.json").read_bytes()
        proc = run_script("import_docx.py", root, "--docx", WORD_FIXTURE, "--force", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        report = json.loads(proc.stdout)
        self.assertEqual(before, (root / report["backup"] / "drafts" / "introduction.md").read_bytes())
        # Реестр фикстуры структурирован и содержит confirmed: --force его не заменяет (правило verify_sources --extract).
        self.assertEqual(registry_before, (root / "sources.json").read_bytes())
        self.assertFalse((root / report["backup"] / "sources.json").exists())
        self.assertTrue(report["sources"]["registry_preserved"])
        self.assertEqual(9, report["sources"]["docx_entries"])
        self.assertEqual("audit/sources-extracted.json", report["sources"]["extracted"])
        extracted = read_json(root / "audit" / "sources-extracted.json")
        self.assertEqual(list(range(1, 10)), [entry["id"] for entry in extracted])
        self.assertEqual({"id", "raw_text", "status", "hints"}, set(extracted[0]))
        self.assertEqual({"pending"}, {entry["status"] for entry in extracted})
        warning = next(w for w in report["warnings"] if "sources.json не заменён" in w)
        self.assertIn("audit/sources-extracted.json", warning)
        self.assertIn("--replace-sources", warning)
        # Явная замена — только --force --replace-sources; прежний реестр в резервной копии.
        proc = run_script("import_docx.py", root, "--docx", WORD_FIXTURE, "--force", "--replace-sources", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        report = json.loads(proc.stdout)
        self.assertEqual(registry_before, (root / report["backup"] / "sources.json").read_bytes())
        self.assertNotIn("registry_preserved", report["sources"])
        registry = read_json(root / "sources.json")
        self.assertEqual(list(range(1, 10)), [entry["id"] for entry in registry])
        self.assertEqual({"id", "raw_text", "status"}, set(registry[0]))

    def test_force_replaces_registry_of_pending_raw_entries(self):
        # Реестр только из записей raw_text/pending (прежний импорт) --force по-прежнему заменяет.
        target = new_project(self.folder)
        self.assertEqual(0, run_script("import_docx.py", target, "--docx", WORD_FIXTURE).returncode)
        (target / "sources.json").write_text(json.dumps(
            [{"id": 1, "raw_text": "Старая запись", "status": "pending", "hints": {"year": "2020"}}], ensure_ascii=False), encoding="utf-8")
        proc = run_script("import_docx.py", target, "--docx", WORD_FIXTURE, "--force", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        report = json.loads(proc.stdout)
        self.assertEqual(9, report["sources"]["imported"])
        self.assertEqual(9, len(read_json(target / "sources.json")))
        self.assertEqual("Старая запись", read_json(target / report["backup"] / "sources.json")[0]["raw_text"])
        self.assertFalse((target / "audit" / "sources-extracted.json").exists())

    def test_registry_protection_rule_matches_verify_sources_extract(self):
        importer = load_module("import_docx")
        verify = load_module("verify_sources")
        cases = {
            "pending-raw": [{"id": 1, "raw_text": "x", "status": "pending", "hints": {"year": "2020"}, "notes": ""}],
            "no-status": [{"id": 1, "raw_text": "x"}],
            "verified": [{"id": 1, "raw_text": "x", "status": "verified"}],
            "suspicious": [{"id": 1, "raw_text": "x", "status": "suspicious"}],
            "rejected": [{"id": 2, "raw_text": "x", "status": "rejected"}],
            "structured": [{"id": "ivanov2020", "type": "book", "title": "T", "status": "pending"}],
            "wrapped": {"sources": [{"id": 1, "raw_text": "x", "status": "pending"}]},
        }
        expected_protected = {"verified", "suspicious", "rejected", "structured", "wrapped"}
        for name, data in cases.items():
            path = self.folder / name / "sources.json"
            path.parent.mkdir()
            path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(name in expected_protected, importer.registry_protection(path) is not None, name)
            self.assertEqual(verify.protected_registry_reason(path) is None, importer.registry_protection(path) is None, name)
        broken = self.folder / "broken" / "sources.json"
        broken.parent.mkdir()
        broken.write_text("[{", encoding="utf-8")
        self.assertIsNotNone(importer.registry_protection(broken))
        self.assertIsNotNone(verify.protected_registry_reason(broken))

    def test_lost_content_gives_exit_code_1(self):
        target = new_project(self.folder)
        doc = Document()
        doc.add_paragraph("Введение", style="Heading 1")
        paragraph = doc.add_paragraph("Формула: ")
        math = etree.SubElement(paragraph._p, "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
        run = etree.SubElement(math, "{http://schemas.openxmlformats.org/officeDocument/2006/math}r")
        etree.SubElement(run, "{http://schemas.openxmlformats.org/officeDocument/2006/math}t").text = "E=mc2"
        doc.add_paragraph("Глава 1. Теория", style="Heading 1")
        doc.add_paragraph("Текст главы.")
        doc.add_paragraph("Заключение", style="Heading 1")
        doc.add_paragraph("Итог.")
        source = self.folder / "foreign.docx"
        doc.save(str(source))
        proc = run_script("import_docx.py", target, "--docx", source, "--json")
        self.assertEqual(1, proc.returncode, proc.stdout)
        report = json.loads(proc.stdout)
        self.assertEqual("equation", report["unrecognized"][0]["kind"])
        self.assertTrue(report["unrecognized"][0]["lost_content"])
        annotation = (target / "drafts" / "annotation.md").read_text(encoding="utf-8")
        self.assertIn("[ЗАПОЛНИТЬ: аннотация не найдена в DOCX]", annotation)
        self.assertIn("Текст главы.", (target / "drafts" / "chapter-1.md").read_text(encoding="utf-8"))
        self.assertFalse((target / "drafts" / "chapter-3.md").exists())

    def test_merge_keeps_comments_and_untouched_blocks(self):
        importer = load_module("import_docx")
        old = ("# Глава I. Теория\n<!-- комментарий автора -->\nПервый абзац\nс переносом [@a].\n\n"
               "Второй абзац.\n\n```python\nx = [1]\n```\n\n- пункт один\n- пункт два\n")
        canon_old = lambda text: importer.canonical_citations(text, lambda kind, ref: ref if kind == "id" else None)
        canon_new = lambda text: importer.canonical_citations(text, lambda kind, ref: {1: "a"}.get(ref) if kind == "num" else ref)
        heading = {"kind": "heading", "level": 1, "text": "Глава I. Теория"}
        listing = {"kind": "listing", "caption": "", "code": "x = [1]", "language": ""}
        items = {"kind": "list", "ordered": False, "items": ["пункт один", "пункт два"]}
        same = [heading, {"kind": "paragraph", "text": "Первый абзац с переносом [1]."},
                {"kind": "paragraph", "text": "Второй абзац."}, listing, items]
        text, changed = importer.merge_draft(old, same, "chapter", canon_old, canon_new)
        self.assertEqual(0, changed)
        self.assertEqual(old, text)
        edited = [heading, {"kind": "paragraph", "text": "Вставленный абзац."},
                  {"kind": "paragraph", "text": "Первый абзац изменён [1]."}, listing, items]
        text, changed = importer.merge_draft(old, edited, "chapter", canon_old, canon_new)
        self.assertEqual(("# Глава I. Теория\n<!-- комментарий автора -->\nВставленный абзац.\n\n"
                          "Первый абзац изменён [@a].\n\n```python\nx = [1]\n```\n\n- пункт один\n- пункт два\n"), text)
        removed = [heading, {"kind": "paragraph", "text": "Первый абзац с переносом [1]."}, listing, items]
        text, changed = importer.merge_draft(old, removed, "chapter", canon_old, canon_new)
        self.assertEqual(1, changed)
        self.assertNotIn("Второй абзац", text)
        self.assertNotIn("\n\n\n", text)
        self.assertIn("Первый абзац\nс переносом [@a].", text)

    def test_foreign_docx_structures_are_recognized(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        target = new_project(self.folder)
        doc = Document()
        doc.add_paragraph("МИНИСТЕРСТВО ПРОСВЕЩЕНИЯ")
        doc.add_paragraph("Тема работы")
        body = doc.element.body
        sdt = parse_xml(
            f'<w:sdt {nsdecls("w")}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
            '<w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent>'
            '<w:p><w:r><w:t>ОГЛАВЛЕНИЕ</w:t></w:r></w:p>'
            '<w:p><w:r><w:t xml:space="preserve">ВВЕДЕНИЕ</w:t></w:r><w:r><w:tab/><w:t>3</w:t></w:r></w:p>'
            '<w:p><w:r><w:t xml:space="preserve">ГЛАВА 1. ТЕОРИЯ</w:t></w:r><w:r><w:tab/><w:t>5</w:t></w:r></w:p>'
            '</w:sdtContent></w:sdt>')
        body.insert(len(body) - 1, sdt)
        intro = doc.add_paragraph("ВВЕДЕНИЕ")
        intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("Текст введения со ссылкой [2, с. 4].")
        chapter = doc.add_paragraph()
        chapter.add_run("ГЛАВА 1. ТЕОРИЯ").bold = True
        doc.add_paragraph("1.1. Первый параграф", style="Heading 2")
        doc.add_paragraph("Задачи:")
        doc.add_paragraph("первая задача", style="List Number")
        doc.add_paragraph("вторая задача", style="List Number")
        for line in ("def f(x):", "    return x[1]"):
            run = doc.add_paragraph().add_run(line)
            run.font.name = "Courier New"
        doc.add_paragraph("Листинг 1 – Функция")
        doc.add_paragraph("Таблица 1 – Сравнение")
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, (("A", "B"), ("1", "2 | 3"))):
            for cell, value in zip(row.cells, values):
                cell.text = value
        conclusion = doc.add_paragraph()
        conclusion.add_run("ЗАКЛЮЧЕНИЕ").bold = True
        doc.add_paragraph("Итоги работы.")
        bib = doc.add_paragraph()
        bib.add_run("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ").bold = True
        doc.add_paragraph("Иванов И.И. Первая книга. – М., 2020.", style="List Number")
        doc.add_paragraph("Петров П.П. Вторая книга. – М., 2021.", style="List Number")
        label = doc.add_paragraph("ПРИЛОЖЕНИЕ А")
        label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        title = doc.add_paragraph()
        title.add_run("Анкета").bold = True
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("Вопросы анкеты.")
        source = self.folder / "foreign.docx"
        doc.save(str(source))

        proc = run_script("import_docx.py", target, "--docx", source, "--json")
        self.assertEqual(0, proc.returncode, proc.stdout + proc.stderr)
        drafts = {p.name: p.read_text(encoding="utf-8") for p in (target / "drafts").glob("*.md")}
        self.assertEqual({"annotation.md", "introduction.md", "chapter-1.md", "conclusion.md", "appendix-1.md"}, set(drafts))
        self.assertNotIn("ОГЛАВЛЕНИЕ", "".join(drafts.values()))
        self.assertIn("Текст введения со ссылкой [2, с. 4].", drafts["introduction.md"])
        chapter1 = drafts["chapter-1.md"]
        self.assertTrue(chapter1.startswith("# Глава I. ТЕОРИЯ"), chapter1)
        self.assertIn("## Первый параграф", chapter1)
        self.assertIn("1. первая задача\n2. вторая задача", chapter1)
        self.assertIn("Листинг: Функция\n\n```\ndef f(x):\n    return x[1]\n```", chapter1)
        self.assertIn("Таблица: Сравнение\n\n| A | B |\n|---|---|\n| 1 | 2 \\| 3 |", chapter1)
        self.assertTrue(drafts["appendix-1.md"].startswith("# Анкета"))
        registry = read_json(target / "sources.json")
        self.assertEqual([1, 2], [entry["id"] for entry in registry])
        self.assertEqual("Иванов И.И. Первая книга. – М., 2020.", registry[0]["raw_text"])
        reparsed = draft_summary(target)
        self.assertEqual([("ТЕОРИЯ", ["Первый параграф"], False)], reparsed["chapters"])

    def test_import_usage_errors(self):
        target = new_project(self.folder)
        self.assertEqual(2, run_script("import_docx.py", target, "--docx", self.folder / "missing.docx").returncode)
        broken = self.folder / "broken.docx"
        broken.write_bytes(b"not a zip")
        self.assertEqual(2, run_script("import_docx.py", target, "--docx", broken).returncode)
        self.assertEqual(2, run_script("import_docx.py", target, "--docx", WORD_FIXTURE, "--update").returncode)
        proc = run_script("import_docx.py", target, "--docx", WORD_FIXTURE, "--update", "--replace-sources", "--json")
        self.assertEqual(2, proc.returncode)
        self.assertIn("--replace-sources", json.loads(proc.stdout)["errors"][0])
        proc = run_script("import_docx.py", target, "--docx", WORD_FIXTURE, "--overwrite-drafts", "--json")
        self.assertEqual(2, proc.returncode)
        self.assertIn("--overwrite-drafts", json.loads(proc.stdout)["errors"][0])


# ---------------------------------------------------------------------------
# Fix-A (приёмка V2/V1): запись сборки, сверка импорта с последней сборкой
# ---------------------------------------------------------------------------

AAKER = {"id": "aaker2024", "type": "book", "authors": ["Аакер Д.А."], "title": "Стратегическое планирование",
         "place": "М.", "publisher": "Наука", "year": 2024, "status": "confirmed",
         "verification": {"method": "catalog", "checked_at": "2026-09-01T10:00:00Z", "note": "РГБ"}}


class word_like_lock:
    """Держит файл открытым так, как Word держит открытый документ: чтение и запись, общий доступ только на чтение."""

    def __init__(self, path: Path) -> None:
        self.path = Path(path)
        self.handle = None

    def __enter__(self):
        import ctypes
        from ctypes import wintypes

        self.kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        create = self.kernel32.CreateFileW
        create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                           wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        create.restype = wintypes.HANDLE
        self.kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
        handle = create(str(self.path), 0x80000000 | 0x40000000, 0x1, None, 3, 0x80, None)
        if handle is None or handle == ctypes.c_void_p(-1).value:
            raise OSError(ctypes.get_last_error(), f"не удалось открыть {self.path}")
        self.handle = handle
        return self

    def __exit__(self, *exc):
        self.kernel32.CloseHandle(self.handle)
        return False


def add_first_source(root: Path) -> None:
    """Новый источник, встающий в списке первым, и ссылка на него: номера всех записей сдвигаются."""
    sources = read_json(root / "sources.json")
    sources.append(dict(AAKER))
    (root / "sources.json").write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding="utf-8")
    chapter = root / "drafts" / "chapter-2.md"
    chapter.write_text(chapter.read_text(encoding="utf-8") + "\nСтратегия внедрения описана в пособии [@aaker2024, с. 3].\n",
                       encoding="utf-8")


def edit_docx_paragraph(path: Path, startswith: str, text: str) -> None:
    doc = Document(str(path))
    paragraph = next(p for p in doc.paragraphs if p.text.startswith(startswith))
    for run in paragraph.runs[1:]:
        run._element.getparent().remove(run._element)
    paragraph.runs[0].text = text
    doc.save(str(path))


def drafts_bytes(root: Path) -> dict:
    return {p.name: p.read_bytes() for p in sorted((root / "drafts").glob("*.md"))}


def leftovers(root: Path) -> list:
    return sorted(p.relative_to(root).as_posix() for folder in ("final", "exports") if (root / folder).is_dir()
                  for p in (root / folder).iterdir() if p.name.endswith(".tmp") or p.name.startswith(".vkr-"))


class BuildWriteOrderTest(unittest.TestCase):
    """A1: служебные файлы сборки пишутся только после DOCX; --content-only их не трогает."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = copy_fixture(Path(self.tmp.name))
        proc = run_script("build_vkr.py", self.root, "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        self.first = json.loads(proc.stdout)

    def tearDown(self):
        self.tmp.cleanup()

    TRACKED = ("final/vkr.docx", "exports/citation-map.json", "exports/vkr-content.json", "exports/build-manifest.json")

    def snapshot(self) -> dict:
        return {rel: sha(self.root / rel) for rel in self.TRACKED}

    def test_citation_map_has_build_id_and_entry_texts(self):
        citation_map = read_json(self.root / "exports" / "citation-map.json")
        manifest = read_json(self.root / "exports" / "build-manifest.json")
        self.assertTrue(citation_map["build_id"])
        self.assertEqual(citation_map["build_id"], manifest["build_id"])
        self.assertEqual(self.first["build_id"], manifest["build_id"])
        self.assertEqual("final/vkr.docx", manifest["docx"]["path"])
        self.assertEqual(sorted(citation_map["numbers"]), sorted(citation_map["entries"]))
        bibliography = [p.text for p in Document(str(self.root / "final" / "vkr.docx")).paragraphs if p.style.name == "VKR Bibliography"]
        self.assertEqual(bibliography, [f"{n}. {citation_map['entries'][str(n)]}" for n in range(1, len(bibliography) + 1)])
        self.assertEqual([], common_manifest_problems(self.root))
        self.assertEqual([], leftovers(self.root))

    @unittest.skipUnless(os.name == "nt", "блокировка файла, как у открытого в Word документа, — только Windows")
    def test_docx_open_in_word_refuses_before_any_write(self):
        add_first_source(self.root)
        before = self.snapshot()
        backups = sorted(p.name for p in (self.root / "backups" / "checkpoints").iterdir()) if (self.root / "backups" / "checkpoints").is_dir() else []
        with word_like_lock(self.root / "final" / "vkr.docx"):
            proc = run_script("build_vkr.py", self.root, "--json")
            self.assertEqual(2, proc.returncode, proc.stdout)
            report = json.loads(proc.stdout)
            self.assertEqual(["OUTPUT_LOCKED"], [error["code"] for error in report["errors"]], report["errors"])
            self.assertIn("закрой файл в Word", report["errors"][0]["message"])
            self.assertIn("ничего не изменено", report["errors"][0]["message"])
            self.assertEqual({}, report["outputs"])
        self.assertEqual(before, self.snapshot(), "при открытом DOCX сборка не должна менять ни DOCX, ни карту, ни манифест")
        after_backups = sorted(p.name for p in (self.root / "backups" / "checkpoints").iterdir()) if (self.root / "backups" / "checkpoints").is_dir() else []
        self.assertEqual(backups, after_backups)
        self.assertEqual([], leftovers(self.root))
        proc = run_script("build_vkr.py", self.root, "--json")  # файл закрыт — сборка проходит
        self.assertEqual(0, proc.returncode, proc.stdout)
        citation_map = read_json(self.root / "exports" / "citation-map.json")
        self.assertEqual("aaker2024", citation_map["numbers"]["1"])
        self.assertEqual(read_json(self.root / "exports" / "build-manifest.json")["build_id"], citation_map["build_id"])

    @unittest.skipUnless(os.name == "nt", "блокировка файла — только Windows")
    def test_locked_citation_map_refuses_before_docx(self):
        add_first_source(self.root)
        before = self.snapshot()
        with word_like_lock(self.root / "exports" / "citation-map.json"):
            proc = run_script("build_vkr.py", self.root, "--json")
        self.assertEqual(2, proc.returncode, proc.stdout)
        self.assertIn("exports/citation-map.json", json.loads(proc.stdout)["errors"][0]["message"])
        self.assertEqual(before, self.snapshot())

    def test_content_only_does_not_touch_last_build_files(self):
        add_first_source(self.root)
        before = self.snapshot()
        proc = run_script("build_vkr.py", self.root, "--content-only", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        report = json.loads(proc.stdout)
        self.assertEqual({"content": "exports/vkr-content-preview.json"}, report["outputs"])
        self.assertEqual([], report["next_steps"])
        self.assertEqual(before, self.snapshot(), "--content-only не должен перезаписывать файлы последней сборки")
        preview = read_json(self.root / "exports" / "vkr-content-preview.json")
        self.assertTrue(preview["bibliography"][0].startswith("1. Аакер"), preview["bibliography"][0])
        proc = run_script("build_vkr.py", self.root, "--output", "exports/preview.docx", "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        self.assertEqual(before, self.snapshot(), "промежуточный DOCX не меняет карту и манифест final/vkr.docx")

    def test_multiple_citation_numbers_ascending(self):
        # A7: [@yolkin2021; @abramova2022] → «[1; 3]», а не «[3; 1]».
        texts = [p.text for p in Document(str(self.root / "final" / "vkr.docx")).paragraphs]
        self.assertTrue(any(text.endswith("этой способности [1; 3].") for text in texts), texts)
        self.assertEqual("[1; 3]", build.sort_citation_numbers("[3; 1]"))
        self.assertEqual("см. [см. 2; 9, с. 5] и [1–3; 7]", build.sort_citation_numbers("см. [см. 9, с. 5; 2] и [7; 1–3]"))
        for unchanged in ("[3, 5]", "[Цит. по: 4; см. 2]", "`[3; 1]`", "[@b; @a]", "[2; 4]"):
            self.assertEqual(unchanged, build.sort_citation_numbers(unchanged), unchanged)

    def test_stale_verification_warning(self):
        # A9: реквизиты confirmed-источника изменены после verify_sources.py --mark.
        verify = load_module("verify_sources")
        sources = read_json(self.root / "sources.json")
        record = next(item for item in sources if item["id"] == 3)
        record["verification"]["fields_sha256"] = verify.bibliographic_fields_sha256(record)
        (self.root / "sources.json").write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding="utf-8")
        report = json.loads(run_script("build_vkr.py", self.root, "--json").stdout)
        self.assertNotIn("SOURCE_VERIFICATION_STALE", {w["code"] for w in report["warnings"]})
        record["year"] = 2018
        (self.root / "sources.json").write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding="utf-8")
        proc = run_script("build_vkr.py", self.root, "--json")
        self.assertEqual(0, proc.returncode, proc.stdout)
        stale = [w for w in json.loads(proc.stdout)["warnings"] if w["code"] == "SOURCE_VERIFICATION_STALE"]
        self.assertEqual(1, len(stale), stale)
        self.assertRegex(stale[0]["message"], r": 3$")


def common_manifest_problems(root: Path) -> list:
    import vkr_common  # noqa: WPS433

    return vkr_common.build_manifest_problems(root)


class ImportUpdateGuardTest(unittest.TestCase):
    """A1, A3: --update не портит ссылки и не затирает черновики, изменённые после сборки."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.root = copy_fixture(self.folder)
        proc = run_script("build_vkr.py", self.root)
        self.assertEqual(0, proc.returncode, proc.stderr)
        self.docx = self.root / "final" / "vkr.docx"

    def tearDown(self):
        self.tmp.cleanup()

    def update(self, docx: Path, *extra: str):
        proc = run_script("import_docx.py", self.root, "--docx", docx, "--update", *extra, "--json")
        return proc.returncode, json.loads(proc.stdout)

    def test_docx_from_previous_build_is_refused(self):
        # Сценарий lockcopy2/3: правка в Word идёт по старой сборке, а сборка с новым источником уже прошла.
        edited = self.folder / "word-copy.docx"
        shutil.copyfile(str(self.docx), str(edited))
        edit_docx_paragraph(edited, "Глава описывает", "Глава описывает основания тренажёров [2, с. 7].")
        add_first_source(self.root)
        self.assertEqual(0, run_script("build_vkr.py", self.root).returncode)
        drafts_before, sources_before = drafts_bytes(self.root), (self.root / "sources.json").read_bytes()
        code, report = self.update(edited)
        self.assertEqual(1, code, report)
        self.assertEqual("fail", report["status"])
        self.assertIn("список литературы в DOCX не совпадает с последней сборкой", report["errors"][0])
        self.assertIn("пересобери", report["errors"][0])
        self.assertTrue(report["sources"]["differences"])
        self.assertEqual(drafts_before, drafts_bytes(self.root), "черновики не должны меняться")
        self.assertEqual(sources_before, (self.root / "sources.json").read_bytes())
        self.assertEqual([], report["files"]["updated"])
        self.assertFalse((self.root / "exports" / "import-sync.json").exists())

    def test_bibliography_edited_in_word_is_refused(self):
        doc = Document(str(self.docx))
        entry = next(p for p in doc.paragraphs if p.style.name == "VKR Bibliography" and p.text.startswith("2. "))
        entry.runs[0].text = entry.runs[0].text.replace("Питер", "Питер Пресс")
        doc.save(str(self.docx))
        edit_docx_paragraph(self.docx, "Глава описывает", "Глава описывает основания тренажёров.")
        drafts_before = drafts_bytes(self.root)
        code, report = self.update(self.docx)
        self.assertEqual(1, code, report)
        self.assertTrue(any(item.startswith("запись 2 отличается") for item in report["sources"]["differences"]), report)
        self.assertEqual(drafts_before, drafts_bytes(self.root))

    def test_citation_map_of_previous_version_is_refused(self):
        citation_map = read_json(self.root / "exports" / "citation-map.json")
        (self.root / "exports" / "citation-map.json").write_text(json.dumps({"numbers": citation_map["numbers"]}), encoding="utf-8")
        edit_docx_paragraph(self.docx, "Глава описывает", "Глава описывает основания тренажёров.")
        drafts_before = drafts_bytes(self.root)
        code, report = self.update(self.docx)
        self.assertEqual(1, code, report)
        self.assertIn("прежней версией build_vkr.py", report["errors"][0])
        self.assertEqual(drafts_before, drafts_bytes(self.root))

    def test_forged_numbers_in_citation_map_are_refused(self):
        # T4: номера переставлены на другие id (entries и DOCX не тронуты) — прежде --update молча
        # менял [@egorov2020] на [@abramova2022].
        path = self.root / "exports" / "citation-map.json"
        citation_map = read_json(path)
        numbers = dict(citation_map["numbers"])
        first, second = sorted(numbers, key=int)[:2]
        numbers[first], numbers[second] = numbers[second], numbers[first]
        citation_map["numbers"] = numbers
        path.write_text(json.dumps(citation_map, ensure_ascii=False), encoding="utf-8")
        edit_docx_paragraph(self.docx, "Глава описывает", "Глава описывает основания тренажёров.")
        drafts_before = drafts_bytes(self.root)
        code, report = self.update(self.docx)
        self.assertEqual(1, code, report)
        self.assertEqual("fail", report["status"])
        self.assertIn("карта ссылок не соответствует реестру", report["errors"][0])
        self.assertIn("пересобери", report["errors"][0])
        self.assertEqual(2, len(report["sources"]["citation_map_problems"]), report["sources"])
        self.assertEqual(drafts_before, drafts_bytes(self.root), "черновики не должны меняться")
        self.assertEqual([], report["files"]["updated"])

    def test_unknown_id_in_citation_map_is_refused(self):
        path = self.root / "exports" / "citation-map.json"
        citation_map = read_json(path)
        citation_map["numbers"]["1"] = "нет-такого-источника"
        path.write_text(json.dumps(citation_map, ensure_ascii=False), encoding="utf-8")
        drafts_before = drafts_bytes(self.root)
        code, report = self.update(self.docx)
        self.assertEqual(1, code, report)
        self.assertIn("которого нет в sources.json", report["sources"]["citation_map_problems"][0])
        self.assertEqual(drafts_before, drafts_bytes(self.root))

    def test_honest_update_passes_the_citation_map_check(self):
        edit_docx_paragraph(self.docx, "Глава описывает", "Глава описывает основания тренажёров.")
        code, report = self.update(self.docx)
        self.assertEqual(0, code, report)
        self.assertNotIn("citation_map_problems", report["sources"])

    def test_draft_changed_after_build_is_not_overwritten(self):
        # Сценарий copy4: черновик поправлен после сборки, DOCX не правили.
        chapter = self.root / "drafts" / "chapter-1.md"
        edited = chapter.read_text(encoding="utf-8").replace("Глава описывает", "Глава подробно описывает")
        chapter.write_text(edited, encoding="utf-8")
        edited_bytes = chapter.read_bytes()
        code, report = self.update(self.docx)
        self.assertEqual(1, code, report)
        self.assertEqual(["drafts/chapter-1.md (изменён после сборки)"], report["drafts_changed_after_build"])
        self.assertIn("--overwrite-drafts", report["errors"][0])
        self.assertIn("пересобери", report["errors"][0])
        self.assertEqual(edited_bytes, chapter.read_bytes(), "правка черновика не должна пропасть")
        code, report = self.update(self.docx, "--overwrite-drafts")
        self.assertEqual(0, code, report)
        self.assertEqual(["drafts/chapter-1.md"], report["files"]["updated"])
        self.assertIn("Глава описывает психолого", chapter.read_text(encoding="utf-8"))
        self.assertEqual(edited_bytes, (self.root / report["backup"] / "drafts" / "chapter-1.md").read_bytes())
        self.assertTrue(any(w.startswith("--overwrite-drafts") for w in report["warnings"]), report["warnings"])

    def test_repeated_update_before_rebuild_and_later_manual_edit(self):
        edit_docx_paragraph(self.docx, "Глава описывает", "Глава описывает основания тренажёров (правка 1).")
        code, report = self.update(self.docx)
        self.assertEqual(0, code, report)
        self.assertEqual(["drafts/chapter-1.md"], report["files"]["updated"])
        edit_docx_paragraph(self.docx, "Глава описывает", "Глава описывает основания тренажёров (правка 2).")
        code, report = self.update(self.docx)
        self.assertEqual(0, code, report)  # черновик изменён прошлым --update этой же сборки — не конфликт
        chapter = self.root / "drafts" / "chapter-1.md"
        self.assertIn("(правка 2).", chapter.read_text(encoding="utf-8"))
        chapter.write_text(chapter.read_text(encoding="utf-8") + "\nРучная правка после импорта.\n", encoding="utf-8")
        edit_docx_paragraph(self.docx, "Глава описывает", "Глава описывает основания тренажёров (правка 3).")
        code, report = self.update(self.docx)
        self.assertEqual(1, code, report)
        self.assertIn("Ручная правка после импорта.", chapter.read_text(encoding="utf-8"))

    def test_lost_draft_is_recreated_without_override(self):
        # failure-recovery.md: черновик потерян, final/vkr.docx цел — --update создаёт его заново.
        original = (self.root / "drafts" / "chapter-2.md").read_text(encoding="utf-8")
        (self.root / "drafts" / "chapter-2.md").unlink()
        (self.root / "exports" / "build-manifest.json").unlink()  # и манифест потерян: перезаписывать нечего
        code, report = self.update(self.docx)
        self.assertEqual(0, code, report)
        self.assertEqual(["drafts/chapter-2.md"], report["files"]["created"])
        recreated = (self.root / "drafts" / "chapter-2.md").read_text(encoding="utf-8")
        self.assertIn("Сравнение VR-платформ", recreated)
        self.assertIn("[@smith2019, p. 105]", recreated)
        self.assertIn("[@smith2019, p. 105]", original)
        self.assertEqual(0, run_script("build_vkr.py", self.root).returncode)
        (self.root / "drafts" / "chapter-3.md").unlink()
        code, report = self.update(self.docx)
        self.assertEqual(0, code, report)
        self.assertTrue(any("удалены после сборки" in w and "chapter-3.md" in w for w in report["warnings"]), report["warnings"])

    def test_untitled_listings_round_trip_without_split_warning(self):
        # Подпись «Листинг N» без названия (листинг без «Листинг: …» в черновике) — всё равно граница листинга.
        appendix = self.root / "drafts" / "appendix-1.md"
        appendix.write_text(appendix.read_text(encoding="utf-8") + "\n```\nimport csv\n\nimport json\n```\n\n```\nprint(1)\n```\n",
                            encoding="utf-8")
        self.assertEqual(0, run_script("build_vkr.py", self.root).returncode)
        drafts_before = drafts_bytes(self.root)
        code, report = self.update(self.docx)
        self.assertEqual(0, code, report)
        self.assertEqual([], report["files"]["updated"])
        self.assertEqual(drafts_before, drafts_bytes(self.root))
        self.assertFalse([w for w in report["warnings"] if "листинг" in w], report["warnings"])

    def test_sorted_numbers_do_not_count_as_edit(self):
        # A7 + --update: «[1; 3]» в DOCX и «[@yolkin2021; @abramova2022]» в черновике — одна ссылка.
        drafts_before = drafts_bytes(self.root)
        code, report = self.update(self.docx)
        self.assertEqual(0, code, report)
        self.assertEqual([], report["files"]["updated"])
        self.assertEqual(drafts_before, drafts_bytes(self.root))


def student_docx(path: Path, bibliography: list, code_lines=None, picture: bool = False) -> Path:
    """DOCX «от студента»: разделы заголовками, ссылки [N], список литературы абзацами."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_paragraph("Введение", style="Heading 1")
    doc.add_paragraph("Актуальность темы подтверждают работы [1] и [7, с. 5].")
    doc.add_paragraph("Глава 1. Теоретические основы", style="Heading 1")
    doc.add_paragraph("Теория опирается на исследования [6; 7].")
    if picture:
        doc.add_picture(str(FIXTURE / "evidence" / "figures" / "architecture.png"), width=Mm(40))
        doc.add_paragraph("Рисунок 1 — Архитектура")
    for line in code_lines or []:
        paragraph = doc.add_paragraph()
        if isinstance(line, tuple):  # (текст, шрифт)
            paragraph.add_run(line[0]).font.name = line[1]
        elif line:
            paragraph.add_run(line).font.name = "Consolas"
        else:  # пустая строка кода: шрифт у Word только в знаке абзаца
            mark = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), "Consolas")
            fonts.set(qn("w:hAnsi"), "Consolas")
            mark.append(fonts)
            paragraph._p.get_or_add_pPr().append(mark)
    doc.add_paragraph("Заключение", style="Heading 1")
    doc.add_paragraph("Итоги работы.")
    doc.add_paragraph("Список литературы", style="Heading 1")
    for entry in bibliography:
        if isinstance(entry, tuple):  # (текст, полужирный)
            doc.add_paragraph().add_run(entry[0]).bold = entry[1]
        else:
            doc.add_paragraph(entry)
    doc.save(str(path))
    return path


STUDENT_LIST = [
    "1. Андреев, А.А. Дистанционное обучение: сущность, технология, организация [Текст] / А.А. Андреев. – М.: МЭСИ, 1999. – 196 с.",
    "2. Босова, Л.Л. Информатика. 7 класс [Текст]: учебник / Л.Л. Босова. – М.: БИНОМ, 2019. – 224 с.",
    "3. Полат, Е.С. Современные педагогические технологии [Текст] / Е.С. Полат. – М.: Академия, 2007. – 368 с.",
    "4. Об образовании в Российской Федерации [Электронный ресурс]: федер. закон от 29.12.2012 № 273-ФЗ.",
    "5. Moodle Docs [Электронный ресурс]. – URL: https://docs.moodle.org/ (дата обращения: 10.03.2027).",
    "6. Роберт, И.В. Теория и методика информатизации образования (психолого-педагогический и",
    "технологический аспекты) [Текст] / И.В. Роберт. – М.: БИНОМ. Лаборатория знаний, 2014. – 398 с.",
    "7. Кравченко, Г.В. Использование LMS Moodle в учебном процессе вуза [Текст] / Г.В. Кравченко // Известия АлтГУ. – 2016. – № 2. – С. 30–34.",
]


class ImportBibliographyAndCodeTest(unittest.TestCase):
    """A2: тот же разбор списка литературы, что у verify_sources --extract; A6: листинги с пустыми строками."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.target = new_project(self.folder)

    def tearDown(self):
        self.tmp.cleanup()

    def run_import(self, source: Path):
        proc = run_script("import_docx.py", self.target, "--docx", source, "--json")
        return proc.returncode, json.loads(proc.stdout)

    def test_broken_entry_is_joined_and_numbers_match_extract(self):
        source = student_docx(self.folder / "student.docx", STUDENT_LIST)
        code, report = self.run_import(source)
        self.assertEqual(0, code, report)
        registry = read_json(self.target / "sources.json")
        self.assertEqual(list(range(1, 8)), [entry["id"] for entry in registry])
        self.assertTrue(registry[5]["raw_text"].endswith("Лаборатория знаний, 2014. – 398 с."), registry[5])
        self.assertTrue(registry[6]["raw_text"].startswith("Кравченко"), registry[6])
        self.assertTrue(any("приклеен к предыдущей записи" in w for w in report["warnings"]), report["warnings"])
        extracted = load_module("verify_sources").extract_bibliography(source)
        self.assertEqual([(r["id"], r["raw_text"]) for r in extracted["records"]],
                         [(r["id"], r["raw_text"]) for r in registry])
        self.assertNotIn("1007", json.dumps(report, ensure_ascii=False))

    def test_number_collision_and_stray_paragraph_stop_import(self):
        drafts_before = drafts_bytes(self.target)
        collision = STUDENT_LIST[:5] + ["6. Роберт, И.В. Теория и методика информатизации образования [Текст]. – М., 2014.",
                                        "7. Горбунова, Т.Н. Сравнительный анализ СДО [Текст] // Информатика и образование. – 2020.",
                                        STUDENT_LIST[7]]
        stray = STUDENT_LIST[:5] + ["Роберт, И.В. Теория и методика информатизации образования [Текст]. – М., 2014."] + STUDENT_LIST[7:]
        for name, entries, needle in (("collision", collision, "номер 7 стоит у 2 записей"),
                                      ("stray", stray, "абзац без номера «Роберт")):
            with self.subTest(name):
                source = student_docx(self.folder / f"{name}.docx", entries, picture=True)
                code, report = self.run_import(source)
                self.assertEqual(1, code, report)
                self.assertEqual("fail", report["status"])
                self.assertTrue(any(needle in problem for problem in report["sources"]["problems"]), report["sources"])
                self.assertIn("файлы проекта не изменены", report["errors"][0])
                self.assertEqual(drafts_before, drafts_bytes(self.target))
                self.assertEqual("[]\n", (self.target / "sources.json").read_text(encoding="utf-8"))
                self.assertFalse((self.target / "evidence" / "figures").exists(), "рисунки не пишутся при отказе")
                self.assertNotIn("1007", json.dumps(report, ensure_ascii=False))

    def test_bold_group_heading_inside_bibliography_is_not_an_entry(self):
        entries = [("Нормативные правовые акты", True), STUDENT_LIST[3].replace("4. ", "1. "),
                   ("ЭЛЕКТРОННЫЕ РЕСУРСЫ", False), STUDENT_LIST[4].replace("5. ", "2. ")]
        code, report = self.run_import(student_docx(self.folder / "groups.docx", entries))
        self.assertEqual(0, code, report)
        self.assertEqual([1, 2], [entry["id"] for entry in read_json(self.target / "sources.json")])
        self.assertEqual(2, sum(1 for w in report["bibliography"]["warnings"] if w["code"] == "SUBHEADING_SKIPPED"))

    def test_blank_lines_inside_monospace_code_keep_one_listing(self):
        code_lines = ["import csv", "", "MOODLE_URL = 'https://moodle.example.ru'", "", "",
                      "def get_grades(course_id):", "    return course_id", "",
                      ("Этот абзац набран обычным шрифтом", "Times New Roman"), "print(get_grades(1))"]
        code, report = self.run_import(student_docx(self.folder / "code.docx", STUDENT_LIST[:5], code_lines=code_lines))
        self.assertEqual(0, code, report)
        chapter = (self.target / "drafts" / "chapter-1.md").read_text(encoding="utf-8")
        self.assertIn("```\nimport csv\n\nMOODLE_URL = 'https://moodle.example.ru'\n\n\ndef get_grades(course_id):\n"
                      "    return course_id\n```", chapter)
        self.assertEqual(4, chapter.count("```"), chapter)  # два листинга: до и после абзаца обычным шрифтом
        self.assertTrue(any("листинг разделён абзацем" in w for w in report["warnings"]), report["warnings"])


if __name__ == "__main__":
    unittest.main()
