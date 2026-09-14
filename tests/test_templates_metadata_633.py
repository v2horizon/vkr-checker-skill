"""WP4 templates-metadata -- тесты для 6.33-production.

unittest, без сети, без Word во время запуска (число страниц из Word
записано заранее в tests/fixtures/templates/word-pages.json и здесь только
сверяется по sha256 файла). Работает на PY39 и PYCODEX.

Запуск:
    python -m unittest discover -s tests -t .
    python tests/test_templates_metadata_633.py
"""
from __future__ import annotations

import contextlib
import hashlib
import importlib.util
import io
import json
import os
import re
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / "vkr-mpgu"
ASSETS = SKILL / "assets"
FIXTURES = Path(__file__).resolve().parent / "fixtures" / "templates"


def load_module(name):
    spec = importlib.util.spec_from_file_location(name, SKILL / "scripts" / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


cleaner = load_module("clean_docx_metadata")
integrity = load_module("docx_integrity")

# Шаблоны WP4 (без официальных methodology-09-03-02-2024.docx и
# implementation-certificate-imo-2024.docx -- их WP4 не имеет права менять).
OWNED_TEMPLATES = [
    "template-mpgu.docx",
    "titlepage-universal.docx",
    "titlepage-09-03-02-imo.docx",
    "oath-template.docx",
    "oath-template-collective.docx",
    "permission-template.docx",
    "implementation-certificate.docx",
]

EMOJI_RE = re.compile(
    "[\U0001F300-\U0001FAFF\U00002600-\U000027BF\U0001F000-\U0001F0FF←-⇿✀-➿]"
)

EMU_PER_MM = 36000  # 914400 EMU/in / 25.4 mm/in


def _mm(length):
    return length / EMU_PER_MM


def _all_paragraph_texts(doc):
    texts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    return texts


# ---------------------------------------------------------------------------
# A. Геометрия и гигиена шаблонов (без Word)
# ---------------------------------------------------------------------------

class TemplateGeometryTests(unittest.TestCase):

    def test_all_owned_templates_exist(self):
        for name in OWNED_TEMPLATES:
            self.assertTrue((ASSETS / name).exists(), name)

    def test_a4_and_margins(self):
        for name in OWNED_TEMPLATES:
            with self.subTest(template=name):
                doc = Document(str(ASSETS / name))
                sec = doc.sections[0]
                self.assertAlmostEqual(_mm(sec.page_width), 210.0, delta=0.5)
                self.assertAlmostEqual(_mm(sec.page_height), 297.0, delta=0.5)
                self.assertAlmostEqual(_mm(sec.left_margin), 35.0, delta=0.5)
                self.assertAlmostEqual(_mm(sec.right_margin), 10.0, delta=0.5)
                self.assertAlmostEqual(_mm(sec.top_margin), 20.0, delta=0.5)
                self.assertAlmostEqual(_mm(sec.bottom_margin), 20.0, delta=0.5)

    def test_official_templates_untouched_margins(self):
        # Не наши файлы -- только читаем, чтобы убедиться, что WP4 их не
        # трогал (поля официальных документов другие: 30/15/20/20).
        for name in ("methodology-09-03-02-2024.docx", "implementation-certificate-imo-2024.docx"):
            doc = Document(str(ASSETS / name))
            sec = doc.sections[0]
            self.assertAlmostEqual(_mm(sec.left_margin), 30.0, delta=0.5)
            self.assertAlmostEqual(_mm(sec.right_margin), 15.0, delta=0.5)

    def test_times_new_roman_throughout(self):
        for name in OWNED_TEMPLATES:
            with self.subTest(template=name):
                doc = Document(str(ASSETS / name))
                normal = doc.styles["Normal"]
                self.assertEqual(normal.font.name, "Times New Roman")

    def test_no_emoji_or_service_lines(self):
        service_markers = ("TODO", "FIXME", "XXX-", "�")
        for name in OWNED_TEMPLATES:
            with self.subTest(template=name):
                doc = Document(str(ASSETS / name))
                for text in _all_paragraph_texts(doc):
                    self.assertIsNone(EMOJI_RE.search(text), (name, text))
                    for marker in service_markers:
                        self.assertNotIn(marker, text, (name, text))



class TitlePageRegressionTests(unittest.TestCase):
    """UNF-D-11 / R2-08: титульные листы не должны расползаться на лишние
    страницы (одиночная лесенка пустых абзацев + завершающий разрыв
    страницы). Эти тесты падают при откате исправления WP4."""

    def _paragraphs(self, name):
        return Document(str(ASSETS / name)).paragraphs

    def test_no_trailing_forced_page_break(self):
        for name in ("titlepage-universal.docx", "titlepage-09-03-02-imo.docx"):
            with self.subTest(template=name):
                paras = self._paragraphs(name)
                last = paras[-1]
                has_break = any(
                    br.get(qn("w:type")) == "page"
                    for br in last._p.findall(".//" + qn("w:br"))
                )
                self.assertFalse(has_break, f"{name}: последний абзац содержит разрыв страницы")

    def test_no_duplicated_blank_paragraphs(self):
        for name in ("titlepage-universal.docx", "titlepage-09-03-02-imo.docx"):
            with self.subTest(template=name):
                texts = [p.text for p in self._paragraphs(name)]
                for i in range(len(texts) - 1):
                    self.assertFalse(
                        texts[i] == "" and texts[i + 1] == "",
                        f"{name}: два пустых абзаца подряд ({i}, {i + 1})",
                    )

    def test_single_line_spacing_everywhere(self):
        for name in ("titlepage-universal.docx", "titlepage-09-03-02-imo.docx"):
            with self.subTest(template=name):
                for i, p in enumerate(self._paragraphs(name)):
                    pf = p.paragraph_format
                    self.assertEqual(pf.line_spacing, 1.0, (name, i))
                    self.assertEqual(pf.space_after, Pt(0), (name, i))

    def test_imo_chair_block_uses_bracket_hints(self):
        texts = [p.text for p in self._paragraphs("titlepage-09-03-02-imo.docx")]
        joined = "\n".join(texts)
        self.assertIn("[название кафедры]", joined)
        self.assertIn("[степень/звание]", joined)
        # два формально одинаковых "[И.О. Фамилия]" -- для руководителя и
        # для заведующего кафедрой -- это ожидаемо.
        self.assertGreaterEqual(texts.count("[И.О. Фамилия]"), 2)


class CertificateRegressionTests(unittest.TestCase):
    """REG-D-10: implementation-certificate.docx не должен заканчиваться
    пустыми абзацами (пустая 4-я страница)."""

    def test_no_trailing_empty_paragraphs(self):
        doc = Document(str(ASSETS / "implementation-certificate.docx"))
        paras = doc.paragraphs
        self.assertNotEqual(paras[-1].text.strip(), "")

    def test_underscore_blanks_fit_text_width(self):
        # R2-08 suggested_fix: строки подчёркиваний короче 165 мм (измерено
        # эмпирически в Word: ~66 симв. пустой строки на TNR14).
        for name in ("permission-template.docx", "implementation-certificate.docx"):
            with self.subTest(template=name):
                doc = Document(str(ASSETS / name))
                for p in doc.paragraphs:
                    for run in p.runs:
                        if set(run.text) == {"_"}:
                            self.assertLessEqual(len(run.text), 66, (name, p.text))


class TemplateMpguStyleTests(unittest.TestCase):
    """Стиль «Обычный» в template-mpgu.docx + колонтитул (SPEC / задание A)."""

    def setUp(self):
        self.doc = Document(str(ASSETS / "template-mpgu.docx"))

    def test_normal_style(self):
        pf = self.doc.styles["Normal"].paragraph_format
        self.assertEqual(pf.space_after, Pt(0))
        self.assertEqual(pf.line_spacing, 1.5)
        self.assertEqual(self.doc.styles["Normal"].font.size, Pt(14))

    def test_normal_style_xml_values(self):
        styles_el = self.doc.styles.element
        # Ищем стиль вручную, а не через XPath-предикат: так надёжнее между
        # версиями lxml, которые тянут за собой python-docx 1.0.1 и 1.2.0.
        normal = None
        for style_el in styles_el.findall(qn("w:style")):
            if style_el.get(qn("w:styleId")) == "Normal":
                normal = style_el
                break
        self.assertIsNotNone(normal)
        spacing = normal.find(qn("w:pPr")).find(qn("w:spacing"))
        self.assertEqual(spacing.get(qn("w:after")), "0")
        self.assertEqual(spacing.get(qn("w:line")), "360")
        ind = normal.find(qn("w:pPr")).find(qn("w:ind"))
        self.assertEqual(ind.get(qn("w:firstLine")), "709")
        jc = normal.find(qn("w:pPr")).find(qn("w:jc"))
        self.assertEqual(jc.get(qn("w:val")), "both")

    def test_titlepg_and_page_footer(self):
        sectPr = self.doc.sections[0]._sectPr
        self.assertIsNotNone(sectPr.find(qn("w:titlePg")))
        footer_ref = sectPr.find(qn("w:footerReference"))
        self.assertIsNotNone(footer_ref)
        self.assertEqual(footer_ref.get(qn("w:type")), "default")
        rid = footer_ref.get(qn("r:id"))
        footer_part = self.doc.part.rels[rid].target_part
        footer_xml = footer_part._element
        instr_texts = [
            n.text or "" for n in footer_xml.iter() if n.tag == qn("w:instrText")
        ]
        self.assertTrue(any("PAGE" in t for t in instr_texts))
        jc = footer_xml.find(".//" + qn("w:jc"))
        self.assertEqual(jc.get(qn("w:val")), "center")


# ---------------------------------------------------------------------------
# B. Замер страниц в Word -- сверяем sha256, не пересчитываем страницы
# ---------------------------------------------------------------------------

class WordPageCountLedgerTests(unittest.TestCase):

    def setUp(self):
        with open(FIXTURES / "word-pages.json", "r", encoding="utf-8-sig") as fh:
            self.ledger = json.load(fh)["files"]

    def test_ledger_covers_all_owned_templates(self):
        for name in OWNED_TEMPLATES:
            self.assertIn(name, self.ledger)

    def test_files_unchanged_since_measurement(self):
        for name, entry in self.ledger.items():
            path = ASSETS / name
            with self.subTest(template=name):
                self.assertTrue(path.exists(), name)
                actual = hashlib.sha256(path.read_bytes()).hexdigest()
                self.assertEqual(
                    actual, entry["sha256"],
                    f"{name}: файл изменился после замера страниц в Word "
                    f"({entry['pages']} стр.) -- обнови word-pages.json и "
                    "перезамерь через Word COM",
                )


# ---------------------------------------------------------------------------
# C. clean_docx_metadata.py
# ---------------------------------------------------------------------------

def _text_el(tag, text=None, **attrs):
    el = ET.Element(tag)
    for k, v in attrs.items():
        el.set(k, v)
    if text is not None:
        el.text = text
    return el


def _make_dirty_docx(path):
    """Строит .docx с комментарием, принятыми/непринятыми правками,
    custom.xml и Company в app.xml -- через прямую сборку XML, как просит
    задание ("создай через XML"), без Word."""
    doc = Document()
    doc.add_paragraph("Титульный абзац.")
    doc.add_paragraph("Первый абзац без изменений.")
    doc.add_paragraph("Абзац с комментарием научного руководителя.")
    doc.add_paragraph("Последний абзац документа.")
    doc.core_properties.author = "Иванов Иван Иванович"
    doc.core_properties.last_modified_by = "Иванов Иван Иванович"
    doc.core_properties.keywords = "диплом, черновик"
    doc.core_properties.revision = 17
    doc.save(str(path))

    with zipfile.ZipFile(str(path), "r") as zin:
        names = zin.infolist()
        parts = {zi.filename: zin.read(zi.filename) for zi in names}
        order = [zi.filename for zi in names]

    W = cleaner.W_NS
    # Части собираются namespace-сохраняющим разбором/сериализацией самого
    # скрипта: ET.tostring теряет объявления префиксов из mc:Ignorable, и
    # такой «грязный» вход Word уже не открыл бы (дефект W0).
    parse, serialize = cleaner._parse_xml, cleaner._serialize_xml

    # --- document.xml: комментарий + принятые/непринятые правки ---------
    doc_root, doc_ns = parse(parts["word/document.xml"])
    body = doc_root.find("{%s}body" % W)
    paragraphs = body.findall("{%s}p" % W)
    target = None
    for p in paragraphs:
        text = "".join(t.text or "" for t in p.iter("{%s}t" % W))
        if "комментарием" in text:
            target = p
            break
    assert target is not None
    run = target.find("{%s}r" % W)
    ref_run = ET.Element("{%s}r" % W)
    ref_run.append(_text_el("{%s}commentReference" % W, **{"{%s}id" % W: "0"}))
    children = list(target)
    run_pos = children.index(run)
    new_children = (
        children[:run_pos]
        + [_text_el("{%s}commentRangeStart" % W, **{"{%s}id" % W: "0"}), run,
           _text_el("{%s}commentRangeEnd" % W, **{"{%s}id" % W: "0"}), ref_run]
        + children[run_pos + 1:]
    )
    for c in list(target):
        target.remove(c)
    for c in new_children:
        target.append(c)

    revision_p = ET.Element("{%s}p" % W)
    ins = ET.SubElement(revision_p, "{%s}ins" % W, {
        "{%s}id" % W: "1", "{%s}author" % W: "Иванов И.И.", "{%s}date" % W: "2024-01-01T00:00:00Z"})
    ins_run = ET.SubElement(ins, "{%s}r" % W)
    ins_run.append(_text_el("{%s}t" % W, "Новое предложение, вставленное с правками. "))
    delete = ET.SubElement(revision_p, "{%s}del" % W, {
        "{%s}id" % W: "2", "{%s}author" % W: "Иванов И.И.", "{%s}date" % W: "2024-01-01T00:00:00Z"})
    del_run = ET.SubElement(delete, "{%s}r" % W)
    del_run.append(_text_el("{%s}delText" % W, "устаревший текст, который удалили."))
    body.insert(list(body).index(target) + 1, revision_p)

    for i, p in enumerate(body.findall("{%s}p" % W)[:2]):
        p.set("{%s}rsidR" % W, "00AB12C%d" % i)
        p.set("{%s}rsidRDefault" % W, "00AB12C%d" % i)

    parts["word/document.xml"] = serialize(doc_root, doc_ns)

    # --- settings.xml: trackChanges + rsids ------------------------------
    settings_root, settings_ns = parse(parts["word/settings.xml"])
    settings_root.insert(0, ET.Element("{%s}trackChanges" % W))
    # python-docx's own default settings.xml already ships a <w:rsids> list
    # (its own template's editing history) -- reuse it instead of adding a
    # second, schema-invalid sibling.
    rsids = settings_root.find("{%s}rsids" % W)
    if rsids is None:
        rsids = ET.SubElement(settings_root, "{%s}rsids" % W)
    ET.SubElement(rsids, "{%s}rsid" % W, {"{%s}val" % W: "00AB12C1"})
    parts["word/settings.xml"] = serialize(settings_root, settings_ns)

    # --- app.xml: Company/Manager/Template/TotalTime/HyperlinkBase ------
    app_root, app_ns = parse(parts["docProps/app.xml"])
    EP = cleaner.EP_NS
    for local, value in (
        ("Company", "ООО Ромашка"),
        ("Manager", "П.С. Петров"),
        ("Template", "Diplom-MPGU.dotm"),
        ("TotalTime", "245"),
    ):
        el = app_root.find("{%s}%s" % (EP, local))
        if el is None:
            el = ET.SubElement(app_root, "{%s}%s" % (EP, local))
        el.text = value
    hb = app_root.find("{%s}HyperlinkBase" % EP)
    if hb is None:
        hb = ET.SubElement(app_root, "{%s}HyperlinkBase" % EP)
    hb.text = "https://example.invalid/"
    parts["docProps/app.xml"] = serialize(app_root, app_ns)

    # --- docProps/custom.xml + связь + Override --------------------------
    VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"
    custom_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="%s">'
        '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="StudentId">'
        '<vt:lpwstr>stud-12345</vt:lpwstr></property></Properties>' % VT
    ).encode("utf-8")
    parts["docProps/custom.xml"] = custom_xml
    order.append("docProps/custom.xml")

    rels_root, rels_ns = parse(parts["_rels/.rels"])
    PR = cleaner.PR_NS
    rel = ET.SubElement(rels_root, "{%s}Relationship" % PR)
    rel.set("Id", "rIdCustom1")
    rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties")
    rel.set("Target", "docProps/custom.xml")
    parts["_rels/.rels"] = serialize(rels_root, rels_ns)

    ct_root, ct_ns = parse(parts["[Content_Types].xml"])
    CT = cleaner.CT_NS
    override = ET.SubElement(ct_root, "{%s}Override" % CT)
    override.set("PartName", "/docProps/custom.xml")
    override.set("ContentType", "application/vnd.openxmlformats-officedocument.custom-properties+xml")

    # --- word/comments.xml + commentsExtended/Ids + people.xml -----------
    comments_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
        '<w:comments xmlns:w="%s">'
        '<w:comment w:id="0" w:author="П.С. Петров" w:date="2024-01-01T00:00:00Z" w:initials="ПП">'
        '<w:p><w:r><w:t>Комментарий научрука: переписать этот абзац.</w:t></w:r></w:p>'
        '</w:comment></w:comments>' % W
    ).encode("utf-8")
    parts["word/comments.xml"] = comments_xml
    order.append("word/comments.xml")

    for extra_name, extra_xml in (
        ("word/commentsExtended.xml",
         '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
         '<w15:commentsEx xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
         '</w15:commentsEx>'),
        ("word/commentsIds.xml",
         '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
         '<w16cid:commentsIds xmlns:w16cid="http://schemas.microsoft.com/office/word/2016/wordml/cid">'
         '</w16cid:commentsIds>'),
        ("word/people.xml",
         '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
         '<w15:people xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">'
         '<w15:person w15:author="П.С. Петров"/></w15:people>'),
    ):
        parts[extra_name] = extra_xml.encode("utf-8")
        order.append(extra_name)

    doc_rels_root, doc_rels_ns = parse(parts["word/_rels/document.xml.rels"])
    for i, (rid, target, rtype) in enumerate((
        ("rIdComments1", "comments.xml",
         "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"),
        ("rIdComments2", "commentsExtended.xml",
         "http://schemas.microsoft.com/office/2011/relationships/commentsExtended"),
        ("rIdComments3", "commentsIds.xml",
         "http://schemas.microsoft.com/office/2016/09/relationships/commentsIds"),
        ("rIdPeople1", "people.xml",
         "http://schemas.microsoft.com/office/2011/relationships/people"),
    )):
        rel = ET.SubElement(doc_rels_root, "{%s}Relationship" % PR)
        rel.set("Id", rid)
        rel.set("Type", rtype)
        rel.set("Target", target)
    parts["word/_rels/document.xml.rels"] = serialize(doc_rels_root, doc_rels_ns)

    for partname, ctype in (
        ("/word/comments.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"),
        ("/word/commentsExtended.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"),
        ("/word/commentsIds.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsIds+xml"),
        ("/word/people.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.people+xml"),
    ):
        override = ET.SubElement(ct_root, "{%s}Override" % CT)
        override.set("PartName", partname)
        override.set("ContentType", ctype)
    parts["[Content_Types].xml"] = serialize(ct_root, ct_ns)

    with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as zout:
        for name in order:
            data = parts[name]
            if isinstance(data, str):
                data = data.encode("utf-8")
            zout.writestr(name, data)

    return path


class DirtyFixtureCleanTests(unittest.TestCase):

    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.addCleanup(self._td.cleanup)
        self.src = Path(self._td.name) / "dirty.docx"
        _make_dirty_docx(self.src)

    def _clean(self, **kwargs):
        out = Path(self._td.name) / "clean.docx"
        rc = cleaner.main([str(self.src), "-o", str(out)] + kwargs.pop("extra_args", []))
        return rc, out

    def test_dirty_fixture_is_detected_by_check(self):
        rc = cleaner.main([str(self.src), "--check"])
        self.assertEqual(rc, 1)

    def test_dirty_fixture_is_a_valid_package(self):
        # «Грязный» вход должен быть целым для Word (W0): иначе самопроверка
        # очистки законно откажется записывать результат.
        self.assertEqual([], integrity.docx_package_problems(self.src))

    def test_cleaned_result_is_a_valid_package(self):
        rc, out = self._clean(extra_args=["--strip-rsid"])
        self.assertEqual(rc, 0)
        self.assertEqual([], integrity.docx_package_problems(out))

    def test_clean_removes_everything(self):
        rc, out = self._clean()
        self.assertEqual(rc, 0)
        with zipfile.ZipFile(str(out)) as z:
            names = set(z.namelist())
            for gone in ("docProps/custom.xml", "word/comments.xml",
                         "word/commentsExtended.xml", "word/commentsIds.xml",
                         "word/people.xml"):
                self.assertNotIn(gone, names)
            core = z.read("docProps/core.xml").decode("utf-8")
            self.assertNotIn("Иванов", core)
            self.assertIn("<cp:revision>1</cp:revision>", core)
            app = z.read("docProps/app.xml").decode("utf-8")
            self.assertNotIn("Ромашка", app)
            self.assertNotIn("Петров", app)
            self.assertNotIn("HyperlinkBase", app)
            self.assertIn("<Template>Normal.dotm</Template>", app)
            self.assertIn("<TotalTime>0</TotalTime>", app)
            settings = z.read("word/settings.xml").decode("utf-8")
            self.assertNotIn("trackChanges", settings)
            document = z.read("word/document.xml").decode("utf-8")
            for gone_tag in ("commentRangeStart", "commentRangeEnd", "commentReference",
                              "w:ins", "w:del", "w:delText"):
                self.assertNotIn(gone_tag, document, gone_tag)

        # Rels and content types no longer mention the removed parts.
        with zipfile.ZipFile(str(out)) as z:
            ct = z.read("[Content_Types].xml").decode("utf-8")
            self.assertNotIn("custom.xml", ct)
            self.assertNotIn("comments", ct)
            rels = z.read("_rels/.rels").decode("utf-8")
            self.assertNotIn("custom.xml", rels)

    def test_accepted_insertion_text_kept_deletion_dropped(self):
        rc, out = self._clean()
        self.assertEqual(rc, 0)
        doc = Document(str(out))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Новое предложение, вставленное с правками.", full_text)
        self.assertNotIn("устаревший текст, который удалили.", full_text)

    def test_result_opens_with_python_docx(self):
        rc, out = self._clean()
        self.assertEqual(rc, 0)
        doc = Document(str(out))
        self.assertGreaterEqual(len(doc.paragraphs), 4)

    def test_strip_rsid_flag(self):
        out_no = Path(self._td.name) / "no_strip.docx"
        out_yes = Path(self._td.name) / "with_strip.docx"
        self.assertEqual(cleaner.main([str(self.src), "-o", str(out_no)]), 0)
        self.assertEqual(cleaner.main([str(self.src), "-o", str(out_yes), "--strip-rsid"]), 0)
        with zipfile.ZipFile(str(out_no)) as z:
            self.assertIn(b'w:rsidR', z.read("word/document.xml"))
        with zipfile.ZipFile(str(out_yes)) as z:
            document = z.read("word/document.xml")
            settings = z.read("word/settings.xml")
            self.assertNotIn(b"w:rsidR=", document)
            self.assertNotIn(b"<w:rsids>", settings)

    def test_check_on_cleaned_file_is_clean(self):
        rc, out = self._clean()
        self.assertEqual(rc, 0)
        rc2 = cleaner.main([str(out), "--check"])
        self.assertEqual(rc2, 0)

    def test_json_report_is_valid_json_only(self):
        out = Path(self._td.name) / "clean_json.docx"
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = cleaner.main([str(self.src), "-o", str(out), "--json"])
        self.assertEqual(rc, 0)
        payload = json.loads(buf.getvalue())
        self.assertIn("core_fields_cleared", payload)
        self.assertFalse(payload["deterministic"])

    def test_check_json_report(self):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = cleaner.main([str(self.src), "--check", "--json"])
        self.assertEqual(rc, 1)
        payload = json.loads(buf.getvalue())
        self.assertTrue(payload["dirty"])
        self.assertIn("word/comments.xml", payload["comment_parts_present"])


class ModesAndExitCodeTests(unittest.TestCase):

    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.addCleanup(self._td.cleanup)
        self.src = Path(self._td.name) / "dirty.docx"
        _make_dirty_docx(self.src)

    def test_missing_file_exit_2(self):
        rc = cleaner.main([str(Path(self._td.name) / "nope.docx")])
        self.assertEqual(rc, 2)

    def test_bad_zip_exit_2(self):
        bad = Path(self._td.name) / "bad.docx"
        bad.write_bytes(b"not a zip at all")
        rc = cleaner.main([str(bad)])
        self.assertEqual(rc, 2)

    def test_not_a_docx_zip_exit_2(self):
        not_docx = Path(self._td.name) / "notdocx.docx"
        with zipfile.ZipFile(str(not_docx), "w") as z:
            z.writestr("hello.txt", "hi")
        rc = cleaner.main([str(not_docx)])
        self.assertEqual(rc, 2)

    def test_in_place_and_output_conflict_exit_2(self):
        with self.assertRaises(SystemExit) as ctx:
            cleaner.main([str(self.src), "--in-place", "-o", str(self.src)])
        self.assertEqual(ctx.exception.code, 2)

    def test_check_and_output_conflict_exit_2(self):
        with self.assertRaises(SystemExit) as ctx:
            cleaner.main([str(self.src), "--check", "-o", "x.docx"])
        self.assertEqual(ctx.exception.code, 2)

    def test_default_output_refused_inside_final_dir(self):
        final_dir = Path(self._td.name) / "final"
        final_dir.mkdir()
        target = final_dir / "vkr.docx"
        target.write_bytes(self.src.read_bytes())
        with self.assertRaises(SystemExit) as ctx:
            cleaner.main([str(target)])
        self.assertEqual(ctx.exception.code, 2)
        self.assertEqual(["vkr.docx"], sorted(p.name for p in final_dir.iterdir()))

    def test_in_place_mode(self):
        before = self.src.read_bytes()
        rc = cleaner.main([str(self.src), "--in-place"])
        self.assertEqual(rc, 0)
        after = self.src.read_bytes()
        self.assertNotEqual(before, after)
        doc = Document(str(self.src))
        self.assertEqual(doc.core_properties.author, "")

    def test_default_output_name(self):
        rc = cleaner.main([str(self.src)])
        self.assertEqual(rc, 0)
        expected = self.src.with_name(self.src.stem + "-clean" + self.src.suffix)
        self.assertTrue(expected.exists())

    def test_output_same_path_different_case(self):
        # Windows: то же самое имя файла, но другой регистр -- атомарная
        # замена через temp + os.replace должна сработать без SameFileError.
        if os.name != "nt":
            self.skipTest("регистронезависимость файловой системы проверяем только на Windows")
        upper = Path(str(self.src).upper())
        rc = cleaner.main([str(self.src), "-o", str(upper)])
        self.assertEqual(rc, 0)
        doc = Document(str(self.src))
        self.assertEqual(doc.core_properties.author, "")

    def test_help_without_python_docx_dependency(self):
        # Скрипт не импортирует python-docx вовсе -- --help должен работать
        # даже если бы его не было. Проверяем напрямую: модуль не содержит
        # top-level "import docx"/"from docx".
        src_text = (SKILL / "scripts" / "clean_docx_metadata.py").read_text(encoding="utf-8")
        self.assertNotRegex(src_text, re.compile(r"^\s*(import docx\b|from docx\b)", re.MULTILINE))
        with self.assertRaises(SystemExit) as ctx:
            cleaner.main(["--help"])
        self.assertEqual(ctx.exception.code, 0)

    def test_help_examples_use_final_docx_and_exports(self):
        # 6.33: сдаётся единственный final/vkr.docx; копии очистки — только в exports/.
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer), self.assertRaises(SystemExit):
            cleaner.main(["--help"])
        text = buffer.getvalue()
        self.assertIn("clean_docx_metadata.py final/vkr.docx --in-place", text)
        self.assertRegex(text, r"clean_docx_metadata\.py final/vkr\.docx -o exports/\S+\.docx")
        self.assertNotIn("vkr-final", text)
        for line in text.splitlines():
            if "clean_docx_metadata.py " in line and " -o " in line:
                self.assertIn("-o exports/", line)


class DeterminismTests(unittest.TestCase):

    def setUp(self):
        self._td = tempfile.TemporaryDirectory()
        self.addCleanup(self._td.cleanup)
        self.src = Path(self._td.name) / "dirty.docx"
        _make_dirty_docx(self.src)

    def _sha(self, path):
        return hashlib.sha256(Path(path).read_bytes()).hexdigest()

    def test_same_seed_same_input_byte_identical(self):
        out1 = Path(self._td.name) / "seedA_1.docx"
        out2 = Path(self._td.name) / "seedA_2.docx"
        self.assertEqual(cleaner.main([str(self.src), "-o", str(out1), "--seed", "s1"]), 0)
        self.assertEqual(cleaner.main([str(self.src), "-o", str(out2), "--seed", "s1"]), 0)
        self.assertEqual(self._sha(out1), self._sha(out2))

    def test_seed_value_does_not_invent_dates(self):
        # Значение seed не порождает дат: результат зависит только от входного документа.
        out1 = Path(self._td.name) / "seedA.docx"
        out2 = Path(self._td.name) / "seedB.docx"
        self.assertEqual(cleaner.main([str(self.src), "-o", str(out1), "--seed", "s1"]), 0)
        self.assertEqual(cleaner.main([str(self.src), "-o", str(out2), "--seed", "s2"]), 0)
        self.assertEqual(self._sha(out1), self._sha(out2))

    def test_without_seed_is_reported_as_nondeterministic(self):
        # Прямая проверка контракта на уровне report, а не сравнение SHA
        # двух прогонов: без --seed created/modified берутся из реального
        # текущего времени с разрешением в 1 с -- два быстрых подряд запуска
        # могут попасть в одну и ту же секунду, так что сравнение хешей было
        # бы нестабильным тестом.
        parts, _infolist = cleaner.read_docx(self.src)
        _, report = cleaner.clean_docx_parts(parts)
        self.assertFalse(report["deterministic"])

        out = Path(self._td.name) / "plain.docx"
        self.assertEqual(cleaner.main([str(self.src), "-o", str(out)]), 0)

    def _core_dates(self, path):
        with zipfile.ZipFile(str(path)) as archive:
            root = ET.fromstring(archive.read("docProps/core.xml"))
        created = root.find("{%s}created" % cleaner.DCTERMS_NS).text
        modified = root.find("{%s}modified" % cleaner.DCTERMS_NS).text
        return created, modified

    def test_created_equals_modified_now_without_seed(self):
        # F14/B8: дата создания не сдвигается в прошлое (раньше -- на 1-48 часов).
        import datetime

        out = Path(self._td.name) / "now.docx"
        before = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer):
            self.assertEqual(cleaner.main([str(self.src), "-o", str(out), "--json"]), 0)
        after = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
        created, modified = self._core_dates(out)
        self.assertEqual(created, modified)
        stamp = datetime.datetime.strptime(created, "%Y-%m-%dT%H:%M:%SZ")
        self.assertTrue(before - datetime.timedelta(seconds=2) <= stamp <= after + datetime.timedelta(seconds=2), (before, stamp, after))
        report = json.loads(buffer.getvalue())
        self.assertEqual(report["created"], report["modified"])

    def test_created_equals_modified_with_seed(self):
        outputs = []
        for name in ("seeded_1.docx", "seeded_2.docx"):
            out = Path(self._td.name) / name
            self.assertEqual(cleaner.main([str(self.src), "-o", str(out), "--seed", "s1"]), 0)
            outputs.append(self._core_dates(out))
        self.assertEqual(outputs[0], outputs[1])
        self.assertEqual(outputs[0][0], outputs[0][1])
        parts, _infolist = cleaner.read_docx(self.src)
        modified, created = cleaner._deterministic_times(parts)
        self.assertEqual(modified, created)
        if modified is not None:
            # дата взята из самого документа, а не вычислена из seed
            self.assertEqual(outputs[0][0], cleaner._fmt_z(modified))
        # документ без даты: даты не выдумываются
        self.assertEqual((None, None), cleaner._deterministic_times({"docProps/core.xml": b"<cp:coreProperties/>"}))

    def test_help_and_docs_do_not_promise_created_shift(self):
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer), self.assertRaises(SystemExit):
            cleaner.main(["--help"])
        text = buffer.getvalue()
        self.assertIn("created = modified", text)
        for needle in ("1-48", "1–48", "сдвигается на"):
            self.assertNotIn(needle, text)
        for relative in ("references/friend-sharing.md", "references/mpgu-formatting.md", "SKILL.md",
                         "references/quickstart.md", "references/submission-checklist.md"):
            doc = (SKILL / relative).read_text(encoding="utf-8")
            for needle in ("1-48 час", "1–48 час", "created` назад", "сдвигает `created`"):
                self.assertNotIn(needle, doc, relative)


class RepeatedCleanTests(unittest.TestCase):
    """T6: повторная очистка уже чистого файла не переписывает байты (иначе — AUDIT_SNAPSHOT_STALE)."""

    WORD_SAVED = ROOT / "tests" / "fixtures" / "build" / "vkr-word-saved.docx"

    def setUp(self):
        self._td = tempfile.TemporaryDirectory(prefix="vkr633-repeat-clean-")
        self.addCleanup(self._td.cleanup)
        self.tmp = Path(self._td.name)

    def _run(self, *args):
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer), contextlib.redirect_stderr(io.StringIO()):
            code = cleaner.main([str(a) for a in args])
        return code, buffer.getvalue()

    def test_second_in_place_run_changes_nothing(self):
        # Файл после Word (creator/lastModifiedBy заполнены) — первая очистка обязана его чистить.
        target = self.tmp / "vkr.docx"
        target.write_bytes(self.WORD_SAVED.read_bytes())
        with zipfile.ZipFile(str(target)) as archive:
            core = archive.read("docProps/core.xml").decode("utf-8")
        self.assertIn("Nikolay Skorikov", core, "фикстура должна быть «после Word»")
        before = target.read_bytes()

        code, output = self._run(target, "--in-place", "--json")
        self.assertEqual(0, code, output)
        first = json.loads(output)
        self.assertEqual("cleaned", first["status"])
        self.assertIn("lastModifiedBy", first["core_fields_cleared"])
        after_first = target.read_bytes()
        self.assertNotEqual(before, after_first, "первая очистка обязана переписать файл")

        code, output = self._run(target, "--in-place", "--json")
        self.assertEqual(0, code, output)
        second = json.loads(output)
        self.assertEqual("already_clean", second["status"], second)
        self.assertEqual(after_first, target.read_bytes(), "повторная очистка не должна менять байты")
        self.assertEqual([], integrity.docx_package_problems(target))
        # --check на очищенном файле подтверждает: чистить нечего.
        self.assertEqual(0, self._run(target, "--check", "--json")[0])
        # человекочитаемый вывод объясняет, почему файл не переписан
        code, text = self._run(target, "--in-place")
        self.assertEqual(0, code)
        self.assertIn("чистить нечего", text)
        self.assertIn("snapshot", text)

    def test_copy_to_another_path_is_still_written(self):
        target = self.tmp / "vkr.docx"
        target.write_bytes(self.WORD_SAVED.read_bytes())
        self.assertEqual(0, self._run(target, "--in-place", "--json")[0])
        copy = self.tmp / "exports" / "vkr-print.docx"
        code, output = self._run(target, "-o", copy, "--json")
        self.assertEqual(0, code, output)
        self.assertEqual("cleaned", json.loads(output)["status"])
        self.assertTrue(copy.is_file())

    def test_differing_dates_are_still_normalized(self):
        # created != modified — очистке есть что сделать даже при пустых полях.
        target = self.tmp / "dates.docx"
        target.write_bytes(self.WORD_SAVED.read_bytes())
        self.assertEqual(0, self._run(target, "--in-place", "--json")[0])
        parts, infolist = cleaner.read_docx(target)
        core = parts["docProps/core.xml"].decode("utf-8")
        parts["docProps/core.xml"] = re.sub(
            r"(<dcterms:created[^>]*>)[^<]*(<)", r"\g<1>2020-01-01T00:00:00Z\g<2>", core).encode("utf-8")
        cleaner.write_docx(target, parts, infolist, False, None)
        code, output = self._run(target, "--in-place", "--json")
        self.assertEqual(0, code, output)
        report = json.loads(output)
        self.assertEqual("cleaned", report["status"], report)
        self.assertEqual(report["created"], report["modified"])

    def test_future_modified_date_is_reported_with_seed(self):
        # T7: дата из будущего не меняется, но о ней предупреждают.
        target = self.tmp / "future.docx"
        parts, infolist = cleaner.read_docx(self.WORD_SAVED)
        core = parts["docProps/core.xml"].decode("utf-8")
        parts["docProps/core.xml"] = re.sub(
            r"(<dcterms:modified[^>]*>)[^<]*(<)", r"\g<1>2099-12-31T23:59:00Z\g<2>", core).encode("utf-8")
        cleaner.write_docx(target, parts, infolist, False, None)
        code, output = self._run(target, "--in-place", "--seed", "s1", "--json")
        self.assertEqual(0, code, output)
        report = json.loads(output)
        self.assertTrue(any("больше текущего времени" in w for w in report["warnings"]), report["warnings"])
        self.assertEqual("2099-12-31T23:59:00Z", report["modified"], "дату не меняем, только предупреждаем")
        self.assertEqual(report["created"], report["modified"])
        code, text = self._run(target, "-o", self.tmp / "future-copy.docx", "--seed", "s1")
        self.assertIn("ПРЕДУПРЕЖДЕНИЕ", text)
        # обычная дата предупреждения не даёт
        clean_report = json.loads(self._run(self.WORD_SAVED, "-o", self.tmp / "ok.docx", "--seed", "s1", "--json")[1])
        self.assertEqual([], clean_report["warnings"])


class CompatShimTests(unittest.TestCase):
    """Старый API 6.31/6.32, всё ещё вызывается из tests/test_audit_fixes_632.py
    (cleaner.clean_metadata(path, path, seed=...)) -- должен продолжать работать."""

    def test_clean_metadata_in_place_same_as_old_api(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "document.docx"
            doc = Document()
            doc.core_properties.author = "Author"
            doc.add_paragraph("Text")
            doc.save(str(path))
            cleaner.clean_metadata(path, path, seed="test")
            cleaned = Document(str(path))
            self.assertEqual("", cleaned.core_properties.author)


if __name__ == "__main__":
    unittest.main()
